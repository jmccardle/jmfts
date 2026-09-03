"""A workbook in, one object per row out. ``INGEST_SPEC.md`` 8.4's ``records`` shape.

Two modes, and the difference between them is the point of the script.

``--read`` (the default) needs no database, no model and no queue. It runs the two pure
pieces — :func:`jmfts_core.office.cells.read_rows` and
:func:`jmfts_core.sheet_records.build_records` — over a file on disk and prints what comes
out. That is the transformation itself, and it works on any machine with the ``office``
extra installed.

``--ingest`` runs the real path: upload the file, drain the queue with the real worker, and
read the record nodes back out of the tree. It needs the appliance's database and it writes
to it. What it prints is what a retrieval hit on one of those rows would return.

Output is JSON on stdout and progress on stderr, so it composes::

    python -m scripts.demo_sheet_records book.xlsx | jq '.sheets[].records[0].record'
    python -m scripts.demo_sheet_records book.xlsx --sheet Deals --limit 3

Research code: ``scripts/`` is outside the lint gate on purpose (see CLAUDE.md).
"""

from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _note(message: str) -> None:
    """Progress goes to stderr so that stdout stays a JSON document."""
    print(message, file=sys.stderr)


# ---------------------------------------------------------------------------
# A workbook to demonstrate on, when the caller has none to hand
# ---------------------------------------------------------------------------


def sample_workbook() -> bytes:
    """A workbook holding every case the reader has a rule for.

    Generated rather than committed, and the caller is TOLD it is a sample — a demo that
    quietly invented its own input would be showing that the code runs, not that it works
    on anything.
    """
    import datetime

    import openpyxl

    workbook = openpyxl.Workbook()
    deals = workbook.active
    deals.title = "Deals"
    deals.append(["Deal ID", "Account", "Value", "Close Date", "Won", "Region"])
    deals.append(["D-4471", "Northwind Freight", 128000, datetime.datetime(2026, 9, 30), True, "NE"])
    deals.append(["D-4472", "Contoso", 96500.5, datetime.datetime(2026, 10, 15), False, "NW"])
    deals.append([None, None, None, None, None, None])
    # A leading apostrophe in Excel. The value keeps its leading zeros either way; the flag
    # is the author saying this is an identifier and not a number.
    deals["A5"] = "0012345"
    deals["A5"].quotePrefix = True
    deals["B5"] = "Fabrikam"
    deals["C5"] = "=C2+C3"
    deals["E5"] = True
    deals["F5"] = "SE"

    lookup = workbook.create_sheet("Regions")
    lookup.append(["Code", "Region"])
    for code, name in (("NE", "Northeast"), ("NW", "Northwest"), ("SE", "Southeast")):
        lookup.append([code, name])

    notes = workbook.create_sheet("Notes")
    notes["B3"] = "Free text with no header row above it, so this sheet yields no records."

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# --read: the transformation, with nothing else attached
# ---------------------------------------------------------------------------


def read_mode(data: bytes, *, sheets: list, max_rows: int, limit) -> dict:
    from jmfts_core.config import get_settings
    from jmfts_core.office.cells import read_rows
    from jmfts_core.office.sheets import measure_sheet
    from jmfts_core.office.workbook import read_sheets
    from jmfts_core.sheet_records import NO_HEADER_REASON, build_records

    budget = get_settings().embedding_doc_window
    output = []
    for sheet in read_sheets(data):
        if sheets and sheet.name not in sheets:
            continue
        _note(f"  {sheet.name}: measuring")
        # The same order the queue runs them in, and for the same reason: the header
        # labels come from what the profile measured, never from a second look at row 1.
        measurement = measure_sheet(
            data, sheet.name, render_cell_budget=budget, with_sketches=False
        )
        if not measurement.header_row.verdict:
            _note(f"  {sheet.name}: no header row, so no records")
            output.append(
                {"sheet": sheet.name, "records": [], "no_records": NO_HEADER_REASON}
            )
            continue

        rows = read_rows(data, sheet.name, max_rows=max_rows)
        records = build_records(rows, header=[column.name for column in measurement.columns])
        _note(f"  {sheet.name}: {len(records)} record(s) from {len(rows.rows)} row(s)")
        output.append(
            {
                "sheet": sheet.name,
                "columns": [column.name for column in measurement.columns],
                "rows_read": len(rows.rows),
                "records": [_as_json(record) for record in records[:limit]],
            }
        )
    return {"mode": "read", "sheets": output}


def _as_json(record) -> dict:
    stored = {"row_index": record.row_index, "record": record.record, "content": record.content}
    if record.cells:
        stored["cells"] = record.cells
    return stored


# ---------------------------------------------------------------------------
# --ingest: the real path, through the queue
# ---------------------------------------------------------------------------


def ingest_mode(data: bytes, *, filename: str, limit) -> dict:
    from sqlalchemy import select

    import jmfts_core.ingest_tasks  # noqa: F401  (registers the handlers)
    import jmfts_core.sheet_tasks  # noqa: F401
    from jmfts_client.contracts.upload import UploadedFile
    from jmfts_core.config import get_settings
    from jmfts_core.database import get_session
    from jmfts_core.ingest_worker import IngestWorker
    from jmfts_core.models.document import Document
    from jmfts_core.repositories.evidence import EvidenceRepository
    from jmfts_core.services.ingest_service import IngestService
    from jmfts_core.settling import NO_ROLLUP
    from jmfts_core.models.document import USETYPE_RECORD, USETYPE_SHEET

    settings = get_settings()
    # Said out loud before anything is written. This mode writes to whatever JMFTS_DB_*
    # points at, and a demo that turned out to have been running against the appliance's
    # real database is not a surprise anybody should get afterwards.
    _note(
        f"  writing to database {settings.db_name!r} on "
        f"{settings.db_host}:{settings.db_port}"
    )

    with get_session() as session:
        response = IngestService(session).upload_file(
            UploadedFile(data=data, filename=filename, content_type=XLSX_MIME)
        )
        session.commit()
        _note(f"  uploaded as document {response.document_id}")

        worker = IngestWorker(
            worker_id="demo-sheet-records",
            session_factory=get_session,
            planner=NO_ROLLUP,
        )
        ran = worker.drain(max_tasks=10_000)
        _note(f"  drained {ran} task(s)")

        sheets = (
            session.execute(
                select(Document)
                .where(
                    Document.parent_id == response.document_id,
                    Document.usetype == USETYPE_SHEET,
                )
                .order_by(Document.id)
            )
            .scalars()
            .all()
        )

        # Evidence is rows since `SPRINT_JOBS.md` Phase 2b, so one read per node here
        # rather than an attribute access. `read_many` is one query for the whole set.
        found = EvidenceRepository(session).read_many([s.id for s in sheets])
        output = []
        for sheet in sheets:
            block = found.get(sheet.id, {}).get("sheet") or {}
            records = (
                session.execute(
                    select(Document)
                    .where(Document.parent_id == sheet.id, Document.usetype == USETYPE_RECORD)
                    .order_by(Document.id)
                )
                .scalars()
                .all()
            )
            _note(f"  {block.get('name')}: {len(records)} record node(s)")
            output.append(
                {
                    "sheet": block.get("name"),
                    "sheet_node_id": sheet.id,
                    "shape": block.get("shape"),
                    "rung": block.get("rung"),
                    "records": [
                        {
                            "node_id": node.id,
                            "row_index": row["row_index"],
                            "record": row["record"],
                            "content": node.content,
                            **({"cells": row["cells"]} if "cells" in row else {}),
                        }
                        for node, row in (
                            (node, EvidenceRepository(session).read_all(node.id))
                            for node in records[:limit]
                        )
                    ],
                }
            )
        return {"mode": "ingest", "document_id": response.document_id, "sheets": output}


# ---------------------------------------------------------------------------


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "path",
        nargs="?",
        help="an .xlsx file; omit it to run on a generated sample workbook",
    )
    parser.add_argument(
        "--ingest",
        action="store_true",
        help="run the real path: upload, drain the queue, read the nodes back. Needs the "
        "database and WRITES to it.",
    )
    parser.add_argument(
        "--sheet",
        action="append",
        default=[],
        help="only this sheet; repeatable. --read only.",
    )
    parser.add_argument(
        "--max-rows",
        type=int,
        default=10_000,
        help="fail rather than truncate past this many rows per sheet (default: 10000)",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="print at most this many records per sheet; all of them by default",
    )
    args = parser.parse_args(argv)

    if args.path:
        source = Path(args.path)
        if not source.is_file():
            parser.error(f"{source} is not a file")
        data = source.read_bytes()
        filename = source.name
        _note(f"reading {source}")
    else:
        data = sample_workbook()
        filename = "sample.xlsx"
        _note("NO FILE GIVEN — running on a generated sample workbook, not real data.")

    if args.ingest:
        if args.sheet:
            parser.error("--sheet applies to the read path only; --ingest runs every sheet")
        result = ingest_mode(data, filename=filename, limit=args.limit)
    else:
        result = read_mode(data, sheets=args.sheet, max_rows=args.max_rows, limit=args.limit)

    json.dump(result, sys.stdout, indent=2, ensure_ascii=False)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
