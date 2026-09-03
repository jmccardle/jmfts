#!/usr/bin/env python3
"""The 0.3.0 job-system release, run rather than described.

``docs/SPRINT_JOBS.md`` phases 1, 2, 2a, 2b, 3 and 4 are what this release contains. Each
one is specified in that document and asserted by a test, and neither form SHOWS a
reader anything: a spec states an intent and a test states a pass or a fail. This script
runs the shipped code and prints what came back.

Seven sections, and the first six need no database, no model and no optional dependency
beyond ``openpyxl`` for the workbook one:

    python -m scripts.demo_release_0_3_0 list
    python -m scripts.demo_release_0_3_0 run [SECTION ...]
    python -m scripts.demo_release_0_3_0 markdown --out demo.md

``markdown`` writes the same output as fenced blocks, each preceded by the command that
produced it. That file is what the release page on the documentation site embeds, so the
page and the appliance cannot drift: regenerating it is how the page is updated.

``stamp`` is the one section that touches a database, and it touches the one
``JMFTS_DB_*`` names. Point it somewhere throwaway:

    docker run -d --name jmfts-demo-pg -e POSTGRES_USER=jmfts -e POSTGRES_PASSWORD=jmfts \\
        -e POSTGRES_DB=jmfts -p 127.0.0.1:5434:5432 pgvector/pgvector:pg16
    JMFTS_DB_PORT=5434 jmfts-init-db
    JMFTS_DB_PORT=5434 python -m scripts.demo_release_0_3_0 run stamp

WHAT THIS IS NOT. It is not a benchmark and it reports no timings — ``benchmarks/`` is
where measurement lives. It is not a test either: nothing here asserts, and a section
that printed something unexpected would print it rather than fail. What it is for is
showing that the six phases produce the behaviour the documents claim, in the reader's
own terminal, against files the reader can rebuild.
"""

from __future__ import annotations

import argparse
import contextlib
import io
import sys
from pathlib import Path
from typing import Callable, Optional

from jmfts_core.ingest_options import (
    OPTION_CHECKS,
    TASK_PARAM_DEFAULTS,
    resolve_options,
)
from jmfts_core.ingest_tasks import (
    CHAR_COUNT,
    GUARDABLE_PATTERNS,
    HAS_TEXT_LAYER,
    OP_GT,
    PATTERNS_NOT_PROBED,
    ROOT_SCOPE,
    SCOPE_CHILDREN,
    TASK_ROWS,
    TASK_STRUCTURE_SHEETS,
    Scope,
    TaskRow,
    children_of,
    explain_plan,
    option,
    plan_after_probe,
    term,
)
from jmfts_core.models.document import USETYPE_CHUNK, USETYPE_SHEET
from jmfts_core.probe import detect_format, probe_patterns

# Importing a handler module is what puts its rows' handlers in `TASK_HANDLERS` and its
# declarations in `ATOMS`. `ingest_tasks` holds the table but not the handlers.
import jmfts_core.citation_tasks  # noqa: F401,E402
import jmfts_core.embed_tasks  # noqa: F401,E402
import jmfts_core.rollup_tasks  # noqa: F401,E402
import jmfts_core.sheet_tasks  # noqa: F401,E402
import jmfts_core.structure_tasks  # noqa: F401,E402

# ---------------------------------------------------------------------------
# printing
# ---------------------------------------------------------------------------


def _rule(title: str) -> None:
    print(f"\n{title}")
    print("-" * len(title))


def _params(params: dict) -> str:
    return ", ".join(f"{k}={v!r}" for k, v in sorted(params.items())) or "—"


# ---------------------------------------------------------------------------
# fixtures — real bytes, built here, deterministic
# ---------------------------------------------------------------------------

NOTE_MD = """\
# Retrieval notes

Hybrid fuses vector and BM25 by rank, so what comes back is an ordering.

## Late interaction

MaxSim scores a query token against every document token and keeps the max.

## Segmentation

PELT finds change points in an ordered sequence of embeddings.
"""


def _markdown_bytes() -> bytes:
    return NOTE_MD.encode("utf-8")


def _workbook_bytes() -> Optional[bytes]:
    """A workbook written by the library that will read it back.

    ``tests/corpus`` commits no binaries and its ``.xlsx`` is the smallest package that
    satisfies tier 1 — ``probe`` reads the ZIP manifest and never opens a worksheet part.
    ``structure:sheets`` is tier 2 and calls ``openpyxl.load_workbook``, which wants the
    relationship parts that fixture leaves out.
    """
    try:
        import openpyxl
    except ImportError:
        return None

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Measurements"
    sheet.append(["run", "corpus", "recall@10", "seconds"])
    for row, corpus in enumerate(("multihop", "wiki", "arxiv", "email", "notes"), start=1):
        sheet.append([row, corpus, round(0.60 + row / 40, 3), 12 * row])
    book.create_sheet("Notes").append(["free text, no header row"])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _probe(data: bytes, filename: str) -> tuple[str, dict]:
    detection = detect_format(data, filename=filename)
    patterns, _detail = probe_patterns(data, detection)
    return detection.format, patterns


# ---------------------------------------------------------------------------
# 1. one ingest path
# ---------------------------------------------------------------------------


def demo_one_path() -> None:
    """`pipeline.py` is deleted; a plan is a pure function of what probe measured."""
    data = _markdown_bytes()
    fmt, patterns = _probe(data, "retrieval-notes.md")
    plan = plan_after_probe(fmt, patterns)

    _rule(f"probe measured {len(data)} bytes of retrieval-notes.md")
    print(f"  format    {fmt}")
    true_patterns = sorted(name for name, value in patterns.items() if value is True)
    print(f"  patterns  {', '.join(true_patterns) or '(none true)'}")

    _rule("What Part 4's table scheduled, at the file node")
    for spec in plan.eligible:
        line = f"  {spec.task_type:22} write {spec.write_mode}"
        if spec.after:
            line += f", after {', '.join(spec.after)}"
        print(line)
        if spec.params:
            print(f"  {'':22} params {_params(spec.params)}")

    _rule("What it deliberately did not schedule, and why")
    for task, reason in list(plan.not_applicable.items())[:4]:
        print(f"  {task:22} {reason}")
    remaining = len(plan.not_applicable) - 4
    if remaining > 0:
        print(f"  ... and {remaining} more rows, each with its own measured reason")

    print(
        "\n  Every row of the table is decided. A task missing from a plan would be a\n"
        "  wrong answer rather than a short one, so nothing is omitted for brevity here\n"
        "  except by this script."
    )


# ---------------------------------------------------------------------------
# 2. scope on a rule
# ---------------------------------------------------------------------------


def demo_scope() -> None:
    """Phase 3: a rule names WHICH nodes it applies to, including nodes that do not exist."""
    data = _markdown_bytes()
    fmt, patterns = _probe(data, "retrieval-notes.md")

    _rule("Every scope Part 4's table declares")
    for scope in dict.fromkeys(row.scope for row in TASK_ROWS):
        rows = [row.task for row in TASK_ROWS if row.scope == scope]
        print(f"  {scope}")
        print(f"      {', '.join(rows)}")

    _rule("The same table, asked at the file node and at a chunk")
    for scope in (ROOT_SCOPE, _chunk_scope()):
        plan = plan_after_probe(fmt, patterns, scope=scope)
        tasks = ", ".join(spec.task_type for spec in plan.eligible) or "(nothing)"
        print(f"  {scope}")
        print(f"      {tasks}")

    print(
        "\n  The chunk does not exist yet. `structure:declared` has not run, so there is no\n"
        "  node to name — and that is exactly why the row lives in the table rather than\n"
        "  in the handler that creates the node. Before Phase 3 a plan could not report it."
    )


def _chunk_scope() -> Scope:
    """The scope the `embed` row declares, read off the row rather than rebuilt."""
    for row in TASK_ROWS:
        if row.scope.kind == SCOPE_CHILDREN and USETYPE_CHUNK in row.scope.usetypes:
            return row.scope
    raise LookupError("no row is scoped to chunks")


# ---------------------------------------------------------------------------
# 3. EXPLAIN reaches the sheet tier
# ---------------------------------------------------------------------------


def demo_explain() -> None:
    """`EXPLAIN` used to stop at the sheet list, because the per-sheet tasks were not rows."""
    data = _workbook_bytes()
    if data is None:
        print('  openpyxl is not installed — `pip install -e ".[office]"`')
        return

    fmt, patterns = _probe(data, "measurements.xlsx")
    plan = explain_plan(fmt, patterns=patterns, patterns_source="probed")

    _rule(f"POST /ingest/explain — format {plan.format}, patterns {plan.patterns_source}")
    print(f"  {'task':22} {'outcome':12} scope")
    for task in plan.tasks:
        if task.outcome == "not_applicable":
            continue
        print(f"  {task.task:22} {task.outcome:12} {task.scope}")

    latent = [
        task
        for task in plan.tasks
        if task.scope != str(ROOT_SCOPE) and task.outcome != "not_applicable"
    ]
    _rule(f"The {len(latent)} rows that used to be invisible")
    for task in latent:
        print(f"  {task.task}")
        print(f"      scope   {task.scope}")
        print(f"      params  {_params(task.params)}")
    print(
        "\n  None of these three name a node that exists yet. `structure:sheets` has not\n"
        "  run, so there are no worksheets — and `embed` is one scope further down again,\n"
        "  on the records `extract:sheet` will write."
    )


# ---------------------------------------------------------------------------
# 4. the knobs
# ---------------------------------------------------------------------------


def demo_knobs() -> None:
    """A `params_key` on a row is what makes an option group reach a queue row."""
    _rule("Every option group, and which task reads it")
    readers: dict[str, list[str]] = {group: [] for group in TASK_PARAM_DEFAULTS}
    for row in TASK_ROWS:
        if row.params_key:
            readers[row.params_key].append(row.task)
    for group, tasks in sorted(readers.items()):
        print(f"  {group:16} {', '.join(tasks) or '(no row names it)'}")

    _rule("Setting one, and watching it reach the row")
    overrides = {"sheet_records": {"max_rows": 500}}
    resolved = resolve_options("xlsx", overrides)
    print(f"  override   {overrides}")
    print(f"  resolved   {_params(resolved['sheet_records'])}")

    data = _workbook_bytes()
    if data is not None:
        fmt, patterns = _probe(data, "measurements.xlsx")
        plan = explain_plan(fmt, options=overrides, patterns=patterns, patterns_source="probed")
        for task in plan.tasks:
            if task.task == "extract:sheet":
                print(f"  on the row {task.task}: {_params(task.params)}")

    _rule("And refusing a bad one")
    for bad in ({"sheet_records": {"max_rows": 0}}, {"sheet_records": {"max_rows": True}}):
        try:
            resolve_options("xlsx", bad)
        except ValueError as exc:
            print(f"  {bad}")
            print(f"      ValueError: {exc}")

    print(
        f"\n  {len(OPTION_CHECKS)} groups carry per-key checks, and `max_rows` reaching the\n"
        "  handler is what Phase 3 closed. The group was validated before this release; the\n"
        "  handler read its own literal default and never looked at the value."
    )


# ---------------------------------------------------------------------------
# 5. the audit
# ---------------------------------------------------------------------------


def demo_audit() -> None:
    """Part 14 forbids a second list. The table audits itself at import."""
    import jmfts_core.ingest_tasks as it

    _rule("Four mistakes the table refuses, at import time")

    sheets_row = next(row for row in TASK_ROWS if row.task == TASK_STRUCTURE_SHEETS)
    sheet_scope = children_of(TASK_STRUCTURE_SHEETS, usetypes=(USETYPE_SHEET,))

    mistakes: list[tuple[str, tuple[TaskRow, ...]]] = [
        (
            "a duplicate task name",
            (sheets_row, sheets_row),
        ),
        (
            "`after` naming a row that comes later",
            (
                TaskRow(task="a", write_mode="self", after=("b",)),
                TaskRow(task="b", write_mode="self"),
            ),
        ),
        (
            "a child scope naming a producer below it",
            (
                TaskRow(task="a", write_mode="self", scope=sheet_scope),
                sheets_row,
            ),
        ),
        (
            "a `params_key` naming no group",
            (TaskRow(task="a", write_mode="self", params_key="not_a_group"),),
        ),
    ]

    original = it.TASK_ROWS
    try:
        for label, rows in mistakes:
            it.TASK_ROWS = rows
            print(f"  {label}")
            try:
                it._check_task_rows()
            except Exception as exc:  # noqa: BLE001
                print(f"      {type(exc).__name__}: {exc}")
            else:
                print("      ACCEPTED — the audit did not catch this")
    finally:
        it.TASK_ROWS = original

    print(
        "\n  A test over a list only proves the list equals itself. These are stated over\n"
        "  the table the appliance actually runs, so deleting a row moves the audit with it."
    )


# ---------------------------------------------------------------------------
# 6. guards with operators
# ---------------------------------------------------------------------------


def demo_guards() -> None:
    """Phase 4: a guard is a comparison, its right side can be an option, and both lists
    read an unmeasured name differently."""
    _rule("Every pattern a guard may read, and what it holds")
    guarded_by = _guard_readers()
    for name, kind in sorted(GUARDABLE_PATTERNS.items()):
        mark = "   (probe does not write it)" if name in PATTERNS_NOT_PROBED else ""
        print(f"  {name:20} {kind:6}{mark}".rstrip())
        print(f"      {', '.join(guarded_by.get(name, ())) or '(no row reads it)'}")
    print(
        "\n  `matched.patterns` is an open namespace, and `evidence.pattern_type` answers\n"
        "  `bool` for a name it has never heard of. That is right for storage and wrong for\n"
        "  a guard: `has_hedings` would type-check, plan cleanly, and stand its row down on\n"
        "  every document forever."
    )

    _rule("The two unknown policies, on a .docx")
    docx = explain_plan("docx", patterns={"has_text_layer": True})
    print("  patterns supplied   has_text_layer=True")
    print("  has_heading_styles  guardable, and no prober emits it")
    print()
    for name in ("structure:declared", "structure:inferred"):
        task = next(t for t in docx.tasks if t.task == name)
        print(f"  {name:22} {task.outcome}")
        print(f"      {task.reason or 'its condition holds'}")
    print(
        "\n  `structure:declared` requires that name and `structure:inferred` forbids it.\n"
        "  Flatten the two lists into one expression and the second reads\n"
        "  `not (has_heading_styles = true)`, which is false — and every .docx comes out of\n"
        "  ingestion with no children."
    )

    _rule("A threshold the caller sets, on a file whose size was measured")
    data = _markdown_bytes()
    fmt, patterns = _probe(data, "retrieval-notes.md")
    print(f"  retrieval-notes.md probed as {fmt}: char_count={patterns.get(CHAR_COUNT)}")
    for facts in ({}, {"enabled": True}, {"enabled": True, "min_characters": 400}):
        _facts_line(fmt, patterns, facts)

    _rule("The same floor, on a format nothing measures the size of")
    pdf_patterns = {"has_text_layer": True, "has_outline": True}
    print(f"  patterns supplied   {_params(pdf_patterns)}")
    _facts_line("pdf", pdf_patterns, {"enabled": True, "min_characters": 400})
    print(
        "\n  probe emits `char_count` for `text` alone. Written as a `requires`, this floor\n"
        "  would have stopped fact extraction on every PDF, .docx and .pptx — not because\n"
        "  they are short, but because nobody measured. Written as a `forbids` it says what\n"
        "  a caller means by a minimum: do not spend the LLM call on a document measured as\n"
        "  too small, and do spend it where nobody measured a size."
    )

    _rule("Four terms the table refuses, at import time")
    _refuse("a pattern name the vocabulary does not have", term("has_hedings"))
    _refuse("a count compared with a flag", term(CHAR_COUNT, OP_GT, True))
    _refuse("an ordering operator on a flag", term(HAS_TEXT_LAYER, OP_GT, True))
    _refuse("an option nothing declares", term(CHAR_COUNT, OP_GT, option("facts", "min_chars")))
    print(
        "\n  Same shape as the four row mistakes above: the check is stated over the term the\n"
        "  appliance runs, not over a checked-in list of names to compare it against."
    )


def _guard_readers() -> dict[str, list[str]]:
    """Which rows read each guardable pattern, over every format a sentinel resolves for."""
    import jmfts_core.ingest_tasks as it

    readers: dict[str, list[str]] = {}
    for row in TASK_ROWS:
        for t in row.requires + row.forbids:
            if not isinstance(t.left, str):
                continue
            targets = it.SENTINEL_PATTERNS.get(t.left)
            names = sorted(set(targets.values())) if targets else [t.left]
            for name in names:
                readers.setdefault(name, [])
                if row.task not in readers[name]:
                    readers[name].append(row.task)
    return readers


def _facts_line(fmt: str, patterns: dict, facts: dict) -> None:
    """What `extract:facts` does under one `facts` option set, and the sentence if it does not."""
    plan = explain_plan(fmt, options={"facts": facts} if facts else None, patterns=patterns)
    task = next(t for t in plan.tasks if t.task == "extract:facts")
    print(f"  facts={_params(facts) if facts else '(the defaults)'}")
    print(f"      extract:facts   {task.outcome}")
    if task.reason:
        print(f"      {task.reason}")


def _refuse(label: str, bad) -> None:
    """Put one malformed term through the audit and print what it said."""
    import jmfts_core.ingest_tasks as it

    print(f"  {label}")
    try:
        it._check_term(TaskRow(task="a", write_mode="self", requires=(bad,)), bad)
    except Exception as exc:  # noqa: BLE001
        print(f"      {type(exc).__name__}: {exc}")
    else:
        print("      ACCEPTED — the audit did not catch this")


# ---------------------------------------------------------------------------
# 7. the stamp, live
# ---------------------------------------------------------------------------


def demo_stamp() -> None:
    """A node records which rule produced it, and a person's edit clears the record."""
    _require_database()

    from jmfts_client.contracts.upload import UploadedFile
    from jmfts_core.client import LocalJmftsClient
    from jmfts_core.database import get_session
    from jmfts_core.ingest_worker import IngestWorker
    from jmfts_core.models.document import Document
    from jmfts_core.repositories.document import DocumentRepository
    from jmfts_core.repositories.evidence import EvidenceRepository
    from jmfts_core.settling import NO_ROLLUP

    data = _markdown_bytes()
    client = LocalJmftsClient()
    uploaded = client.upload_file(UploadedFile(data=data, filename="retrieval-notes.md"))
    root_id = uploaded.document_id

    worker = IngestWorker(worker_id="demo-0-3-0", session_factory=get_session, planner=NO_ROLLUP)
    ran = 0
    while ran < 200 and worker.run_once():
        ran += 1

    _rule(f"retrieval-notes.md -> node {root_id}, {ran} tasks drained")
    with get_session() as session:
        nodes = (
            session.query(Document)
            .filter(Document.path.contains([root_id]) | (Document.id == root_id))
            .order_by(Document.id)
            .all()
        )
        print(f"  {'id':>5}  {'usetype':10} {'produced_by':22} title")
        for node in nodes:
            stamp = node.produced_by or "(asserted)"
            title = (node.title or "")[:34]
            print(f"  {node.id:>5}  {node.usetype or '':10} {stamp:22} {title}")

        print(
            "\n  `(asserted)` is a NULL column and it means a person, an importer or an\n"
            "  upload made the node. The file node is asserted because the upload wrote it."
        )

        _rule("Evidence is rows, not a column")
        evidence = EvidenceRepository(session)
        names = sorted(evidence.read_all(root_id))
        print(f"  node {root_id} carries {len(names)} evidence names:")
        print(f"      {', '.join(names)}")
        root = session.get(Document, root_id)
        print(f"  structured_content is the caller's: {root.structured_content!r}")

        _rule("A person edits a produced node")
        produced = next((n for n in nodes if n.produced_by), None)
        if produced is None:
            print("  nothing in this tree carries a stamp; skipping")
            return

        repo = DocumentRepository(session)
        print(f"  node {produced.id} before   produced_by={produced.produced_by!r}")
        repo.update(produced.id, title="A title the reader chose")
        session.flush()
        session.refresh(produced)
        print(f"  after a retitle           produced_by={produced.produced_by!r}")
        repo.update(produced.id, content="Text the reader wrote themselves.")
        session.flush()
        session.refresh(produced)
        print(f"  after a content edit      produced_by={produced.produced_by!r}")
        session.rollback()

    print(
        "\n  Only `content` clears it. A retitle does not make the node stop being what\n"
        "  the rule produced, and neither does a metadata patch."
    )


def _require_database() -> None:
    """Fail here, naming the way out, rather than deep inside the first query."""
    from sqlalchemy import text

    from jmfts_core.database import get_engine

    try:
        with get_engine().connect() as conn:
            conn.execute(text("SELECT 1 FROM documents LIMIT 1"))
    except Exception as exc:  # noqa: BLE001
        print(
            f"`stamp` needs a reachable JMFTS database with the schema in place.\n\n"
            f"  {type(exc).__name__}: {exc}\n\n"
            "Point JMFTS_DB_* at one and run `jmfts-init-db` against it. A throwaway:\n"
            "  docker run -d --name jmfts-demo-pg -e POSTGRES_USER=jmfts \\\n"
            "      -e POSTGRES_PASSWORD=jmfts -e POSTGRES_DB=jmfts \\\n"
            "      -p 127.0.0.1:5434:5432 pgvector/pgvector:pg16\n"
            "  JMFTS_DB_PORT=5434 jmfts-init-db\n"
            "  JMFTS_DB_PORT=5434 python -m scripts.demo_release_0_3_0 run stamp",
            file=sys.stderr,
        )
        raise SystemExit(2) from None


# ---------------------------------------------------------------------------
# the driver
# ---------------------------------------------------------------------------


class Section:
    def __init__(self, slug: str, title: str, phase: str, run: Callable[[], None]):
        self.slug = slug
        self.title = title
        self.phase = phase
        self.run = run

    @property
    def command(self) -> str:
        return f"python -m scripts.demo_release_0_3_0 run {self.slug}"


SECTIONS: tuple[Section, ...] = (
    Section("one-path", "One ingest path", "2a", demo_one_path),
    Section("scope", "A rule names its scope", "3", demo_scope),
    Section("explain", "EXPLAIN reaches the sheet tier", "3", demo_explain),
    Section("knobs", "The sheet knobs are settable", "3", demo_knobs),
    Section("audit", "The table audits itself", "3", demo_audit),
    Section("guards", "Guards take operators", "4", demo_guards),
    Section("stamp", "A node names its rule", "2b + 3", demo_stamp),
)

BY_SLUG = {section.slug: section for section in SECTIONS}


def cmd_list(_args: argparse.Namespace) -> int:
    print(f"  {'section':12} {'phase':8} what it shows")
    for section in SECTIONS:
        print(f"  {section.slug:12} {section.phase:8} {section.title}")
    print("\n  `stamp` needs a database. The rest need nothing.")
    return 0


def _selected(names: list[str]) -> list[Section]:
    if not names:
        return [section for section in SECTIONS if section.slug != "stamp"]
    unknown = [name for name in names if name not in BY_SLUG]
    if unknown:
        raise SystemExit(f"unknown section(s): {', '.join(unknown)}")
    return [BY_SLUG[name] for name in names]


def cmd_run(args: argparse.Namespace) -> int:
    for section in _selected(args.sections):
        print(f"\n{'=' * 78}")
        print(f"{section.title}  —  SPRINT_JOBS.md phase {section.phase}")
        print("=" * 78)
        section.run()
    return 0


def cmd_markdown(args: argparse.Namespace) -> int:
    """Emit each section as a fenced block, preceded by the command that produced it.

    The release page embeds this file. Regenerating it is how that page is updated, which
    is what keeps the page from describing a run the appliance no longer does.
    """
    parts: list[str] = []
    for section in _selected(args.sections):
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer):
            section.run()
        body = buffer.getvalue().strip("\n")
        parts.append(f"<!-- {section.slug} -->\n```console\n$ {section.command}\n{body}\n```")

    document = "\n\n".join(parts) + "\n"
    if args.out:
        Path(args.out).write_text(document, encoding="utf-8")
        print(f"wrote {args.out}", file=sys.stderr)
    else:
        sys.stdout.write(document)
    return 0


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("list", help="the sections and what each shows").set_defaults(func=cmd_list)

    p_run = sub.add_parser("run", help="run sections and print to the terminal")
    p_run.add_argument("sections", nargs="*")
    p_run.set_defaults(func=cmd_run)

    p_md = sub.add_parser("markdown", help="the same output, as fenced blocks")
    p_md.add_argument("sections", nargs="*")
    p_md.add_argument("--out")
    p_md.set_defaults(func=cmd_markdown)

    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
