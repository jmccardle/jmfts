"""What a worksheet measures to. ``docs/INGEST_SPEC.md`` 8.3.

One pass over one sheet's used range, producing 8.3's table and nothing else. **No shape
is chosen here and none is chosen by the task that calls this** — 8.8 leaves every
threshold in 8.4 unset pending calibration against real workbooks, so this module's whole
job is to produce the numbers that calibration will be run against.

TIER 2, and the reason this file is in :mod:`jmfts_core.office`: ``openpyxl`` is the
``office`` extra, the import happens inside :func:`~jmfts_core.office.require_openpyxl` at
the point of use, and ``tests/test_office_packaging.py`` fails if it climbs to module
scope.

**Three things openpyxl 3.1.5 does that 8.3 was written without knowing.** All three were
measured rather than assumed, and each one changed a definition here:

1. ``ReadOnlyWorksheet`` HAS NO ``merged_cells`` ATTRIBUTE AT ALL. ``OFFICE_SPEC.md`` Part
   4 note 2 requires ``read_only=True`` *because* ``profile:sheet`` measures merged cells,
   and in this version those two requirements contradict each other. The count is
   therefore read from the worksheet part's own ``<mergeCells>`` element with the standard
   library (:func:`count_merged_cells`), which costs a second scan of that part and is
   still cheaper than materialising the grid. Reporting zero would have been the
   indistinguishable-from-none failure the project's Fail Early rule exists to stop.

2. ``max_row``/``max_column`` COME FROM THE ``<dimension>`` ELEMENT, WHICH IS OPTIONAL.
   A workbook written without one reports ``None`` for both, ``reset_dimensions()`` does
   not help in read-only mode, and the rows come back ragged. So 8.3's "used range
   dimensions" is measured from the cells that actually hold a value, and what the file
   DECLARED is recorded separately — a writer whose dimension disagrees with its own cells
   is a fact about the file, not an error.

3. ``iter_rows`` IS ANCHORED AT A1 whatever the dimension says (``min_col = min_col or 1``
   in ``worksheet.py``), so an enumeration index is a real row number and a position in
   the yielded tuple is a real column number even for a sheet whose data starts at C5.
   That is relied on below and is the reason no cell object is constructed.

**Memory is bounded and the bound is disclosed.** An exact distinct-value set is the only
way to an exact ``distinct_count`` and an exact ``is_unique``, and a column of a million
rows cannot have one held in memory. :data:`DISTINCT_TRACKED_MAX` is where tracking stops;
past it the column reports what it knows — a floor, and a sketch — rather than a number it
did not measure. See the constant for why the value is a resource bound and not one of
8.8's thresholds.

**THE HEADER IS LOOKED FOR IN ROWS 1 TO 8, IN TWO PASSES, AND 8.3 SAID ROW 1.** That is an
amendment to 8.3 and it was chosen from a measurement rather than argued —
``scripts/header_rules.py`` applied six candidate rules to the first eight rows of 30,448
open-web (``datasets/corpus-fuse``) and 8,652 git-corpora sheets on 2026-09-03:

======================================  ================  =============
rule                                    FUSE (open web)   git corpora
======================================  ================  =============
8.3 as written — the rule at row 1      23.9%             35.5%
the same rule, rows 1–8                 45.5%             44.7%
plus the any-type second pass           54.4%             54.8%
======================================  ================  =============

``extract:sheet`` fires on nothing else, so those are the sheets that produce record nodes
at all. See :data:`HEADER_SCAN_ROWS` for why eight rows, :func:`_header_scan` for why two
passes rather than one relaxed rule, and :class:`HeaderRowScan` for what the rows above the
header are (``SPRINT_0_4_0.md`` open question 4.3, answered there).
"""

from __future__ import annotations

import datetime
import io
import re
import zipfile
from dataclasses import dataclass, field, replace
from typing import Any, Optional, Sequence
from xml.etree import ElementTree as ET

from jmfts_core.office import require_openpyxl
from jmfts_core.sketch import SketchBuilder

#: How many distinct values one tracked set holds before it stops growing. A RESOURCE
#: BOUND, not one of ``INGEST_SPEC.md`` 8.8's thresholds: nothing branches on it, no shape
#: is decided by it, and a threshold sweep may pick any closed-set threshold at all up to
#: this without re-reading a blob. What it costs to raise is memory — the tracked values
#: are the canonical strings themselves, so 50,000 of them is single-digit MB per column
#: and a sheet holds one set per column plus two for the interior.
#:
#: What it costs to have at all: a column with more distinct values than this reports
#: ``distinct_count: null`` with a floor beside it, and ``is_unique: null``. That is the
#: honest answer — "at least 50,000 distinct" is what was measured — and it is a real
#: limitation for the identifier column of a very large table, recorded rather than
#: papered over.
DISTINCT_TRACKED_MAX = 50_000

#: How many rows from the top :func:`_header_scan` looks at for a header row.
#:
#: MEASURED, not chosen. ``scripts/header_rules.py`` read the first eight rows of every
#: sheet in both corpora and recorded which row each candidate rule fired at. The rows the
#: shipped rule finds below row 1 are a LONG TAIL rather than a spike at row 2 — FUSE
#: r2:1,646 r3:1,312 r4:1,405 r5:731 r6:645 r7:574 — which is title rows and merged banners
#: of varying height, consistent with the 47.5% merged-cell rate on the same corpus. Two
#: rows would have taken less than half of what eight take.
#:
#: It bounds the memory this module holds for the scan (eight rows of canonical strings,
#: :attr:`_Scan.head_values`) and it bounds nothing else: the per-column accumulators start
#: below it and the rows between the header and row 8 are folded back in at assembly, so a
#: sheet whose header is at row 4 counts rows 5 onward exactly.
HEADER_SCAN_ROWS = 8

#: 8.3's rule as written: every cell of the used width is text, and all are distinct.
HEADER_RULE_ALL_TEXT = "all_text"

#: The second pass: every cell is non-empty and all are distinct, of ANY type. A year or
#: period header — ``2019, 2020, 2021`` — is a perfectly good field list that 8.3 rejects
#: for being numeric.
HEADER_RULE_ANY_TYPE = "any_type"

#: Where a column's ``name`` came from: the header row 8.3 names.
NAME_SOURCE_HEADER = "header_row"

#: Where a column's ``name`` came from: a merged banner above the data, for a sheet that has
#: no header row at all. See :func:`_banner_labels`.
NAME_SOURCE_BANNER = "banner"

#: How many of a column's distinct values are handed back for storage. 8.5 stores
#: ``values`` for a closed-set column and a count for a high-cardinality one, and WHICH IS
#: WHICH IS 8.8'S UNSET THRESHOLD — so this is not that decision. It is the ceiling on how
#: large a set can be and still be replayable from the stored profile, which is what makes
#: any later choice of that threshold a query rather than a re-ingest.
VALUES_RETAINED_MAX = 1_000

#: Cell value kinds, as ``dominant_type`` reports them. Open strings; named here so the
#: measurer and the prose renderer cannot come to spell them differently.
TYPE_TEXT = "text"
TYPE_NUMBER = "number"
TYPE_DATE = "date"
TYPE_BOOL = "bool"
#: What a column with no value in any row reports. Not "unknown" — it was measured, and
#: what was measured is that there is nothing there.
TYPE_EMPTY = "empty"

#: SpreadsheetML's main namespace, and the two relationships namespaces, for the tier-1
#: read of ``<mergeCells>``. Spelled here rather than imported from :mod:`jmfts_core.probe`:
#: probe must stay importable with nothing but the standard library and gains nothing from
#: exporting its constants to a tier-2 module.
_S_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PR_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

#: Escapes a cell value for a markdown table cell: a pipe would close the cell and a
#: newline would end the row.
_MARKDOWN_UNSAFE = re.compile(r"[|\r\n]")

#: One ``<mergeCell ref="A5:B6"/>``. The two-ended form only: a merge is a rectangle, and
#: ``$`` absolute markers do not appear in this attribute — Excel writes a plain rectangle.
_MERGE_REF = re.compile(r"^([A-Za-z]{1,3})([0-9]{1,7}):([A-Za-z]{1,3})([0-9]{1,7})$")

#: How fast :func:`_scan` runs over the empty rows a wide ``<dimension>`` pads with.
#: MEASURED, 2026-09-05, openpyxl 3.1.5 on CPython 3.11: three synthetic sheets of five
#: cells declaring 8.2e6, 3.3e7 and 8.2e7 cells scanned in 0.66s, 2.87s and 6.87s, which is
#: 1.25e7, 1.14e7 and 1.19e7 declared cells per second. It is here so the refusal below can
#: say what the scan WOULD have cost — "2.75e11 cells" means nothing to an operator and
#: "6 hours" means something — and so the derivation of :data:`DECLARED_CELLS_MAX` is
#: arithmetic somebody can redo rather than a number to trust. A cell holding a value costs
#: more than this; the rate is the padding rate, which is what the pathological case is
#: made of.
SCAN_CELLS_PER_SECOND = 1.2e7

#: The largest ``declared_rows * declared_cols`` :func:`measure_sheet` will scan.
#: ``docs/SPRINT_0_4_0.md`` Block B step 4, and a RESOURCE BOUND like
#: :data:`DISTINCT_TRACKED_MAX` rather than one of 8.8's thresholds — nothing branches on
#: it and no shape is decided by it.
#:
#: **Why the declaration and not a clock.** ``ReadOnlyWorksheet._cells_by_row`` fills the
#: gap between two ``<row r="...">`` indices by yielding an empty row per missing index,
#: each one ``max_column`` wide, and ``max_column``/``max_row`` come from ``<dimension>``.
#: So a five-cell sheet whose rows are numbered 1, 16777214 and 16777217 under
#: ``<dimension ref="A1:XFE16777217"/>`` makes :func:`_scan`'s inner loop run 2.75e11
#: times. Nothing downstream ends it: ``repositories/task_queue.py:458`` says
#: ``lease_seconds`` bounds how long a live worker may go without beating and not how long
#: a task may run, the heartbeat thread keeps beating, and ``ingest_worker.py:288`` says
#: the thread cannot interrupt the handler — so that one file removes a worker from the
#: fleet permanently. A wall clock would stop it and would also be untestable and
#: unreproducible: two workers on different hardware would disagree about the same file,
#: and the same file would pass and fail on the same worker under different load. The
#: declared extent is a property of the bytes, so the answer is the same everywhere.
#:
#: **Where the number comes from.** Two measurements, both taken 2026-09-05:
#:
#: * ``<dimension>`` read out of 35,547 worksheet parts in 11,447 real workbooks
#:   (``datasets/corpus-fuse`` open-web, ``datasets/lo-qa``, ``datasets/poi-testdata``).
#:   Sheets declaring more than 1e6 cells: 249. More than 1e7: 17. More than 5e7: 5. More
#:   than 1e8: 4. More than 1e9: 3 — which is exactly the three failures Block B measured.
#:   The four above this bound are 2.7e11, 1.7e10, 1.1e9 and 3.3e8 cells and every one of
#:   them declares the whole grid or more; the largest accepted is 6.3e7. The cut lands in
#:   a 5x gap in the distribution rather than through a cluster.
#: * :data:`SCAN_CELLS_PER_SECOND`, the rate the padded rows are scanned at.
#:
#: 1e8 cells is therefore about 8 seconds of CPU for one sheet — under a tenth of
#: ``Settings.worker_lease_seconds`` (90) — while the four refused sheets are 27 seconds,
#: 90 seconds, 24 minutes and 6.4 hours. Raising this costs exactly that time, per sheet.
DECLARED_CELLS_MAX = 100_000_000


class WorksheetPartMissing(ValueError):
    """The workbook does not resolve this sheet's name to a part in the package.

    Raised rather than returning a merged-cell count of zero. Zero merges and "the sheet's
    XML could not be found" are different facts, and a profile that recorded the second as
    the first would put a measurement into the record that nothing measured.
    """


class DeclaredExtentTooLarge(ValueError):
    """The sheet declares more cells than :data:`DECLARED_CELLS_MAX`. Step 4.

    A ``ValueError`` subclass so ``task_errors.classify_exception`` grades it PERMANENT,
    which is the correct grade and not a convenient one: a declared extent is a string in
    the file's own ``<dimension>`` element and does not shrink between attempts. Retrying
    it would spend the budget three times over on the one input where a single attempt
    already costs more than the worker.

    It is a refusal and not a truncation. Scanning the first ``DECLARED_CELLS_MAX`` cells
    and reporting the result would put a measurement of part of a sheet into a record whose
    every other field means the whole sheet — ``fill_ratio``, ``distinct_count`` and
    ``is_unique`` would each be a number nobody could tell from the real one.
    """


class SheetIsNotAWorksheet(ValueError):
    """The workbook resolves this name to something with no cell grid. Step 5.

    A chartsheet — 135 of 10,702 open-web workbooks carry one (``docs/SPRINT_0_4_0.md``
    Block B) — is a ``sheet`` in ``xl/workbook.xml`` and is named in
    ``workbook.sheetnames``, so it reaches :func:`measure_sheet` exactly like a worksheet
    and then has no ``max_row`` at all. Left alone that is an ``AttributeError`` reading
    ``'Chartsheet' object has no attribute 'max_row'``, which is already PERMANENT and
    tells the operator nothing about the file. This says what the sheet is.
    """


# ---------------------------------------------------------------------------
# Canonical values
# ---------------------------------------------------------------------------


def canonical(value: Any) -> Optional[str]:
    """One cell as the string every count, comparison and sketch sees, or ``None``.

    Canonicalisation is the whole correctness of a cross-sheet comparison (8.6): two
    columns holding the same identifiers agree only if ``128000`` and ``128000.0`` reached
    the counter as one string. So it happens once, here, and every consumer downstream
    reads the result rather than the cell.

    ``None`` means the cell is empty, and a string of nothing but whitespace is empty too.
    That is a choice and it is recorded: a cell holding three spaces is not a value a
    person put there to mean something, and counting it as one would make ``fill_ratio``
    depend on a workbook's whitespace.

    A ``datetime`` at exactly midnight renders as a plain date. openpyxl returns
    ``datetime.datetime`` for a cell formatted as a date, with no way to tell a date from a
    date-time whose time is zero, and ``2026-09-30`` is what such a column holds.
    """
    if value is None:
        return None
    # Before `int`: `bool` is a subclass of it, and `str(True)` is "True" while a
    # spreadsheet's own rendering is TRUE.
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        # `1.0` and `1` are the same value to a spreadsheet and must sketch to one string.
        # `is_integer()` is exact and `repr` round-trips, so neither branch loses a digit.
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, datetime.timedelta):
        return str(value)
    text = str(value).strip()
    return text or None


def value_type(value: Any) -> str:
    """Which of :data:`TYPE_TEXT` and friends a non-empty cell is."""
    if isinstance(value, bool):
        return TYPE_BOOL
    if isinstance(value, (int, float)):
        return TYPE_NUMBER
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        return TYPE_DATE
    return TYPE_TEXT


def column_letter(index: int) -> str:
    """1 -> ``A``, 27 -> ``AA``. The address a person reads off the top of the sheet."""
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def column_index(reference: str) -> int:
    """``"A2"`` -> 1, ``"AB17"`` -> 28. :func:`column_letter`'s inverse.

    IT LIVES HERE AND NOT IN :mod:`jmfts_core.office.cells`, which is where it was written
    and which still exports it. This module is the lower of the two — ``cells`` imports
    ``canonical``, ``column_letter`` and ``_worksheet_part`` from here and says why: one
    escaping rule and one part-resolution walk, in one place, because two copies are free to
    disagree. :func:`_merged_ranges` needs to turn ``A5:B6`` into column numbers, and a
    second letters-to-index loop here would have been exactly that second copy.
    """
    total = 0
    for character in reference:
        if not character.isalpha():
            break
        total = total * 26 + (ord(character.upper()) - ord("A") + 1)
    return total


# ---------------------------------------------------------------------------
# The shapes the measurement comes back in
# ---------------------------------------------------------------------------


@dataclass
class _DistinctSet:
    """An exact distinct set that stops growing, and says that it stopped.

    The two fields a caller reads are :attr:`count` and :attr:`exact`. ``exact`` false
    means ``count`` is a floor and the true cardinality is unknown to this object — which
    is a measurement, not a failure, and is reported as one.
    """

    limit: int = DISTINCT_TRACKED_MAX
    values: set = field(default_factory=set)
    exact: bool = True

    def add(self, value: str) -> None:
        if not self.exact:
            return
        self.values.add(value)
        if len(self.values) >= self.limit:
            self.exact = False

    def plus(self, values: Sequence[Optional[str]]) -> "_DistinctSet":
        """A copy holding the head rows' values too, for the fold-in in :func:`_column`.

        A set cannot be un-added from, so the cells of rows 1 to :data:`HEADER_SCAN_ROWS`
        are held out of the scan and folded in here instead: a header label may also appear
        in its own column's body, and "subtract one" would then be wrong by exactly the case
        that matters.

        IT TAKES A SEQUENCE SINCE THE HEADER SCAN, where it took one value. The row held out
        used to be row 1 alone, because 8.3 defined the header as a property of row 1; the
        scan now holds out eight and folds back the ones that turned out to be below the
        header, which is between zero and eight values rather than one.
        """
        present = [value for value in values if value is not None]
        if not present:
            return self
        copied = _DistinctSet(limit=self.limit, values=set(self.values), exact=self.exact)
        for value in present:
            copied.add(value)
        return copied

    @property
    def count(self) -> int:
        return len(self.values)


@dataclass(frozen=True)
class ColumnMeasurement:
    """8.3's per-column row, plus what a threshold sweep needs to replay it.

    ``distinct_count`` is ``None`` exactly when :data:`DISTINCT_TRACKED_MAX` was reached,
    and ``distinct_at_least`` is then the floor that was measured. ``is_unique`` is
    ``None`` in the same case: it is ``distinct_count == non_empty == body_rows``, and only
    two of those three are known.

    ``body_rows`` is the denominator ``fill_ratio`` was taken over — the rows below the
    header row, or below the banner where there is one and no header, or every row where
    there is neither. It is carried rather than left to be re-derived because a sweep that
    redefines ``header_row`` (see :class:`HeaderEvidence`) changes the denominator, and it
    must be able to see which one was used. The header scan MOVED this number for two
    populations at once — a header at row 4 takes three rows out of it, and a headerless
    sheet under a merged banner takes the banner's rows out — so it is now carried for a
    reason it was only anticipating before.

    ``name_source`` says which of the two rules named this column: :data:`NAME_SOURCE_HEADER`
    for the header row 8.3 defines, :data:`NAME_SOURCE_BANNER` for the merged banner a
    header-less sheet is labelled from. ``None`` exactly when ``name`` is. It is not
    cosmetic: ``sheet_records.header_labels`` accepts only the first, because a banner label
    names a GROUP of columns and a record key has to name one.
    """

    index: int
    letter: str
    name: Optional[str]
    name_source: Optional[str]
    dominant_type: str
    type_counts: dict
    non_empty: int
    body_rows: int
    fill_ratio: float
    distinct_count: Optional[int]
    distinct_at_least: int
    distinct_exact: bool
    values: Optional[list]
    is_unique: Optional[bool]
    sketch: Optional[dict]


@dataclass(frozen=True)
class HeaderEvidence:
    """Why :attr:`SheetMeasurement.header_row` (or ``header_col``) came out as it did.

    8.3 gives one boolean per header and 8.4 branches on it. The components are stored
    beside the verdict because **the verdict as 8.3 defines it cannot see the sheet 8.4's
    own example is drawn from**: a crossing table has an empty top-left corner, so its
    first row is not "all text" and ``header_row`` is false — and the ``matrix`` shape,
    which requires both headers, could never fire for the shape it was written for.

    That is a gap in the spec, not a licence to invent a looser rule here, so the rule is
    implemented exactly as written and ``leading_empty`` records the corner case as a
    measurement. Step 7 can redefine the verdict from these numbers with no blob read.
    """

    verdict: bool
    cells: int
    non_empty: int
    text_cells: int
    numeric_cells: int
    other_cells: int
    distinct: int
    all_non_empty: bool
    all_text: bool
    all_distinct: bool
    #: The first cell is empty and every other cell satisfies the rule. The crossing
    #: table's corner, measured.
    leading_empty: bool

    @property
    def any_type_verdict(self) -> bool:
        """The second pass's rule: non-empty and distinct, of any type.

        DERIVED AND NOT STORED, because it is exactly ``all_non_empty and all_distinct`` and
        both were already measured for their own sake. Registering a fourth component would
        have been a fourth number that has to stay equal to two others.
        """
        return self.all_non_empty and self.all_distinct


@dataclass(frozen=True)
class BannerRow:
    """One row above the data that is not a header. ``SPRINT_0_4_0.md`` open question 4.3.

    **That question asked what rows 1 to N−1 are when the header is at row N, and offered
    three readings: discard them, keep them as the sheet node's own content, or read a label
    out of them. This is the answer, and it takes the second and third and refuses the
    first.** They are a MEASUREMENT — this class, carried on
    :attr:`HeaderRowScan.banner` — and they are spent twice: ``sheet_profile`` states them
    in the profile node's prose, so the words a person put at the top of the sheet are
    embedded and retrievable, and :func:`_banner_labels` reads a column label out of the
    merged ones for a sheet that has no header row at all.

    Discarding was never available once the scan existed. A header at row 4 means rows 1–3
    were READ, and a measurement that is read and dropped is a fact the appliance had and
    threw away; the rendered table leaves them out of its body because a markdown table has
    one header row, which is a rendering decision rather than a decision about the rows.
    """

    index: int
    #: Canonical strings over ``1..cols``, ``None`` where the cell is empty.
    values: tuple


@dataclass(frozen=True)
class MergedRange:
    """One ``<mergeCell ref="A5:B6"/>``, as row and column numbers. 1-based, inclusive."""

    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def cols(self) -> int:
        return self.max_col - self.min_col + 1


@dataclass(frozen=True)
class MergedCells:
    """What ``<mergeCells>`` declares: the count 8.3 asks for, and the ranges near the top.

    ``ranges`` is NOT every merge in the sheet. Only the ones inside the scanned prefix can
    be a banner — a merge at row 900 is a formatting choice in the middle of the data — so
    only those are kept, and the whole element is already materialised by the parse that
    counts them, so keeping a filtered tuple of four-int records costs nothing the count did
    not already cost.
    """

    count: int
    ranges: tuple


@dataclass(frozen=True)
class HeaderRowScan:
    """Which row of the top :data:`HEADER_SCAN_ROWS` is the header, and under which rule.

    8.3 defines ``header_row`` as a property of row 1 and :attr:`verdict` keeps that name and
    that meaning — is there a header row at all — while :attr:`row` says WHERE it is and
    :attr:`rule` says which of the two passes found it. A sweep that wants to re-decide reads
    :attr:`candidates`, which holds one :class:`HeaderEvidence` per scanned row and is the
    whole input to both passes.

    :attr:`evidence` is the components at :attr:`row`, or row 1's when nothing fired — row 1's
    because that is the row 8.3 asks about, so a reader who does not know about the scan sees
    the number 8.3 promised them.
    """

    verdict: bool
    row: Optional[int]
    rule: Optional[str]
    evidence: HeaderEvidence
    candidates: tuple
    scanned_rows: int
    banner: tuple

    @property
    def first_data_row(self) -> int:
        """The first row that is an instance rather than a label. 1 when there is neither."""
        if self.row is not None:
            return self.row + 1
        return (max((banner.index for banner in self.banner), default=0)) + 1


@dataclass(frozen=True)
class SheetMeasurement:
    """Everything ``INGEST_SPEC.md`` 8.3 asks for about one sheet, and no decision.

    ``rendered_markdown`` is ``None`` when the sheet is provably too large to fit any
    embedding window; see :func:`measure_sheet` for the derivation, which is an inequality
    rather than a chosen size.

    ``header_row`` is a :class:`HeaderRowScan` and ``header_col`` is a bare
    :class:`HeaderEvidence`, and the asymmetry is deliberate: the row scan is what was
    amended, and 8.3's ``header_col`` — "column A is all text and all distinct below row 1" —
    is still implemented exactly as written. Nothing measured says it should move, and moving
    it on the symmetry argument alone would be a rule chosen for tidiness.
    """

    name: str
    rows: int
    cols: int
    non_empty_cells: int
    fill_ratio: float
    declared_rows: Optional[int]
    declared_cols: Optional[int]
    header_row: HeaderRowScan
    header_col: HeaderEvidence
    interior_rows: int
    interior_cols: int
    interior_non_empty: int
    interior_cardinality: Optional[int]
    interior_cardinality_at_least: int
    interior_cardinality_exact: bool
    merged_cells: int
    merged_ranges: tuple
    columns: tuple
    rendered_markdown: Optional[str]
    rendered_unbounded_reason: Optional[str]
    distinct_tracked_max: int
    values_retained_max: int
    header_scan_rows: int


# ---------------------------------------------------------------------------
# The header rule
# ---------------------------------------------------------------------------


def _header_evidence(values: list, kinds: list) -> HeaderEvidence:
    """8.3's header rule over one row or column, with its components kept.

    ``values`` are canonical strings, ``None`` for empty, in address order and covering the
    whole used extent — so a trailing empty cell counts against the rule exactly as an
    interior one does, which is what "all text" means. ``kinds`` is the matching list of
    :func:`value_type` results, ``None`` where the cell is empty; the rule needs it because
    a canonical string cannot say what it came from and "no numerics" is half the rule.
    """
    present = [value for value in values if value is not None]
    text_cells = sum(1 for kind in kinds if kind == TYPE_TEXT)
    numeric_cells = sum(1 for kind in kinds if kind == TYPE_NUMBER)
    other_cells = sum(1 for kind in kinds if kind not in (None, TYPE_TEXT, TYPE_NUMBER))
    distinct = len(set(present))
    all_non_empty = bool(values) and len(present) == len(values)
    all_text = bool(values) and text_cells == len(values)
    all_distinct = bool(present) and distinct == len(present)
    tail_values = values[1:]
    tail_kinds = kinds[1:]
    tail_present = [value for value in tail_values if value is not None]
    leading_empty = (
        bool(tail_values)
        and values[0] is None
        and sum(1 for kind in tail_kinds if kind == TYPE_TEXT) == len(tail_values)
        and len(set(tail_present)) == len(tail_present)
    )
    return HeaderEvidence(
        verdict=all_non_empty and all_text and all_distinct,
        cells=len(values),
        non_empty=len(present),
        text_cells=text_cells,
        numeric_cells=numeric_cells,
        other_cells=other_cells,
        distinct=distinct,
        all_non_empty=all_non_empty,
        all_text=all_text,
        all_distinct=all_distinct,
        leading_empty=leading_empty,
    )


def _header_scan(candidates: Sequence[HeaderEvidence]) -> HeaderRowScan:
    """Two passes over the top rows. The first pass is 8.3's rule; the second allows numerics.

    **TWO PASSES AND NOT ONE RELAXED RULE, and the reason is ordering rather than rate.**
    The any-type rule accepts every row the all-text rule accepts, so as a SET the two
    orderings admit exactly the same sheets — 54.4% / 54.8% either way. What differs is
    which ROW each picks, and a single interleaved pass would let a numeric row above a real
    header win. ``scripts/header_rules.py`` measured what that costs on a neighbouring
    candidate: naming row 1's contiguous prefix fires EARLIER than the shipped rule on 38.9%
    of the sheets both accept, and on those it names 1.1 columns where the shipped rule
    names 7.6, because it finds a one-cell title row and stops. A relaxed rule does not only
    admit more sheets, it admits them sooner, and that moves the outcome from "no records"
    to "wrong records". Running the strict rule over all eight rows first is what buys the
    admissions without the reordering.

    Two other relaxations were measured and are NOT here: naming every present cell fires on
    95.2% of FUSE sheets and admits a merged banner with holes in it as a field list, and a
    majority-text rule gained 235 FUSE sheets and 42 git sheets, which is not worth a code
    change.

    The returned scan carries NO banner. Which rows are the banner is decided here — they
    are the rows above :attr:`HeaderRowScan.row`, or the rows a merged banner covers where
    no pass fired — but their VALUES live in the scan accumulator, so :func:`_assemble`
    attaches them.
    """
    for rule, matches in (
        (HEADER_RULE_ALL_TEXT, lambda evidence: evidence.verdict),
        (HEADER_RULE_ANY_TYPE, lambda evidence: evidence.any_type_verdict),
    ):
        for offset, evidence in enumerate(candidates):
            if matches(evidence):
                row = offset + 1
                return HeaderRowScan(
                    verdict=True,
                    row=row,
                    rule=rule,
                    evidence=evidence,
                    candidates=tuple(candidates),
                    scanned_rows=len(candidates),
                    banner=(),
                )
    return HeaderRowScan(
        verdict=False,
        row=None,
        rule=None,
        # Row 1's, which is the row 8.3 asks about. An empty sheet has no candidate at all
        # and gets the rule's answer over no cells, which is what `_header_evidence` returns
        # for an empty list rather than a shape invented here.
        evidence=candidates[0] if candidates else _header_evidence([], []),
        candidates=tuple(candidates),
        scanned_rows=len(candidates),
        banner=(),
    )


# ---------------------------------------------------------------------------
# The scan
# ---------------------------------------------------------------------------


@dataclass
class _Scan:
    """The mutable accumulator :func:`_scan` fills and :func:`_assemble` freezes.

    **Two interior sets, not four, and the header scan is what took the other two away.**
    The interior begins right of column A only if ``header_col`` holds — a property of every
    used ROW, unknown until the last one has been read — so both column candidates are
    kept, keyed by ``col_offset`` of 0 or 1. The ROW side used to be the same two-way
    choice, row 1 or not; it is now a number between 0 and 8 and keeping nine candidate sets
    for it would be nine bounded sets per sheet. Instead the head rows are held out of the
    interior entirely and folded back in at assembly, which is the treatment the columns
    already had and is exact for the same reason.

    ``head_values`` and ``head_kinds`` are keyed ``(row, col)`` over rows 1 to
    :data:`HEADER_SCAN_ROWS`. They are what both passes of the header scan read, what the
    fold-in below the header replays, and what a merged banner's label is read from.
    """

    rows: int = 0
    cols: int = 0
    non_empty_cells: int = 0
    head_values: dict = field(default_factory=dict)
    head_kinds: dict = field(default_factory=dict)
    first_col: dict = field(default_factory=dict)
    first_col_kind: dict = field(default_factory=dict)
    per_column_non_empty: dict = field(default_factory=dict)
    per_column_types: dict = field(default_factory=dict)
    per_column_distinct: dict = field(default_factory=dict)
    interior: dict = field(default_factory=dict)
    interior_non_empty: dict = field(default_factory=dict)
    rendered_rows: list = field(default_factory=list)
    render_bounded: bool = False


_INTERIOR_COL_OFFSETS = (0, 1)


def _scan(
    worksheet,
    *,
    render_cell_budget: int,
    distinct_tracked_max: int,
    header_scan_rows: int,
) -> _Scan:
    """The single pass. Everything 8.3 measures is accumulated here or not at all."""
    scan = _Scan()
    for col_offset in _INTERIOR_COL_OFFSETS:
        scan.interior[col_offset] = _DistinctSet(limit=distinct_tracked_max)
        scan.interior_non_empty[col_offset] = 0

    for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        rendered: list = []
        row_has_value = False
        in_head = row_index <= header_scan_rows
        for col_index, raw in enumerate(row, start=1):
            text = canonical(raw)
            rendered.append("" if text is None else text)
            if text is None:
                continue
            kind = value_type(raw)
            row_has_value = True
            scan.non_empty_cells += 1
            if col_index > scan.cols:
                scan.cols = col_index

            if in_head:
                scan.head_values[(row_index, col_index)] = text
                scan.head_kinds[(row_index, col_index)] = kind
            if col_index == 1:
                scan.first_col[row_index] = text
                scan.first_col_kind[row_index] = kind

            # THE HEAD ROWS ARE NOT PART OF ANY COLUMN'S STATISTICS while one of them might
            # be the header, and which one is is not known until every used column has been
            # read. So the accumulators below cover rows HEADER_SCAN_ROWS+1..N and
            # `head_values` above keeps the cells; the rows that turn out to be BELOW the
            # header are folded back in at assembly (:func:`_column`, :func:`_assemble`).
            # Counting them here would give a header cell a vote in its own column's
            # `dominant_type`, put the label in the column's distinct set, and make
            # `fill_ratio` exceed 1.
            #
            # It used to be row 1 alone that was held out, because 8.3 defined the header as
            # a property of row 1. Eight rows is the same trade at eight times the size, and
            # the size is what `HEADER_SCAN_ROWS` bounds.
            if in_head:
                continue

            for col_offset in _INTERIOR_COL_OFFSETS:
                if col_index <= col_offset:
                    continue
                scan.interior[col_offset].add(text)
                scan.interior_non_empty[col_offset] += 1

            scan.per_column_non_empty[col_index] = scan.per_column_non_empty.get(col_index, 0) + 1
            types = scan.per_column_types.setdefault(col_index, {})
            types[kind] = types.get(kind, 0) + 1
            distinct = scan.per_column_distinct.get(col_index)
            if distinct is None:
                distinct = _DistinctSet(limit=distinct_tracked_max)
                scan.per_column_distinct[col_index] = distinct
            distinct.add(text)

        if row_has_value:
            scan.rows = row_index
        _retain(scan, rendered, render_cell_budget)

    return scan


def _retain(scan: _Scan, rendered: list, budget: int) -> None:
    """Keep one rendered row while a table of them could still fit an embedding window.

    See :func:`measure_sheet` for the inequality. Once it fails, the rows already kept are
    dropped: they are a prefix of a table nobody can use, and holding them would trade the
    memory this bound exists to save for nothing.
    """
    if scan.render_bounded:
        return
    floor = max(len(scan.rendered_rows) + 1 + 2, scan.non_empty_cells)
    if floor > budget:
        scan.render_bounded = True
        scan.rendered_rows.clear()
        return
    scan.rendered_rows.append(tuple(rendered))


# ---------------------------------------------------------------------------
# The entry point
# ---------------------------------------------------------------------------


def measure_sheet(
    data: bytes,
    sheet_name: str,
    *,
    render_cell_budget: int,
    with_sketches: bool = True,
    distinct_tracked_max: int = DISTINCT_TRACKED_MAX,
    values_retained_max: int = VALUES_RETAINED_MAX,
    header_scan_rows: int = HEADER_SCAN_ROWS,
) -> SheetMeasurement:
    """One pass over one sheet's used range. ``INGEST_SPEC.md`` 8.3.

    **The render budget is an inequality, not a size somebody liked.** A markdown table of
    R data rows spends at least one token on each row's newline and at least one on each
    non-empty cell, so its token count is at least ``max(R + 2, non_empty)``. Once that
    floor passes ``render_cell_budget`` — the embedding document window, which is the
    model's own limit and therefore the largest value 8.4's ``small_table`` test could ever
    be calibrated to — rendering the rest cannot change the answer to any question anybody
    can ask of it, so the rows stop being kept and the reason is recorded. A sheet under
    the floor is rendered exactly and counted exactly.

    ``with_sketches`` false is for a caller that wants the measurements on an install with
    no ``datasketch``. It is not the default: a column with no sketch is invisible to 8.6's
    containment search whatever its cardinality, and going without has to be asked for.

    **Two things it refuses before it scans anything** (``docs/SPRINT_0_4_0.md`` Block B
    steps 4 and 5, both reproduced against real workbooks in 0.3.0). A name that resolves
    to a chartsheet has no grid to measure — :class:`SheetIsNotAWorksheet`. A sheet whose
    ``<dimension>`` declares more than :data:`DECLARED_CELLS_MAX` cells costs more to scan
    than the worker is worth — :class:`DeclaredExtentTooLarge`. Both checks read what is
    already read here, so neither costs a pass over the part, and both are named
    ``ValueError`` subclasses so a caller can tell them from a bug.
    """
    openpyxl = require_openpyxl()
    # Tier 2, at the point of use like `require_openpyxl` itself: `tests/test_office_
    # packaging.py` fails if either import reaches the application's import path.
    from openpyxl.chartsheet import Chartsheet

    # `data_only=True`: a profile measures VALUES. `=SUM(B2:B9)` is not a value, and a
    # column of them would count as text with 1,284 distinct strings. The cost is that a
    # workbook whose formulas were never evaluated by Excel — one openpyxl itself wrote —
    # reports those cells empty, which is what it can honestly say about a file that
    # carries no cached result.
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise WorksheetPartMissing(
                f"the workbook names no sheet {sheet_name!r}; it names "
                f"{workbook.sheetnames!r}, so the node this profile is scoped to and the "
                "stored bytes do not describe the same workbook"
            )
        worksheet = workbook[sheet_name]
        # Step 5. `Chartsheet` is what `read_only=True` hands back for a `<sheet>` whose
        # relationship targets `xl/chartsheets/`; the `iter_rows` arm is the same question
        # asked of whatever openpyxl adds next, and both say which class arrived rather
        # than which attribute was missing.
        if isinstance(worksheet, Chartsheet) or not hasattr(worksheet, "iter_rows"):
            raise SheetIsNotAWorksheet(
                f"the workbook resolves {sheet_name!r} to a "
                f"{type(worksheet).__name__}, which has no cell grid: it is a sheet tab "
                "holding a chart drawn from cells that live on some OTHER sheet. There is "
                "nothing here to measure, and a profile reporting zero rows and zero "
                "columns would be indistinguishable from an empty worksheet"
            )
        declared_rows = worksheet.max_row
        declared_cols = worksheet.max_column
        # Step 4. Both are None for a sheet that declares no `<dimension>`, and that case
        # is deliberately NOT guarded: with no declaration `_cells_by_row` pads with an
        # EMPTY tuple rather than a `max_column`-wide one, so the gap-filling loop costs
        # one iteration per missing row index instead of `max_column` of them, and the
        # scan finishes. The declaration is what makes the pad wide, so the declaration is
        # what is bounded.
        if declared_rows is not None and declared_cols is not None:
            declared_cells = declared_rows * declared_cols
            if declared_cells > DECLARED_CELLS_MAX:
                raise DeclaredExtentTooLarge(
                    f"sheet {sheet_name!r} declares {declared_rows:,} rows by "
                    f"{declared_cols:,} columns = {declared_cells:,} cells, over the "
                    f"{DECLARED_CELLS_MAX:,} this appliance will scan for one sheet "
                    f"(~{declared_cells / SCAN_CELLS_PER_SECOND:,.0f}s against a "
                    f"~{DECLARED_CELLS_MAX / SCAN_CELLS_PER_SECOND:,.0f}s budget). The "
                    "cells that hold a value may be few — a declared extent is what the "
                    "writer put in <dimension>, not what it filled — but the scan is over "
                    "the declared extent, so this sheet cannot be profiled without taking "
                    "the worker out of the fleet for the duration"
                )
        scanned = _scan(
            worksheet,
            render_cell_budget=render_cell_budget,
            distinct_tracked_max=distinct_tracked_max,
            header_scan_rows=header_scan_rows,
        )
    finally:
        # A read-only workbook holds the ZIP open and the caller is inside a database
        # transaction.
        workbook.close()

    return _assemble(
        scanned,
        name=sheet_name,
        declared_rows=declared_rows,
        declared_cols=declared_cols,
        merged=read_merged_cells(data, sheet_name, within_rows=header_scan_rows),
        with_sketches=with_sketches,
        distinct_tracked_max=distinct_tracked_max,
        values_retained_max=values_retained_max,
        header_scan_rows=header_scan_rows,
    )


def _assemble(
    scan: _Scan,
    *,
    name: str,
    declared_rows: Optional[int],
    declared_cols: Optional[int],
    merged: MergedCells,
    with_sketches: bool,
    distinct_tracked_max: int,
    values_retained_max: int,
    header_scan_rows: int,
) -> SheetMeasurement:
    rows, cols = scan.rows, scan.cols
    cells = rows * cols
    fill_ratio = (scan.non_empty_cells / cells) if cells else 0.0

    scanned = min(rows, header_scan_rows)
    header_row = _header_scan(
        [
            _header_evidence(
                [scan.head_values.get((row, index)) for index in range(1, cols + 1)],
                [scan.head_kinds.get((row, index)) for index in range(1, cols + 1)],
            )
            for row in range(1, scanned + 1)
        ]
    )
    # 8.3: "column A is all text and all distinct BELOW ROW 1" — below row 1 outright, not
    # below whatever turned out to be the header row. Implemented as written; see
    # `SheetMeasurement` for why the header scan did not move it.
    header_col = _header_evidence(
        [scan.first_col.get(index) for index in range(2, rows + 1)],
        [scan.first_col_kind.get(index) for index in range(2, rows + 1)],
    )

    banner_ranges = _banner_ranges(merged.ranges, scan=scan, header_row=header_row, cols=cols)
    banner_bottom = max((r.max_row for r in banner_ranges), default=0)
    header_row = replace(
        header_row,
        banner=tuple(
            BannerRow(
                index=row,
                values=tuple(scan.head_values.get((row, index)) for index in range(1, cols + 1)),
            )
            for row in range(1, (header_row.row or banner_bottom + 1))
        ),
    )
    labels = _banner_labels(banner_ranges, scan=scan)

    first_data_row = header_row.first_data_row
    col_offset = 1 if header_col.verdict else 0
    body_rows = max(rows - (first_data_row - 1), 0)
    # The head rows at or below the first data row were held out of the scan and are data,
    # so they are replayed into the interior here — the same fold-in `_column` does, over
    # the same rows, for the same reason the accumulators could not do it in one pass.
    folded = range(first_data_row, scanned + 1)
    interior = scan.interior[col_offset].plus(
        [
            scan.head_values.get((row, index))
            for row in folded
            for index in range(col_offset + 1, cols + 1)
        ]
    )
    interior_non_empty = scan.interior_non_empty[col_offset] + sum(
        1
        for row in folded
        for index in range(col_offset + 1, cols + 1)
        if scan.head_values.get((row, index)) is not None
    )
    named = {
        index: _column_name(index, scan=scan, header_row=header_row, labels=labels)
        for index in range(1, cols + 1)
    }
    columns = tuple(
        _column(
            index=index,
            name=named[index][0],
            name_source=named[index][1],
            scan=scan,
            body_rows=body_rows,
            # A head row at or below the first data row is a value in this column, not a
            # label for it, so it is folded back in.
            folded_rows=folded,
            with_sketches=with_sketches,
            distinct_tracked_max=distinct_tracked_max,
            values_retained_max=values_retained_max,
        )
        for index in range(1, cols + 1)
    )
    rendered, reason = _render(scan, header_row, rows, cols)
    return SheetMeasurement(
        name=name,
        rows=rows,
        cols=cols,
        non_empty_cells=scan.non_empty_cells,
        fill_ratio=fill_ratio,
        declared_rows=declared_rows,
        declared_cols=declared_cols,
        header_row=header_row,
        header_col=header_col,
        interior_rows=body_rows,
        interior_cols=max(cols - col_offset, 0),
        interior_non_empty=interior_non_empty,
        interior_cardinality=interior.count if interior.exact else None,
        interior_cardinality_at_least=interior.count,
        interior_cardinality_exact=interior.exact,
        merged_cells=merged.count,
        merged_ranges=merged.ranges,
        columns=columns,
        rendered_markdown=rendered,
        rendered_unbounded_reason=reason,
        distinct_tracked_max=distinct_tracked_max,
        values_retained_max=values_retained_max,
        header_scan_rows=header_scan_rows,
    )


def _column_name(
    index: int, *, scan: _Scan, header_row: HeaderRowScan, labels: dict
) -> tuple[Optional[str], Optional[str]]:
    """``(name, name_source)`` for one column. The two are decided together on purpose.

    They are one fact in two fields and computing them apart is how they would come to
    disagree — a name with the wrong source is exactly the pair
    :func:`~jmfts_core.sheet_records.header_labels` reads to decide whether the string is a
    record key. The header row wins where there is one; a merged banner names the column
    only where there is not.
    """
    if header_row.row is not None:
        name = scan.head_values.get((header_row.row, index))
        return (name, NAME_SOURCE_HEADER if name else None)
    label = labels.get(index)
    return (label, NAME_SOURCE_BANNER if label else None)


def _banner_ranges(ranges: Sequence, *, scan: _Scan, header_row: HeaderRowScan, cols: int) -> tuple:
    """The merged ranges that are a BANNER over some of the columns rather than a title.

    Three conditions, and each one is a measured property of the file rather than a guess:

    * it spans MORE THAN ONE COLUMN — a merge one column wide is a tall cell;
    * it spans FEWER COLUMNS THAN THE SHEET IS WIDE — a merge across the whole used width is
      the sheet's title, which labels no column in particular and is carried as a
      :class:`BannerRow` instead;
    * its anchor cell holds a value — an empty merge declares a layout and says nothing.

    Where a header row was found, a banner must lie ABOVE it: below the header the same
    shape is a merged data cell. Where none was found, every qualifying range in the scanned
    prefix counts, which is the case step 13 exists for.
    """
    ceiling = header_row.row if header_row.row is not None else None
    return tuple(
        entry
        for entry in ranges
        if entry.cols > 1
        and entry.cols < cols
        and (ceiling is None or entry.max_row < ceiling)
        and scan.head_values.get((entry.min_row, entry.min_col)) is not None
    )


def _banner_labels(ranges: Sequence, *, scan: _Scan) -> dict:
    """``{column index: label}`` from the banner ranges. The LOWEST banner wins.

    ``INGEST_SPEC.md`` 8.5's profile enumerates a column's distinct values, and for a sheet
    with no header row it does so under no label at all — ``Column C has 4 distinct values:
    …`` — which is a sentence an agent cannot use to write its next query. The column letter
    is the honest fallback and is useless for retrieval.

    A merged banner is the one label the FILE states: ``2024 Actuals`` spanning C1:D1 says
    those two columns are that, in the author's own words. The lowest banner wins because a
    stack of them narrows downward — a year over two quarters — and the narrowest is the one
    that describes the column rather than the group it sits in.
    """
    labels: dict = {}
    for entry in sorted(ranges, key=lambda r: (r.max_row, -r.cols)):
        label = scan.head_values.get((entry.min_row, entry.min_col))
        for index in range(entry.min_col, entry.max_col + 1):
            labels[index] = label
    return labels


def _column(
    *,
    index: int,
    name: Optional[str],
    name_source: Optional[str],
    scan: _Scan,
    body_rows: int,
    folded_rows: range,
    with_sketches: bool,
    distinct_tracked_max: int,
    values_retained_max: int,
) -> ColumnMeasurement:
    # THE LIMIT IS THE CALLER'S EVEN ON THE FALLBACK, and the header scan is what made that
    # matter. The scan's accumulators start BELOW `HEADER_SCAN_ROWS`, so on a sheet of eight
    # rows or fewer every column takes this branch and the whole of its distinct set is the
    # fold-in — a default limit here would have silently ignored `distinct_tracked_max` for
    # exactly the sheets a test can build.
    distinct = scan.per_column_distinct.get(index) or _DistinctSet(limit=distinct_tracked_max)
    non_empty = scan.per_column_non_empty.get(index, 0)
    types = dict(scan.per_column_types.get(index, {}))
    folded = [scan.head_values.get((row, index)) for row in folded_rows]
    distinct = distinct.plus(folded)
    for row in folded_rows:
        value = scan.head_values.get((row, index))
        if value is None:
            continue
        non_empty += 1
        kind = scan.head_kinds.get((row, index))
        types[kind] = types.get(kind, 0) + 1
    # Ties break on the type NAME so that two runs over the same bytes cannot disagree;
    # a dict's insertion order is a fact about which row came first, not about the column.
    dominant = max(types, key=lambda kind: (types[kind], kind)) if types else TYPE_EMPTY
    exact = distinct.exact
    count = distinct.count
    # `is_unique` is 8.5's identifier factoid: a distinct value in EVERY row, so a gap
    # disqualifies a column even when what is there is all distinct. A column whose set
    # stopped being tracked cannot answer, and says so rather than guessing.
    if not exact:
        is_unique = None
    else:
        is_unique = bool(body_rows) and non_empty == body_rows and count == body_rows
    values = None
    if exact and count <= values_retained_max:
        values = sorted(distinct.values)
    sketch = None
    if with_sketches and count:
        builder = SketchBuilder()
        # Built from the DISTINCT set. A MinHash is insensitive to repetition, so where the
        # set is exact this is the whole column's sketch; where tracking stopped it is a
        # sketch of the tracked prefix and `partial` says so.
        for value in distinct.values:
            builder.update(value)
        sketch = builder.finish(partial=not exact)
    return ColumnMeasurement(
        index=index,
        letter=column_letter(index),
        name=name,
        name_source=name_source,
        dominant_type=dominant,
        type_counts=types,
        non_empty=non_empty,
        body_rows=body_rows,
        fill_ratio=(non_empty / body_rows) if body_rows else 0.0,
        distinct_count=count if exact else None,
        distinct_at_least=count,
        distinct_exact=exact,
        values=values,
        is_unique=is_unique,
        sketch=sketch,
    )


def _render(scan: _Scan, header_row: HeaderRowScan, rows: int, cols: int) -> tuple:
    """The sheet as one markdown table, or ``None`` and the reason there is none.

    **The banner rows are not in the body, and they are not lost.** A markdown table has one
    header row; putting the rows above it into the body would make a title read as a record.
    They are on :attr:`HeaderRowScan.banner`, the profile node states them in prose, and
    ``SPRINT_0_4_0.md`` 4.3 is answered at :class:`BannerRow` — this is the rendering half of
    that answer and not a decision about the rows.
    """
    if scan.render_bounded:
        return None, (
            "a markdown table of this sheet spends at least one token on each row and one "
            f"on each non-empty cell, so its {rows} rows and {scan.non_empty_cells} filled "
            "cells put its token count above the embedding document window whatever the "
            "tokeniser does; the rendered rows were not kept"
        )
    if not rows or not cols:
        return None, "the sheet's used range holds no value, so there is no table to render"

    lines = [_markdown_row(row, cols) for row in scan.rendered_rows[:rows]]
    first_data_row = header_row.first_data_row
    if header_row.row is not None:
        head = lines[header_row.row - 1]
    else:
        head = _markdown_row(tuple(column_letter(i) for i in range(1, cols + 1)), cols)
    body = lines[first_data_row - 1 :]
    separator = "| " + " | ".join("---" for _ in range(cols)) + " |"
    return "\n".join([head, separator, *body]), None


def _markdown_row(values: tuple, cols: int) -> str:
    padded = list(values[:cols]) + [""] * max(cols - len(values), 0)
    return "| " + " | ".join(_MARKDOWN_UNSAFE.sub(" ", value or "") for value in padded) + " |"


# ---------------------------------------------------------------------------
# Merged cells — tier 1, because tier 2 cannot answer it
# ---------------------------------------------------------------------------


def count_merged_cells(data: bytes, sheet_name: str) -> int:
    """How many merged ranges the sheet declares. 8.3's ``merged_cells``.

    A thin reading of :func:`read_merged_cells`, which is the pass. Kept as its own name
    because 8.3 asks for a count and because this is the call every caller outside this
    module makes.
    """
    return read_merged_cells(data, sheet_name, within_rows=0).count


def read_merged_cells(data: bytes, sheet_name: str, *, within_rows: int) -> MergedCells:
    """The merged ranges the sheet declares, from the worksheet part itself.

    ``ReadOnlyWorksheet`` does not carry ``merged_cells`` (module docstring, note 1), and
    opening the workbook without ``read_only`` to get it would materialise every cell as a
    Python object — which is the cost ``OFFICE_SPEC.md`` Part 4 note 2 exists to avoid. So
    this comes from ``<mergeCells>`` in the part, read with the standard library.

    ``count`` is trusted where the attribute is present and the children are counted where
    it is not; the schema makes the attribute optional, and a writer that omits it has not
    written a broken file. The two can therefore disagree with ``len(ranges)`` and that is a
    fact about the file rather than an error here.

    ``within_rows`` bounds what is RETAINED, not what is parsed: only ranges lying wholly
    inside the first ``within_rows`` rows come back, because only those can be the banner
    :func:`_banner_labels` reads. Zero retains none, which is what
    :func:`count_merged_cells` asks for. The element is fully materialised by the parse that
    counts it either way, so the filter costs nothing the count did not already cost.

    A ``ref`` that is not a two-ended A1 rectangle is SKIPPED rather than repaired. A merge
    is a rectangle by definition; a ``ref`` that is not one came from a writer this module
    cannot interpret, and inventing a rectangle for it would put a banner over columns
    nobody merged.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        path = _worksheet_part(archive, sheet_name)
        with archive.open(path) as stream:
            total = 0
            ranges: list = []
            for _event, element in ET.iterparse(stream, events=("end",)):
                if element.tag == f"{{{_S_NS}}}mergeCells":
                    declared = element.get("count")
                    total = (
                        int(declared)
                        if declared is not None and declared.isdigit()
                        else len(element)
                    )
                    if within_rows:
                        ranges = [
                            entry
                            for child in element
                            if (entry := _merged_range(child.get("ref"))) is not None
                            if entry.max_row <= within_rows
                        ]
                    element.clear()
                    break
                # Rows are the bulk of the part and `mergeCells` follows `sheetData`, so
                # clearing as we go is what keeps this a stream rather than a full parse
                # into memory.
                if element.tag == f"{{{_S_NS}}}row":
                    element.clear()
            return MergedCells(count=total, ranges=tuple(ranges))


def _merged_range(ref: Optional[str]) -> Optional[MergedRange]:
    """``"A5:B6"`` -> a :class:`MergedRange`, or ``None`` for anything else."""
    match = _MERGE_REF.match(ref or "")
    if match is None:
        return None
    first_col, first_row, second_col, second_row = match.groups()
    rows = (int(first_row), int(second_row))
    cols = (column_index(first_col), column_index(second_col))
    return MergedRange(min_row=min(rows), min_col=min(cols), max_row=max(rows), max_col=max(cols))


def _worksheet_part(archive: zipfile.ZipFile, sheet_name: str) -> str:
    """Resolve a sheet name to its part path through the workbook and its relationships.

    openpyxl 3.1.5 leaves ``ReadOnlyWorksheet.path`` as ``None`` and keeps the real path on
    a private attribute, so this walks the two parts that define the mapping rather than
    reading an underscore. They are small — a manifest each — and they are the same two
    parts ``probe`` already reads to count sheets.
    """
    names = set(archive.namelist())
    if "xl/workbook.xml" not in names:
        raise WorksheetPartMissing(
            "the package holds no xl/workbook.xml, so no sheet name resolves to a part"
        )
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationship_id = None
    for element in workbook.iter(f"{{{_S_NS}}}sheet"):
        if element.get("sheetId") is not None and element.get("name") == sheet_name:
            relationship_id = element.get(f"{{{_R_NS}}}id")
            break
    if relationship_id is None:
        raise WorksheetPartMissing(
            f"xl/workbook.xml names no sheet {sheet_name!r} with a relationship id"
        )
    if "xl/_rels/workbook.xml.rels" not in names:
        raise WorksheetPartMissing(
            f"sheet {sheet_name!r} is related through {relationship_id!r} and the package "
            "holds no xl/_rels/workbook.xml.rels to resolve it"
        )
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    for element in rels.iter(f"{{{_PR_NS}}}Relationship"):
        if element.get("Id") != relationship_id:
            continue
        target = element.get("Target") or ""
        path = target[1:] if target.startswith("/") else f"xl/{target}"
        path = path.replace("/./", "/")
        if path not in names:
            raise WorksheetPartMissing(
                f"sheet {sheet_name!r} resolves to {path!r}, which is not in the package"
            )
        return path
    raise WorksheetPartMissing(
        f"sheet {sheet_name!r} names relationship {relationship_id!r}, which "
        "xl/_rels/workbook.xml.rels does not define"
    )


__all__ = [
    "DISTINCT_TRACKED_MAX",
    "HEADER_RULE_ALL_TEXT",
    "HEADER_RULE_ANY_TYPE",
    "HEADER_SCAN_ROWS",
    "NAME_SOURCE_BANNER",
    "NAME_SOURCE_HEADER",
    "TYPE_BOOL",
    "TYPE_DATE",
    "TYPE_EMPTY",
    "TYPE_NUMBER",
    "TYPE_TEXT",
    "VALUES_RETAINED_MAX",
    "BannerRow",
    "ColumnMeasurement",
    "HeaderEvidence",
    "HeaderRowScan",
    "MergedCells",
    "MergedRange",
    "SheetMeasurement",
    "WorksheetPartMissing",
    "canonical",
    "column_index",
    "column_letter",
    "count_merged_cells",
    "measure_sheet",
    "read_merged_cells",
    "value_type",
]
