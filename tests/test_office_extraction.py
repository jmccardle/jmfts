"""``.docx`` and ``.pptx`` become markdown, and the scheduler can plan them.

``INGEST_SPEC.md`` Part 10 step 6. Before this, ``probe`` identified an office file and
nothing downstream could run: ``TEXT_EXTRACTORS`` held ``pdf`` and ``text``, no office
prober reported ``has_text_layer``, and an upload produced a file node with metadata and
no content. These tests cover both halves of closing that — the pattern that lets the row
become eligible, and the reader that makes it worth running.

FIXTURES ARE BUILT BY THE LIBRARIES THEMSELVES, not taken from ``tests/corpus``. The corpus
packages are hand-assembled minimal OOXML for the PROBER, which reads the ZIP directory
with stdlib; they are deliberately missing parts a real reader needs, so python-docx will
not open several of them. What is under test here is the reader, so the input has to be
what a real application writes.

The whole module skips without the ``office`` extra. That is the tier split working: an
install that can probe but not open is a supported deployment, and a suite that failed on
it would be asserting the opposite.
"""

from __future__ import annotations

import io

import pytest

# The readers are the `office` extra. `dev` implies it (pyproject.toml), so the suite job
# runs this file and the base-install job skips it.
pytest.importorskip("docx", reason="the office extra is not installed")
pytest.importorskip("pptx", reason="the office extra is not installed")

import jmfts_core.ingest_tasks  # noqa: E402,F401  (import order; see the module cycle)
from docx import Document  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from jmfts_core.ingest_tasks import explain_plan  # noqa: E402
from jmfts_core.office.extract import (  # noqa: E402
    OfficeReadError,
    docx_to_markdown,
    pptx_to_markdown,
)
from jmfts_core.probe import detect_format, probe_patterns  # noqa: E402
from jmfts_core.structure_tasks import (  # noqa: E402
    DECLARED_SPLITTERS,
    EXTRACTION_DOCX_BODY,
    EXTRACTION_PPTX_SLIDES,
    INFERRED_SPLITTERS,
    TEXT_EXTRACTORS,
)


def _docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(build) -> bytes:
    presentation = Presentation()
    build(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _plan_for(data: bytes, filename: str):
    """What the scheduler would do with these bytes, from the patterns probe measures."""
    detection = detect_format(data, filename=filename, declared_mime=None)
    patterns, _ = probe_patterns(data, detection)
    plan = explain_plan(fmt=detection.format, patterns=patterns, patterns_source="probed")
    return patterns, {task.task: task for task in plan.tasks}


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------


def test_docx_heading_styles_become_atx_headings():
    """The rung that splits a ``.docx`` reads the headings this writes.

    ``DECLARED_STRUCTURE_PATTERN`` gates ``structure:declared`` on ``has_heading_styles``,
    which ``probe`` measures from ``word/styles.xml``. If the reader did not turn those
    same styles into ATX headings, the rung would run and find one untitled region.
    """

    def build(document):
        document.add_heading("Chapter One", level=1)
        document.add_paragraph("Opening prose.")
        document.add_heading("Method", level=2)

    markdown, meta = docx_to_markdown(_docx_bytes(build))
    assert "# Chapter One" in markdown
    assert "## Method" in markdown
    assert "Opening prose." in markdown
    assert meta["heading_count"] == 2


def test_docx_tables_survive_because_the_body_is_walked_in_order():
    """``Document.paragraphs`` omits tables entirely.

    A document whose content is a table would extract as an empty string — a file node
    holding nothing, from bytes that held something. This is why the reader iterates the
    body element rather than that property.
    """

    def build(document):
        document.add_paragraph("Before the table.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "metric"
        table.cell(0, 1).text = "value"
        table.cell(1, 0).text = "nDCG@10"
        table.cell(1, 1).text = "0.7090"
        document.add_paragraph("After the table.")

    markdown, meta = docx_to_markdown(_docx_bytes(build))
    assert meta["table_count"] == 1
    assert "| metric | value |" in markdown
    assert "| nDCG@10 | 0.7090 |" in markdown
    # Order preserved: the table sits between the two paragraphs, not after both.
    assert markdown.index("Before the table.") < markdown.index("| metric")
    assert markdown.index("| metric") < markdown.index("After the table.")


def test_docx_list_styles_are_list_items_even_with_no_numbering_on_the_paragraph():
    """Both list signals are needed; neither subsumes the other.

    Word usually stamps ``w:numPr`` on the paragraph. A document built through
    ``add_paragraph(style="List Bullet")`` carries the numbering in the STYLE and nothing
    on the paragraph, and checking only the paragraph extracts the list as flat prose.
    """

    def build(document):
        document.add_paragraph("First item", style="List Bullet")
        document.add_paragraph("Second item", style="List Bullet")

    markdown, meta = docx_to_markdown(_docx_bytes(build))
    assert meta["list_item_count"] == 2
    assert "- First item" in markdown
    # Consecutive items are one tight list, not separated by blank lines.
    assert "- First item\n- Second item" in markdown


def test_docx_prose_that_looks_like_a_heading_is_escaped():
    """A body paragraph beginning "## " must not become a chunk boundary.

    ``split_on_headings`` cuts on ATX headings, so an unescaped one would create a section
    the author never wrote, out of prose that merely started with a hash.
    """

    def build(document):
        document.add_paragraph("## not a heading, just prose")

    markdown, _ = docx_to_markdown(_docx_bytes(build))
    assert "\\## not a heading" in markdown
    assert "\n## not a heading" not in markdown


def test_an_empty_docx_reports_no_text_layer_and_the_row_does_not_run():
    """The refusal that keeps an empty node from being created silently.

    A ``.docx`` with no text run extracts to an empty string. ``extract:text`` requires
    ``has_text_layer``, so the honest outcome is a row that never runs with a reason
    naming the pattern — not a row that runs and writes nothing.
    """
    data = _docx_bytes(lambda document: None)
    patterns, tasks = _plan_for(data, "empty.docx")
    assert patterns["has_text_layer"] is False
    assert tasks["extract:text"].outcome == "not_applicable"
    assert "has_text_layer" in tasks["extract:text"].reason


def test_a_docx_with_headings_plans_extract_and_the_declared_rung():
    """The end the whole change exists for: an office file that actually ingests."""

    def build(document):
        document.add_heading("Chapter One", level=1)
        document.add_paragraph("Opening prose.")

    patterns, tasks = _plan_for(_docx_bytes(build), "real.docx")
    assert patterns["has_text_layer"] is True
    assert patterns["has_heading_styles"] is True
    assert tasks["extract:text"].outcome == "enqueued"
    assert tasks["structure:declared"].outcome == "enqueued"


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------


def test_pptx_writes_one_heading_per_slide_and_does_not_repeat_the_title():
    """``shapes.title`` builds a fresh proxy on every access.

    Comparing shapes by identity therefore never matches, and the title is written twice —
    once as the heading and once as a bullet beneath it. The reader compares elements.
    """

    def build(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Retrieval Results"
        slide.placeholders[1].text_frame.text = "Hybrid wins three of four"

    markdown, meta = pptx_to_markdown(_pptx_bytes(build))
    assert markdown.startswith("# Retrieval Results")
    assert meta["slide_count"] == 1
    assert markdown.count("Retrieval Results") == 1


def test_a_slide_with_no_title_placeholder_promotes_its_first_line():
    """Measured on 22 real EU decks: 80 of 220 slides carry no title placeholder.

    Those decks are built from a blank layout with the title typed into an ordinary text
    box. Falling straight through to "Slide N" would throw away the heading that
    ``split_on_headings`` turns into the chunk's section title — the label a search result
    is read by — on more than a third of real slides.

    The promoted line must not also appear in the body, or every such slide says its own
    title twice.
    """

    def build(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
        frame.text = "Quarterly Retrieval Results"
        frame.add_paragraph().text = "Hybrid wins three of four"

    markdown, meta = pptx_to_markdown(_pptx_bytes(build))
    assert markdown.startswith("# Quarterly Retrieval Results")
    assert meta["promoted_title_count"] == 1
    assert markdown.count("Quarterly Retrieval Results") == 1
    # The rest of the promoted shape still belongs to the body.
    assert "- Hybrid wins three of four" in markdown


def test_a_slide_with_no_text_at_all_is_numbered():
    """ "Slide N" remains the answer when there is no text to promote.

    An image-only slide has no line to take, and inventing one would be worse than a
    number. This is the case the fallback still exists for.
    """

    def build(presentation):
        presentation.slides.add_slide(presentation.slide_layouts[6])

    markdown, meta = pptx_to_markdown(_pptx_bytes(build))
    assert "# Slide 1" in markdown
    assert meta["promoted_title_count"] == 0


def test_a_real_title_placeholder_is_not_counted_as_promoted():
    """The count must mean what it says, or it cannot be read from an attempt log."""

    def build(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "An Authored Title"

    markdown, meta = pptx_to_markdown(_pptx_bytes(build))
    assert markdown.startswith("# An Authored Title")
    assert meta["promoted_title_count"] == 0


def test_speaker_notes_are_kept_under_their_own_subheading():
    """Notes are authored text; dropping them loses the argument half of many decks."""

    def build(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Results"
        slide.notes_slide.notes_text_frame.text = "Mention the TREC-COVID caveat."

    markdown, meta = pptx_to_markdown(_pptx_bytes(build))
    assert meta["notes_slide_count"] == 1
    assert "## Notes" in markdown
    assert "Mention the TREC-COVID caveat." in markdown


def test_slide_offsets_land_on_their_own_headings():
    """``slide_offsets`` is the record's index INTO the text, like the PDF page map.

    An offset that is off by the separator is the kind of error nothing downstream would
    catch, because every offset would still point somewhere plausible.
    """

    def build(presentation):
        for title in ("First", "Second", "Third"):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = title

    markdown, meta = pptx_to_markdown(_pptx_bytes(build))
    assert meta["slide_count"] == 3
    assert len(meta["slide_offsets"]) == 3
    for offset, title in zip(meta["slide_offsets"], ("First", "Second", "Third")):
        assert markdown[offset:].startswith(f"# {title}")


def test_a_deck_with_slides_plans_extract_and_the_declared_rung():
    def build(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Results"
        slide.placeholders[1].text_frame.text = "Hybrid wins"

    patterns, tasks = _plan_for(_pptx_bytes(build), "real.pptx")
    assert patterns["has_text_layer"] is True
    assert patterns["has_slides"] is True
    assert tasks["extract:text"].outcome == "enqueued"
    assert tasks["structure:declared"].outcome == "enqueued"


# ---------------------------------------------------------------------------
# wiring and failure
# ---------------------------------------------------------------------------


def test_both_formats_are_registered_end_to_end():
    """A reader nothing dispatches to is a reader that never runs.

    Three tables have to agree: the extractor, and both structure rungs' splitters. A
    source missing from either splitter map raises at the rung rather than here.
    """
    assert set(TEXT_EXTRACTORS) >= {"pdf", "text", "docx", "pptx"}
    for source in (EXTRACTION_DOCX_BODY, EXTRACTION_PPTX_SLIDES):
        assert source in DECLARED_SPLITTERS
        assert source in INFERRED_SPLITTERS


def test_the_extractors_produce_the_source_the_splitters_are_keyed_on():
    """The wrappers and the dispatch tables must name the same string."""

    def build(document):
        document.add_paragraph("Prose.")

    assert TEXT_EXTRACTORS["docx"](_docx_bytes(build)).source == EXTRACTION_DOCX_BODY

    def build_deck(presentation):
        presentation.slides.add_slide(presentation.slide_layouts[6])

    assert TEXT_EXTRACTORS["pptx"](_pptx_bytes(build_deck)).source == EXTRACTION_PPTX_SLIDES


def test_a_package_whose_member_will_not_inflate_is_a_permanent_probe_failure():
    """``_open_package`` guards the OPEN. A member can still fail after that.

    ``tests/corpus``'s ``central-directory-mismatch.docx`` is exactly this shape: the
    central directory is complete, ``namelist()`` answers, ``detect_format`` says docx, and
    the archive only falls apart when a part is read. The raw ``zipfile.BadZipFile`` that
    escaped derives straight from ``Exception``, which ``classify_exception`` grades
    RETRYABLE — so probe attempted these bytes three times, and a local header does not
    repair itself between attempts.

    Found by running the real ingestion path over the corpus rather than by reading it.
    """
    import zipfile

    from jmfts_core.task_errors import ErrorType, classify_exception
    from tests.corpus import fixtures

    data = fixtures.central_directory_mismatch_docx()
    detection = detect_format(data, filename="central-directory-mismatch.docx", declared_mime=None)
    assert detection.format == "docx", "the fixture must still LOOK like a docx"

    with pytest.raises(ValueError) as caught:
        probe_patterns(data, detection)
    assert not isinstance(caught.value, zipfile.BadZipFile)
    assert classify_exception(caught.value) is ErrorType.PERMANENT


@pytest.mark.parametrize(
    "reader,label",
    [(docx_to_markdown, "docx"), (pptx_to_markdown, "pptx")],
)
def test_unopenable_bytes_are_a_permanent_failure(reader, label):
    """``OfficeReadError`` is a ValueError so ``classify`` calls it PERMANENT.

    The default for an unrecognised exception is RETRYABLE, which would spend three
    attempts re-opening a package that is not going to open.
    """
    from jmfts_core.task_errors import ErrorType, classify_exception

    with pytest.raises(OfficeReadError) as caught:
        reader(b"PK\x03\x04 this is not a real package")
    assert label in str(caught.value)
    assert classify_exception(caught.value) is ErrorType.PERMANENT
