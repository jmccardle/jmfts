"""Office packages for the REAL-LibreOffice rendition tests, and the converter that renders them.

``tests/test_render_anchors.py`` proves the alignment, the rows and the refusals against a
rendition PyMuPDF wrote out of the package's own strings — which cannot measure the one number
``docs/OFFICE_SPEC.md`` Part 10 is for: how much of a real ``.docx`` or ``.pptx`` the
alignment finds in what LibreOffice really draws. These builders are that measurement's
inputs, and :func:`render_with_soffice` is the renderer.

**NOTHING HERE IS COMMITTED AS BYTES**, for ``tests/corpus/fixtures.py``'s reason: a binary
nobody can review is a fixture nobody can trust. Each package is built from string literals
by python-docx or python-pptx at test time, and rendered by the ``soffice`` on ``PATH``.

Each case exercises one thing a real renderer does that the PyMuPDF fixture cannot:

* ``docx_bullets`` / ``docx_numbered`` — LibreOffice draws ``•``, an OpenSymbol private-use
  glyph, or ``1.`` where the reader writes ``- `` (``/tmp`` audit item #5).
* ``docx_table`` — a table, with a horizontally merged cell python-docx repeats (#16).
* ``docx_merged`` — short paragraphs with no spacing, which PyMuPDF reads as ONE block (#10).
* ``docx_hyphenation`` — automatic hyphenation in a narrow column, soft hyphens the author
  typed, and words a renderer may set with ligatures (#14).
* ``pptx_deck`` — titled and untitled slides, bullets at two levels, a table, and speaker
  notes, which a slide rendition never draws (#21).
"""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path

#: The renderer. ``soffice`` is what ``Dockerfile.libreoffice`` runs behind unoserver, and
#: invoking it directly is the same conversion engine without the server in between.
SOFFICE = shutil.which("soffice")

# A long-enough body sentence that a renderer wraps it, used to pad several cases.
_LONG = (
    "Retrieval quality depends on how faithfully a passage can be traced back to the page "
    "it was printed on, which is the whole purpose of a rendition anchor."
)


def docx_bullets() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Bulleted findings", level=1)
    document.add_paragraph("The review produced the following observations about the corpus.")
    for item in (
        "Every office upload is rendered once and the rendition is kept beside the source.",
        "Short bullet",
        "A passage whose text the renderer prints after a glyph the reader never wrote.",
        "Bullets nested under a parent item are indented further on the page.",
    ):
        document.add_paragraph(item, style="List Bullet")
    document.add_paragraph("A closing paragraph after the list returns to ordinary prose.")
    return _save(document)


def docx_numbered() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Numbered procedure", level=1)
    for item in (
        "Upload the package and wait for the rendition to be written.",
        "Run the anchor task over every chunk the structure rungs produced.",
        "Open the preview and compare the rectangle with the passage text.",
        "Repeat after re-rendering, and confirm the stale anchors are refused.",
    ):
        document.add_paragraph(item, style="List Number")
    document.add_paragraph("That is the whole procedure, and it needs no further steps.")
    return _save(document)


def docx_table() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Quarterly table", level=1)
    document.add_paragraph("The table below summarises the three regions by quarter.")
    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    rows = [
        ["Region", "First quarter", "Second quarter"],
        ["North", "1,204 units", "1,388 units"],
        ["South", "987 units", "1,021 units"],
        ["Combined total across both reporting periods", "", ""],
    ]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    # A horizontal merge: python-docx reports the merged cell once per grid column it spans.
    table.cell(3, 0).merge(table.cell(3, 1))
    table.cell(3, 2).text = "4,600 units"
    document.add_paragraph("Figures are provisional until the audit closes in the spring.")
    return _save(document)


def docx_table_borderless() -> bytes:
    """A table with no ruling lines, which ``find_tables`` has nothing to detect it by."""
    from docx import Document

    document = Document()
    document.add_paragraph("The schedule below lists each stage and the team that owns it.")
    table = document.add_table(rows=3, cols=2)
    for r, values in enumerate(
        [["Stage", "Owner"], ["Rendering", "Platform team"], ["Anchoring", "Search team"]]
    ):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    document.add_paragraph("Ownership changes are announced a week before they take effect.")
    return _save(document)


def docx_merged() -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading("Address block", level=1)
    for line in (
        "Office of the Registrar",
        "Building Seven, Second Floor",
        "Long Road, Riverside District",
        "Telephone enquiries are answered on weekday mornings only.",
    ):
        paragraph = document.add_paragraph(line)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
    document.add_paragraph(_LONG)
    return _save(document)


def docx_hyphenation() -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    document = Document()
    settings = document.settings.element
    auto = OxmlElement("w:autoHyphenation")
    auto.set(qn("w:val"), "true")
    settings.append(auto)
    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(2.6)

    document.add_heading("Hyphenation", level=1)
    document.add_paragraph(
        "Internationalization and telecommunications infrastructure requirements are "
        "characteristically incomprehensible to uninitiated representatives of "
        "interdisciplinary organizations, notwithstanding comprehensive documentation."
    )
    # Soft hyphens the author typed: invisible unless the line breaks there.
    document.add_paragraph(
        "The extra­ordinary circum­stances of the counter­productive "
        "re­organization were docu­mented in the supple­mentary report."
    )
    # Words a renderer may set with fi/fl/ffi/ffl ligatures.
    document.add_paragraph(
        "The office affirmed its official findings: an efficient workflow, affluent "
        "suppliers, and a fluffy waffle served at the final staff briefing."
    )
    return _save(document)


def pptx_deck() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    title_and_body = deck.slide_layouts[1]
    blank = deck.slide_layouts[6]

    slide = deck.slides.add_slide(title_and_body)
    slide.shapes.title.text = "Quarterly review"
    body = slide.placeholders[1].text_frame
    body.text = "Revenue grew in every region we operate in"
    for text, level in (
        ("North led the growth with new contracts", 1),
        ("Costs held flat against the prior quarter", 0),
        ("Hiring is paused until the audit closes", 0),
    ):
        paragraph = body.add_paragraph()
        paragraph.text = text
        paragraph.level = level
    slide.notes_slide.notes_text_frame.text = (
        "Speaker notes: mention the two delayed contracts before the questions start."
    )

    slide = deck.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame
    box.text = "An untitled slide whose first line becomes its heading"
    paragraph = box.add_paragraph()
    paragraph.text = "and a second line in the same text box below it"

    slide = deck.slides.add_slide(blank)
    shape = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    for r, values in enumerate([["Metric", "Value"], ["Churn", "2.1%"], ["Margin", "34%"]]):
        for c, value in enumerate(values):
            shape.table.cell(r, c).text = value

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


#: name -> (builder, suffix). The suffix is what ``soffice`` picks its import filter by.
CASES = {
    "docx_bullets": (docx_bullets, "docx"),
    "docx_numbered": (docx_numbered, "docx"),
    "docx_table": (docx_table, "docx"),
    "docx_table_borderless": (docx_table_borderless, "docx"),
    "docx_merged": (docx_merged, "docx"),
    "docx_hyphenation": (docx_hyphenation, "docx"),
    "pptx_deck": (pptx_deck, "pptx"),
}


def render_with_soffice(data: bytes, suffix: str, *, timeout: float = 120.0) -> bytes:
    """``data`` as a PDF, rendered by the ``soffice`` on ``PATH``.

    A throwaway user profile per call, because two ``soffice`` processes sharing the default
    profile block on its lock — and a test run in parallel with a desktop LibreOffice would
    hang rather than fail.
    """
    if SOFFICE is None:
        raise RuntimeError("soffice is not on PATH")
    with tempfile.TemporaryDirectory(prefix="jmfts-soffice-") as workdir:
        work = Path(workdir)
        source = work / f"input.{suffix}"
        source.write_bytes(data)
        profile = (work / "profile").as_uri()
        completed = subprocess.run(
            [
                SOFFICE,
                f"-env:UserInstallation={profile}",
                "--headless",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                str(work),
                str(source),
            ],
            capture_output=True,
            timeout=timeout,
            check=False,
        )
        produced = work / "input.pdf"
        if completed.returncode != 0 or not produced.exists():
            raise RuntimeError(
                f"soffice exited {completed.returncode} and wrote no PDF: "
                f"{completed.stderr.decode(errors='replace')[-2000:]}"
            )
        return produced.read_bytes()
