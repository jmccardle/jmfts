"""Rows as records. ``INGEST_SPEC.md`` 8.4's ``records`` shape, end to end.

Three things are under test and they are different in kind.

The **reader** (:mod:`jmfts_core.office.cells`) turns bytes into typed values plus the two
facts read-only ``openpyxl`` cannot carry: the formula text and the forced-text flag. No
database.

The **builder** (:mod:`jmfts_core.sheet_records`) turns those rows plus a stored header
into 8.4's three forms. A pure function, so no database and no reader.

The **task**, on the real queue: upload, drain, and look at the tree. That is the half that
would notice a record written under the wrong parent, a header taken from the wrong row, or
a sheet whose shape verdict disagrees with the nodes underneath it.

FIXTURES ARE BUILT BY ``openpyxl`` for the reason ``tests/test_sheet_tasks.py`` gives: the
hand-assembled packages in ``tests/corpus`` exist for the PROBER, which reads the ZIP
directory with the standard library, and that guarantee is worth keeping separate.

**One openpyxl behaviour shapes several fixtures below.** It writes a formula's text and
leaves ``<v/>`` empty, because it does not evaluate. A real workbook saved by Excel carries
the cached result beside the formula. So a formula cell in these fixtures has a note and no
value, which is exactly what the reader should report for a workbook that carries no cached
result — see ``measure_sheet``'s ``data_only`` comment for the same trade.
"""

from __future__ import annotations

import datetime
import io

import pytest
from sqlalchemy import select

pytest.importorskip("openpyxl", reason="the office extra is not installed")

import openpyxl  # noqa: E402

import jmfts_core.ingest_tasks  # noqa: E402,F401  (import order; see the module cycle)
from jmfts_client.contracts.upload import UploadedFile  # noqa: E402
from jmfts_core.ingest_tasks import TASK_EXTRACT_SHEET  # noqa: E402
from jmfts_core.ingest_options import resolve_options  # noqa: E402
from jmfts_core.models.document import (  # noqa: E402
    Document,
    SETTLED_SETTLED,
    USETYPE_RECORD,
    USETYPE_TABLE,
)
from jmfts_core.office.cells import (  # noqa: E402
    ROWS_READ_MAX,
    TooManyRows,
    json_value,
    read_rows,
)
from jmfts_core.office.sheets import (  # noqa: E402
    NAME_SOURCE_BANNER,
    NAME_SOURCE_HEADER,
    measure_sheet,
)
from jmfts_core.repositories.document import DocumentRepository  # noqa: E402
from jmfts_core.repositories.evidence import EvidenceRepository  # noqa: E402
from jmfts_core.services.ingest_service import IngestService  # noqa: E402
from jmfts_core.sheet_profile import sheet_evidence_block  # noqa: E402
from jmfts_core.sheet_records import (  # noqa: E402
    SHAPE_RECORDS,
    SHAPE_SMALL_TABLE,
    HeaderDoesNotCoverTheRow,
    build_records,
    header_labels,
)
from jmfts_core.models.document import USETYPE_SHEET  # noqa: E402
from jmfts_core.sheet_tasks import run_extract_sheet  # noqa: E402
from jmfts_core.structure_tasks import RUNG_INFERRED  # noqa: E402
from tests.conftest import drain_ingest_queue  # noqa: E402

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _workbook_bytes() -> bytes:
    """One sheet of records and one sheet with no header row.

    ``Deals`` holds every type a cell can be, a blank row in the middle so that a row
    NUMBER cannot be confused with a position, a forced-text identifier, and a formula.
    ``Notes`` has a value but no header row under either pass of the scan — every row of it
    has a hole — which is the branch that produces no RECORDS. It still produces a `table`
    node, and that is step 12: on the open web the sheets that match `small_table` and not
    `records` are 56.0% of all of them, against the 4.6% that match `records` and not it.
    """
    workbook = openpyxl.Workbook()
    deals = workbook.active
    deals.title = "Deals"
    deals.append(["Deal ID", "Account", "Value", "Close Date", "Won"])
    deals.append(["D-4471", "Northwind Freight", 128000, datetime.datetime(2026, 9, 30), True])
    deals.append(["D-4472", "Contoso", 96500.5, datetime.datetime(2026, 10, 15), False])
    deals.append([None, None, None, None, None])
    deals["A5"] = "0012345"
    # The apostrophe Excel shows in the formula bar. It is a style flag, not part of the
    # value, and read-only openpyxl has no attribute for it at all.
    deals["A5"].quotePrefix = True
    deals["B5"] = "Fabrikam"
    deals["C5"] = "=C2+C3"
    deals["E5"] = True

    notes = workbook.create_sheet("Notes")
    notes["B3"] = "A paragraph of free text, with no header row above it."

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def workbook_bytes() -> bytes:
    return _workbook_bytes()


def _header(data: bytes, sheet: str) -> dict:
    """The header the way the task gets it: off what ``profile:sheet`` measured.

    Both halves, because the ROW is now a measurement too — a header at row 4 makes rows 1
    to 3 a banner and not instances, and `build_records` cannot skip them without being
    told which row it was.
    """
    measured = measure_sheet(data, sheet, render_cell_budget=8192)
    return {
        "header": [column.name for column in measured.columns],
        "header_row": measured.header_row.row,
    }


def _upload(session, data: bytes, filename: str = "records.xlsx"):
    return IngestService(session).upload_file(
        UploadedFile(data=data, filename=filename, content_type=XLSX_MIME)
    )


def _ingest(session, data: bytes, filename: str = "records.xlsx") -> Document:
    response = _upload(session, data, filename)
    drain_ingest_queue(session)
    return DocumentRepository(session).get(response.document_id)


def _children(session, node_id: int, usetype: str) -> list:
    return list(
        session.execute(
            select(Document)
            .where(Document.parent_id == node_id, Document.usetype == usetype)
            .order_by(Document.id)
        )
        .scalars()
        .all()
    )


def _sheet_node(session, file_node: Document, name: str) -> Document:
    repo = EvidenceRepository(session)
    return next(
        node
        for node in _children(session, file_node.id, USETYPE_SHEET)
        if (repo.read(node.id, "sheet") or {}).get("name") == name
    )


class _ScopedTask:
    """The two fields a per-sheet handler reads off its queue row.

    ``params`` is the RESOLVED ``sheet_records`` group with the test's overrides on top,
    not a bare dict. Since ``SPRINT_JOBS.md`` Phase 3 the group is complete on every queue
    row `plan_frontier` writes, and ``run_extract_sheet`` reads ``params["max_rows"]``
    rather than falling back to a default of its own — so a stand-in row carrying one key
    would exercise a shape the queue never produces.
    """

    def __init__(self, document_id: int, params: dict | None = None):
        self.scope_document_id = document_id
        self.params = {**resolve_options("xlsx")["sheet_records"], **(params or {})}


# ---------------------------------------------------------------------------
# The reader
# ---------------------------------------------------------------------------


class TestJsonValue:
    """The one difference from `canonical`: a number stays a number."""

    def test_a_whole_float_becomes_an_int(self):
        """A spreadsheet has one numeric type and stores every number as a double, so
        reporting `1.0` would be reporting the storage rather than the value."""
        assert json_value(1.0) == 1
        assert isinstance(json_value(1.0), int)
        assert json_value(1.5) == 1.5

    def test_a_boolean_stays_a_boolean(self):
        """Before the int branch: `bool` is a subclass of `int` and a flag column must not
        come back as ones and zeroes."""
        assert json_value(True) is True
        assert json_value(False) is False

    def test_a_midnight_datetime_is_a_plain_date(self):
        assert json_value(datetime.datetime(2026, 9, 30)) == "2026-09-30"
        assert json_value(datetime.datetime(2026, 9, 30, 14, 5)) == "2026-09-30T14:05:00"

    def test_whitespace_only_is_empty(self):
        """The same call `canonical` makes: three spaces is not a value a person put there
        to mean something, and counting it as one would make every ratio depend on a
        workbook's whitespace."""
        assert json_value("   ") is None
        assert json_value("  D-4471 ") == "D-4471"


class TestReadRows:
    def test_values_come_back_typed(self, workbook_bytes):
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        row = next(row for row in rows.rows if row.index == 2)
        assert row.values == ("D-4471", "Northwind Freight", 128000, "2026-09-30", True)

    def test_a_blank_row_yields_nothing_and_does_not_renumber_the_rest(self, workbook_bytes):
        """Row 4 is empty. The rows after it keep their worksheet numbers, because the
        number is how a person finds the row again in Excel."""
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        assert [row.index for row in rows.rows] == [1, 2, 3, 5]

    def test_the_forced_text_flag_is_read_from_the_package(self, workbook_bytes):
        """openpyxl's read-only cells have no `quotePrefix` attribute at all, so this is
        the standard-library pass. The VALUE was never in doubt — it is stored as a string
        and keeps its leading zeros — and what the flag adds is that the author meant it."""
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        row = next(row for row in rows.rows if row.index == 5)
        assert row.values[0] == "0012345"
        assert row.notes[1].text_forced is True

    def test_the_formula_text_is_read_beside_the_value(self, workbook_bytes):
        """`data_only` is an either/or in openpyxl and the worksheet part holds both, so
        the value comes from the reader and the formula from the same pass as the flag."""
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        row = next(row for row in rows.rows if row.index == 5)
        assert row.notes[3].formula == "C2+C3"
        assert row.notes[3].formula_shared is False

    def test_a_shared_formula_takes_the_masters_text_and_says_so(self):
        """Excel writes the text once per group and leaves the members as back references.
        openpyxl never emits that form, so the bytes are assembled here."""
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Calc"
        sheet.append(["n", "doubled"])
        for row in range(2, 5):
            sheet.cell(row, 1).value = row
        buffer = io.BytesIO()
        workbook.save(buffer)
        data = _with_shared_formula(buffer.getvalue())

        rows = read_rows(data, "Calc", max_rows=100)
        notes = {row.index: row.notes.get(2) for row in rows.rows}
        assert notes[2].formula == "A2*2" and notes[2].formula_shared is True
        # The member's own element carries no text, so what it gets is the master's — with
        # the master's cell references, untranslated, and the flag says so.
        assert notes[3].formula == "A2*2" and notes[3].formula_shared is True

    def test_notes_are_not_read_when_the_caller_says_not_to(self, workbook_bytes):
        rows = read_rows(workbook_bytes, "Deals", max_rows=100, with_notes=False)
        assert rows.notes_read is False
        assert all(not row.notes for row in rows.rows)

    def test_more_rows_than_the_bound_raises_rather_than_truncating(self, workbook_bytes):
        """6.6: a named limit that fails the task, not a silent truncation. A caller that
        got the first two rows of a four-row sheet cannot tell that from a two-row sheet."""
        with pytest.raises(TooManyRows, match="more than 2 rows"):
            read_rows(workbook_bytes, "Deals", max_rows=2)

    def test_a_bound_above_the_readers_ceiling_is_refused_at_a_stated_number(self, workbook_bytes):
        with pytest.raises(ValueError, match="ROWS_READ_MAX"):
            read_rows(workbook_bytes, "Deals", max_rows=ROWS_READ_MAX + 1)


def _with_shared_formula(data: bytes) -> bytes:
    """Rewrite column B as a shared formula group, which openpyxl will not write itself."""
    import re
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace(
        '<c r="A2" t="n"><v>2</v></c>',
        '<c r="A2" t="n"><v>2</v></c><c r="B2"><f t="shared" ref="B2:B4" si="0">A2*2</f></c>',
    )
    for row in (3, 4):
        sheet = re.sub(
            rf'(<c r="A{row}" t="n"><v>{row}</v></c>)',
            rf'\1<c r="B{row}"><f t="shared" si="0"/></c>',
            sheet,
        )
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# The builder
# ---------------------------------------------------------------------------


class TestBuildRecords:
    def test_the_header_row_is_not_a_record(self, workbook_bytes):
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        records = build_records(rows, **_header(workbook_bytes, "Deals"))
        assert [record.row_index for record in records] == [2, 3, 5]

    def test_a_record_holds_the_typed_values_under_the_column_names(self, workbook_bytes):
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        records = build_records(rows, **_header(workbook_bytes, "Deals"))
        assert records[0].record == {
            "Deal ID": "D-4471",
            "Account": "Northwind Freight",
            "Value": 128000,
            "Close Date": "2026-09-30",
            "Won": True,
        }

    def test_content_is_labelled_prose_and_not_json(self, workbook_bytes):
        """8.4: `content` is what gets embedded, so braces and quotes would be tokens
        spent on syntax the model gains nothing from."""
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        records = build_records(rows, **_header(workbook_bytes, "Deals"))
        assert records[0].content == (
            "Deal ID: D-4471. Account: Northwind Freight. Value: 128000. "
            "Close Date: 2026-09-30. Won: TRUE."
        )

    def test_an_empty_cell_is_absent_rather_than_null(self, workbook_bytes):
        """Row 5 has no close date. A key holding null and an absent key say the same
        thing about a spreadsheet, and the absent one does not spend a token saying it."""
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        record = build_records(rows, **_header(workbook_bytes, "Deals"))[2]
        assert "Close Date" not in record.record
        assert record.record["Deal ID"] == "0012345"

    def test_notes_land_under_the_column_name_and_only_where_there_are_any(self, workbook_bytes):
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        records = build_records(rows, **_header(workbook_bytes, "Deals"))
        assert records[0].cells == {}
        assert records[2].cells == {
            "Deal ID": {"text_forced": True},
            "Value": {"formula": "C2+C3"},
        }

    def test_the_header_comes_off_the_stored_profile_positionally(self):
        """`header_labels` is the whole coupling between the two tasks. It reads the
        stored `columns` ARRAY in order, so `header[0]` names column A — and a column the
        profile left unnamed stays unnamed here rather than being invented."""
        stored = [
            {"name": "Deal ID", "name_source": NAME_SOURCE_HEADER},
            {"name": "Account", "name_source": NAME_SOURCE_HEADER},
            {"name": None, "name_source": None},
        ]
        assert header_labels(stored) == ["Deal ID", "Account", None]

    def test_a_banner_label_is_not_a_record_key(self):
        """Step 13 gives a header-less sheet's columns a label read from a merged banner,
        and a banner names the GROUP the merge spans — three columns under one banner share
        one string, and three record keys that are one string are one key holding the last
        value. So the SOURCE is checked, not the name."""
        stored = [
            {"name": "2024 Actuals", "name_source": NAME_SOURCE_BANNER},
            {"name": "2024 Actuals", "name_source": NAME_SOURCE_BANNER},
        ]
        assert header_labels(stored) == [None, None]

    def test_a_value_the_header_does_not_name_raises(self, workbook_bytes):
        """The profile and this read ran over the same bytes with the same reader, so a
        disagreement is a real one. Dropping the value would put a record into the tree
        that is missing a field nobody can see is missing."""
        rows = read_rows(workbook_bytes, "Deals", max_rows=100)
        with pytest.raises(HeaderDoesNotCoverTheRow, match="column 4"):
            build_records(rows, header=["Deal ID", "Account", "Value"], header_row=1)


# ---------------------------------------------------------------------------
# The task, on the real queue
# ---------------------------------------------------------------------------


class TestExtractSheetTask:
    def test_upload_produces_one_record_node_per_row(self, db_session, evidence, workbook_bytes):
        """The whole claim, from bytes to nodes: upload, drain, and the rows are there."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")
        records = _children(db_session, sheet.id, USETYPE_RECORD)

        assert [evidence(node)["row_index"] for node in records] == [2, 3, 5]
        first = evidence(records[0])
        assert first["record"]["Value"] == 128000
        assert first["record"]["Won"] is True
        assert first["sheet_name"] == "Deals"

    def test_a_record_node_keeps_the_worksheet_row_number_in_its_title(
        self, db_session, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")
        records = _children(db_session, sheet.id, USETYPE_RECORD)
        assert [node.title for node in records] == [
            "Deals row 2",
            "Deals row 3",
            "Deals row 5",
        ]

    def test_the_sheet_node_records_the_shape_and_its_rung(
        self, db_session, evidence, workbook_bytes
    ):
        """8.7: `rung` is `inferred` for every shape except `unstructured`. The sheet node
        itself was produced at the declared rung; the shape below it was not."""
        file_node = _ingest(db_session, workbook_bytes)
        node = _sheet_node(db_session, file_node, "Deals")
        block = evidence(node)["sheet"]

        # BOTH shapes, and `small_table` first because 8.4 lists it first. `Deals` has a
        # header row AND renders inside the document window, which is the contested
        # population step 12 is about — 29.4% of git-corpora sheets and 12.5% of FUSE ones.
        assert block["shape"] == [SHAPE_SMALL_TABLE, SHAPE_RECORDS]
        assert block["rung"] == RUNG_INFERRED
        assert block["record_count"] == 3
        # A margin says how close a decision was to a boundary, and neither input has one:
        # a boolean has no boundary, and a token count against the model's own window is a
        # hard edge rather than a threshold somebody chose.
        assert block["shape_margin"] is None

    def test_the_shape_verdict_does_not_erase_the_branch_inputs(
        self, db_session, evidence, workbook_bytes
    ):
        """6.2: the measured values a shape decision consumes are what a calibration sweep
        replays against, and this task adds a verdict rather than forgetting the
        evidence."""
        file_node = _ingest(db_session, workbook_bytes)
        node = _sheet_node(db_session, file_node, "Deals")
        decision = evidence(node)["sheet"]["shape_decision"]

        assert decision["decided"] is True
        assert decision["basis"] == "header_row, rendered_fits_doc_window"
        assert "8.4" in decision["reason"]
        assert decision["inputs"]["header_row"] is True
        assert decision["inputs"]["header_row_number"] == 1
        assert "fill_ratio" in decision["inputs"]

    def test_a_sheet_with_no_header_row_gets_no_records_and_says_why(
        self, db_session, evidence, workbook_bytes
    ):
        """Not a failure. A sheet with no header row has no keys, and the two shapes 8.4
        gives it — `matrix` and `unstructured` — read thresholds 8.8 leaves unset.

        It is NOT nothing any more, and that is step 12: `Notes` renders inside the
        document window, so it gets the `small_table` node it always matched and nothing
        ever wrote. On the open web that population is 56.0% of sheets, against the 4.6%
        that match `records` and not `table`.
        """
        file_node = _ingest(db_session, workbook_bytes)
        node = _sheet_node(db_session, file_node, "Notes")

        assert _children(db_session, node.id, USETYPE_RECORD) == []
        assert evidence(node)["sheet"]["shape"] == [SHAPE_SMALL_TABLE]
        attempt = next(
            entry for entry in evidence(node)["attempts"] if entry["task"] == TASK_EXTRACT_SHEET
        )
        assert attempt["rung"] == RUNG_INFERRED
        assert "header_row" in attempt["detail"]["no_records"]


class TestTheTableShape:
    """Step 11 and step 12. 8.4's ``small_table``, and both shapes where both match."""

    def test_the_whole_sheet_arrives_as_one_markdown_table(
        self, db_session, evidence, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")
        tables = _children(db_session, sheet.id, USETYPE_TABLE)

        assert len(tables) == 1
        lines = tables[0].content.splitlines()
        assert lines[0] == "| Deal ID | Account | Value | Close Date | Won |"
        assert lines[1] == "| --- | --- | --- | --- | --- |"
        assert tables[0].title == "Deals — whole sheet"

    def test_the_markdown_is_the_profile_s_own_and_not_a_second_rendering(
        self, db_session, evidence, workbook_bytes
    ):
        """The same rule this task already follows for the header verdict: a second
        derivation could disagree, and then the token count that decided the shape would be
        a count of a string nobody stored."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")
        table = _children(db_session, sheet.id, USETYPE_TABLE)[0]

        assert table.content == evidence(sheet)["sheet"]["measurements"]["rendered_markdown"]

    def test_the_table_node_carries_what_it_is_the_whole_of(
        self, db_session, evidence, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")
        block = evidence(_children(db_session, sheet.id, USETYPE_TABLE)[0])["table"]

        assert block["sheet_name"] == "Deals"
        assert (block["rows"], block["cols"]) == (5, 5)
        assert block["header_row_number"] == 1
        assert block["fits_doc_window"] is True
        assert block["rendered_tokens"] > 0

    def test_both_shapes_are_emitted_and_neither_contains_the_other(
        self, db_session, evidence, workbook_bytes
    ):
        """8.4 says the first match wins. Both match here and both are written: a table
        node and its own row nodes are separate documents, and a query matching both
        returns both."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        tables = _children(db_session, sheet.id, USETYPE_TABLE)
        records = _children(db_session, sheet.id, USETYPE_RECORD)
        assert len(tables) == 1 and len(records) == 3
        assert all(table.parent_id == sheet.id for table in tables)
        assert all(record.parent_id == sheet.id for record in records)

    def test_the_table_node_gets_a_vector_like_any_other_leaf(self, db_session, workbook_bytes):
        """8.4: "embedded as-is". A shape that produced a node no query could reach would
        be `KNOWN-DEFECTS` D1 in a third costume."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")
        table = _children(db_session, sheet.id, USETYPE_TABLE)[0]

        assert table.embed is not None
        assert table.settled == SETTLED_SETTLED

    def test_a_sheet_that_matches_neither_shape_writes_nothing_and_claims_no_rung(
        self, db_session, evidence, workbook_bytes
    ):
        """The branch that used to be the ONLY outcome for a header-less sheet, and is now
        the narrow case: no header row and no render inside the document window.

        Driven through the handler rather than through a fixture whose render genuinely
        misses the window, because such a workbook is 16,000 cells and 400 record nodes to
        embed — the branch under test is a predicate over two stored measurements, and a
        four-hundred-node ingest would be paying for the predicate in forward passes.
        `Notes` has no header under either pass; stripping the stored markdown is what the
        other half of the condition looks like on the node.
        """
        file_node = _ingest(db_session, workbook_bytes)
        node = _sheet_node(db_session, file_node, "Notes")
        block = dict(evidence(node)["sheet"])
        block["measurements"] = dict(
            block["measurements"], rendered_markdown=None, rendered_fits_doc_window=False
        )
        EvidenceRepository(db_session).write(node.id, "sheet", block)
        db_session.flush()

        outcome = run_extract_sheet(db_session, _ScopedTask(node.id))

        assert outcome.rung is None
        assert outcome.detail["shape"] == []
        assert "8.4's `small_table` does not match it" in outcome.detail["no_table"]
        assert outcome.produced["node_count"] == 0
        assert evidence(node)["sheet"]["shape"] == []
        assert evidence(node)["sheet"]["shape_decision"]["decided"] is False

    def test_a_sheet_the_render_missed_says_which_way_it_missed(self, withheld_block):
        """The two ways a sheet misses `small_table` are different facts and are two keys:
        it did not render at all (`rendered_unbounded_reason`), or it rendered over the
        window and the markdown was not kept (`rendered_withheld_reason`). Neither is a
        failure, and a null with no reason beside it would be neither answer."""
        measurements = withheld_block["measurements"]

        assert measurements["rendered_markdown"] is None
        assert measurements["rendered_unbounded_reason"] is None
        assert "over the 8192-token document window" in measurements["rendered_withheld_reason"]

    @pytest.fixture
    def withheld_block(self, workbook_bytes):
        """8.7's block for a sheet whose render is over the document window.

        A pure function of the measurement and a token count, so the over-window case costs
        one call rather than a workbook nobody wants to embed.
        """
        return sheet_evidence_block(
            measure_sheet(workbook_bytes, "Deals", render_cell_budget=8192),
            rendered_tokens=9_000,
            token_window=512,
            doc_window=8192,
            embedding_model="test-model",
            render_cell_budget=8192,
        )

    def test_the_attempt_counts_what_the_standard_library_pass_found(
        self, db_session, evidence, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        node = _sheet_node(db_session, file_node, "Deals")
        detail = next(
            entry for entry in evidence(node)["attempts"] if entry["task"] == TASK_EXTRACT_SHEET
        )["detail"]

        assert detail["cell_notes_read"] is True
        assert detail["cells_with_formula"] == 1
        assert detail["cells_text_forced"] == 1
        assert detail["records_over_token_window"] == 0

    def test_a_row_bound_below_the_sheet_fails_the_task(self, db_session, workbook_bytes):
        """6.6, through the front door: the handler does not catch it and turn it into a
        shorter sheet."""
        file_node = _ingest(db_session, workbook_bytes)
        node = _sheet_node(db_session, file_node, "Deals")
        with pytest.raises(TooManyRows):
            run_extract_sheet(db_session, _ScopedTask(node.id, {"max_rows": 2}))

    def test_it_refuses_a_node_that_is_not_a_sheet(self, db_session, workbook_bytes):
        file_node = _ingest(db_session, workbook_bytes)
        with pytest.raises(ValueError, match="materialises one worksheet's cells"):
            run_extract_sheet(db_session, _ScopedTask(file_node.id))

    def test_it_refuses_a_sheet_the_profile_has_not_measured(
        self, db_session, evidence, workbook_bytes
    ):
        """It reads the header verdict and the column names off the profile rather than
        deriving them a second time, and is ordered after it for that reason."""
        file_node = _ingest(db_session, workbook_bytes)
        node = _sheet_node(db_session, file_node, "Deals")
        block = dict(evidence(node)["sheet"])
        block.pop("measurements")
        EvidenceRepository(db_session).write(node.id, "sheet", block)
        db_session.flush()

        with pytest.raises(ValueError, match="sheet.measurements"):
            run_extract_sheet(db_session, _ScopedTask(node.id))
