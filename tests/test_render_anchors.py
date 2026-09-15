"""``anchor:render`` — the rectangle a passage occupies on a page of the RENDITION.

``docs/OFFICE_SPEC.md`` Part 5's "Two anchors, not one" and Part 11 step 8;
``docs/SPRINT_0_6_0.md`` Block G step 33. Step 32 gave an office document a PDF
(``render:pdf``) and step 34 gave it an exact address into its own package
(``anchor:source``). This is the third anchor and the only approximate one: a page and a
rectangle, recovered by matching the passage's own text against the PDF LibreOffice made.

**WHAT MUST NOT HAPPEN IS THE THING MOST OF THIS FILE IS ABOUT.** The recovered rectangle
goes in a row of its OWN — ``render_anchor``, not ``source_anchor`` — and
:class:`TestTheTwoAnchorsStayApart` is the test that says so from both sides: running this
task leaves the ``ooxml`` address untouched, and the rectangle it writes says
``"of": "rendition"`` so that a reader holding a parsed anchor can tell LibreOffice's
pagination from a PDF's own text layer. Part 5: "Storing both under one key would make a
recovered rectangle indistinguishable from an exact one."

**NO LIBREOFFICE RAN IN ANY TEST HERE, AND THAT IS STATED RATHER THAN HIDDEN.** Tier 3 is a
500 MB system package behind ``JMFTS_CONVERT_URL``, which is blank by default and blank in
CI. What these tests install as the rendition is a PDF built by ``pymupdf`` — a base
dependency — holding the fixture package's own paragraph strings, laid out across two pages
with a running head, one paragraph deliberately omitted and one heading deliberately
repeated. It is stored through the real :func:`run_render_pdf` with the converter stubbed,
which is ``test_render_tasks.py``'s own fixture, so everything downstream of the conversion
is the shipped code path: a real ``document_renditions`` row, a real ``rendition`` evidence
block, a real ``pdf_for``, a real ``pdf_to_markdown``.

What that therefore does NOT prove is the one number Part 10 exists to measure: the RATE at
which real LibreOffice output aligns with the office readers' markdown on real documents.
This fixture's rendered text is the package's text by construction, so it can prove the
alignment, the anchors, the row separation and every refusal — and it cannot prove that a
``.docx`` full of tables, text boxes and hyphenation comes back with a useful recovery rate.
The handler reports ``located_blocks`` against ``source_blocks`` on every attempt so that
number is readable the day a real converter is in front of it. A fixture that had quietly
passed with no rendition present at all would have been worse than no test; the schedule
tests below are what rule that out.
"""

from __future__ import annotations

import io

import pytest

pytest.importorskip("docx", reason="the office extra is not installed")
pymupdf = pytest.importorskip("pymupdf")

import jmfts_core.ingest_tasks  # noqa: E402,F401  (import order; see the module cycle)
from docx import Document as DocxDocument  # noqa: E402
from sqlalchemy import select  # noqa: E402

from jmfts_client.contracts.anchor import (  # noqa: E402
    RENDER_ANCHOR_ROW,
    RENDER_ANCHOR_UNRESOLVED_ROW,
    OoxmlAnchor,
    PdfAnchor,
    parse_anchor,
)
from jmfts_client.contracts.upload import UploadedFile  # noqa: E402
from jmfts_core.atoms import ATOMS, EV_RENDER_ANCHOR, EV_RENDITION  # noqa: E402
from jmfts_core.citation_tasks import (  # noqa: E402
    ANCHOR_KIND_OOXML,
    ANCHOR_KIND_PDF,
    ANCHOR_NAME,
    ANCHOR_OF_RENDITION,
    ANCHOR_OF_SOURCE,
    MIN_PREFIX_CHARACTERS,
    RENDER_ANCHOR_NAME,
    RENDER_ANCHOR_UNRESOLVED_NAME,
    UNRESOLVED_NO_BLOCK,
    UNRESOLVED_NO_SPAN,
    UNRESOLVED_NOT_IN_RENDITION,
    align_to_rendition,
    anchor_for_rendition_span,
    normalize_for_match,
    run_anchor_render,
)
from jmfts_core.ingest_tasks import (  # noqa: E402
    TASK_ANCHOR_RENDER,
    TASK_ANCHOR_SOURCE,
    TASK_RENDER_PDF,
    TASK_ROWS,
    explain_plan,
)
from jmfts_core.models.document import (  # noqa: E402
    SETTLED_SETTLED,
    USETYPE_CHUNK,
    Document,
)
from jmfts_core.models.task_queue import (  # noqa: E402
    ADVISORY_TASK_TYPES,
    WRITE_SELF,
    WRITE_SUBTREE,
)
from jmfts_core.probe import detect_format, probe_patterns  # noqa: E402
from jmfts_core.render_tasks import run_render_pdf  # noqa: E402
from jmfts_core.repositories.document import DocumentRepository  # noqa: E402
from jmfts_core.repositories.evidence import EvidenceRepository  # noqa: E402
from jmfts_core.repositories.rendition import KIND_PDF, RenditionRepository  # noqa: E402
from jmfts_core.repositories.task_queue import TaskQueueRepository  # noqa: E402
from jmfts_core.rollup_tasks import IngestRollupPlanner  # noqa: E402
from jmfts_core.services.document_service import (  # noqa: E402
    DocumentService,
    RegionNotAddressable,
)
from jmfts_core.services.ingest_service import IngestService  # noqa: E402
from tests.conftest import drain_ingest_queue  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# ---------------------------------------------------------------------------
# The fixture package, and the fixture rendition of it
# ---------------------------------------------------------------------------
#
# The strings are shared by the two builders on purpose: what the alignment matches is the
# document's own text, so the rendition has to hold the document's own text. What makes it a
# TEST rather than a tautology is the three things the rendition does differently, each of
# which is something a real renderer does:
#
#   * a RUNNING HEAD on every page, which no source block produced;
#   * ONE PARAGRAPH MISSING, which is a block the alignment must report rather than guess at;
#   * ONE HEADING REPEATED, which is the case Part 5 lists as a search failure ("the text
#     appears more than once") and which a monotone walk resolves correctly.

RUNNING_HEAD = "Draft - internal"
HEADING_ONE = "Fusion tree search"
PROSE_ONE = (
    "MaxSim scores a document as the sum of per-query-token maxima against the document's "
    "own token vectors, which makes it a reranker rather than a retriever."
)
HEADING_TWO = "Retrieval"
PROSE_TWO = "Vector recall is high and precision is not, which is why BM25 supplies the precision."
PROSE_THREE = "Retention at fifty percent of tokens preserves the ranking to within a point."
PROSE_UNRENDERED = (
    "This paragraph is in the package and the renderer never put it on any page at all."
)
PROSE_FOUR = (
    "The second occurrence of that heading is what makes an isolated text search ambiguous."
)

#: What the rendition holds, per page. ``PROSE_UNRENDERED`` is absent from both.
RENDITION_PAGES = [
    [RUNNING_HEAD, HEADING_ONE, PROSE_ONE, HEADING_TWO, PROSE_TWO],
    [RUNNING_HEAD, PROSE_THREE, HEADING_TWO, PROSE_FOUR],
]


def _anchored_docx() -> bytes:
    document = DocxDocument()
    document.add_heading(HEADING_ONE, level=1)
    document.add_paragraph(PROSE_ONE)
    document.add_heading(HEADING_TWO, level=2)
    document.add_paragraph(PROSE_TWO)
    document.add_paragraph(PROSE_THREE)
    document.add_paragraph(PROSE_UNRENDERED)
    document.add_heading(HEADING_TWO, level=2)
    document.add_paragraph(PROSE_FOUR)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _rendition_pdf() -> bytes:
    """A two-page PDF standing in for what LibreOffice would have produced.

    One text box per block, spaced far enough apart that PyMuPDF's own block analysis
    separates them — which is the granularity the handler aligns against, so a fixture that
    merged two paragraphs into one block would be testing a different thing than the one
    under test.
    """
    doc = pymupdf.open()
    for blocks in RENDITION_PAGES:
        page = doc.new_page()
        top = 60.0
        for text in blocks:
            page.insert_textbox(pymupdf.Rect(72, top, 523, top + 60), text, fontsize=11)
            top += 90
    return doc.tobytes()


@pytest.fixture(scope="module")
def rendition_pdf() -> bytes:
    return _rendition_pdf()


def _upload(session, data: bytes, filename: str = "anchors.docx"):
    return IngestService(session).upload_file(
        UploadedFile(data=data, filename=filename, content_type=DOCX_MIME)
    )


def _chunks(session, root_id: int) -> list[Document]:
    return list(
        session.execute(
            select(Document)
            .where(Document.path.contains([root_id]))
            .where(Document.usetype == USETYPE_CHUNK)
            .order_by(Document.id)
        )
        .scalars()
        .all()
    )


def _run_task(session, task_type, node_id, write_mode, handler, params=None):
    """Enqueue, claim, run and COMPLETE one task the way ``IngestWorker.run_once`` does.

    The completion is not ceremony, and neither are the ``params``. ``complete`` is what
    appends the attempt record (spec 5.6), the attempt log is the authority for 6.1's re-run
    diff, and the diff's key is ``(task, param_fingerprint)`` — so a helper that called the
    handler, left the queue row pending, or queued it with the wrong parameters would test
    the handler and quietly disable the diff. That is what
    :class:`TestTheSettlingBoundaryOffersIt` exists to notice.
    """
    tasks = TaskQueueRepository(session)
    task = tasks.enqueue(task_type, node_id, write_mode, params=params, service_badge=None)
    claimed = tasks.claim_next(f"test-{task_type}")
    assert claimed is not None and claimed.id == task.id, (
        f"claim_next took {claimed and claimed.task_type!r} rather than {task_type!r}; "
        "something else is pending on this node"
    )
    tasks.mark_running(claimed)
    outcome = handler(session, claimed)
    tasks.complete(
        claimed,
        detail=outcome.detail,
        produced=outcome.produced,
        rung=outcome.rung,
        status=outcome.status,
    )
    session.flush()
    return outcome


def _install_rendition(session, node, rendition_pdf, monkeypatch):
    """Give ``node`` a rendition through the real ``render:pdf``, with the converter stubbed.

    ``test_render_tasks.py``'s fixture, reused rather than reinvented: everything after
    ``convert_bytes`` is the shipped path, so the row, the evidence block, the page geometry
    and the hash are all the ones a real conversion would have written.
    """
    monkeypatch.setattr("jmfts_core.render_tasks.convert_bytes", lambda data, *, to: rendition_pdf)
    _run_task(session, TASK_RENDER_PDF, node.id, WRITE_SELF, run_render_pdf)


def _run_anchor_render(session, node):
    """Run the task the way the settling walk would, PARAMS AND ALL.

    The spec comes from the planner rather than being written out here, which is what makes
    the ``anchored`` fixture and :class:`TestTheSettlingBoundaryOffersIt` talk about the same
    object: the fingerprint the diff will look for later is the one this run recorded.
    """
    specs = [s for s in IngestRollupPlanner()(session, node) if s.task_type == TASK_ANCHOR_RENDER]
    assert specs, "the planner did not offer anchor:render for a node that has a rendition"
    spec = specs[0]
    return _run_task(
        session, spec.task_type, node.id, spec.write_mode, run_anchor_render, params=spec.params
    )


def _first_words(text: str, count: int = 5) -> str:
    """A needle short enough that no line wrap can split it.

    ``search_for`` matches glyphs laid out on a page, and a whole sentence in a text box
    routinely wraps. What this assertion is for — one paragraph's rectangle against the whole
    page's — is tens of points wide, and five words are enough to name a paragraph uniquely
    in a fixture of eight.
    """
    return " ".join(text.split()[:count])


@pytest.fixture
def anchored(db_session, rendition_pdf, monkeypatch):
    """An ingested ``.docx`` with a rendition and both anchor tasks run over it."""
    response = _upload(db_session, _anchored_docx())
    drain_ingest_queue(db_session, max_tasks=200)
    node = DocumentRepository(db_session).get(response.document_id)
    assert node.settled == SETTLED_SETTLED
    _install_rendition(db_session, node, rendition_pdf, monkeypatch)
    outcome = _run_anchor_render(db_session, node)
    return node, outcome


# ---------------------------------------------------------------------------
# 1. The declaration — a task, an atom, and deliberately not a row
# ---------------------------------------------------------------------------


class TestDeclaration:
    def test_anchor_render_is_declared_advisory(self):
        """``models.task_queue`` cannot import ``ingest_tasks``, so the name is spelled
        twice. A rename touching only one of them would turn a missing rectangle back into a
        failed document — which for this task is the least defensible failure of the three
        anchors, because the passage still has its exact address."""
        assert TASK_ANCHOR_RENDER in ADVISORY_TASK_TYPES

    def test_it_is_not_a_row_of_the_part_four_table(self):
        """The decision the whole step turns on. Its precondition is "the chunks exist AND a
        rendition exists", and the second half written as ``after render:pdf`` would gate the
        claim on a COMPLETED row of an ADVISORY task — which on an appliance with no
        LibreOffice fails permanently, leaving this task unclaimable, permanently pending,
        and therefore permanently unfinished. The file node would never settle."""
        assert TASK_ANCHOR_RENDER not in {row.task for row in TASK_ROWS}

    def test_the_atom_declares_the_rendition_it_cannot_run_without(self):
        """The one fact that distinguishes it from ``anchor:source``, whose ``consumes`` is
        otherwise identical. It is also why they are two atoms: one atom over both formats
        would declare a dependency on a rendition that a source PDF never has."""
        atom = ATOMS[TASK_ANCHOR_RENDER]
        assert EV_RENDITION in {fact.name for fact in atom.consumes}
        assert {fact.name for fact in atom.produces} == {EV_RENDER_ANCHOR}
        assert atom.write_mode == WRITE_SUBTREE

    def test_the_source_anchor_row_is_still_scheduled_for_the_same_document(self):
        """The two are not alternatives. A ``.docx`` gets its exact ``ooxml`` address from
        the Part 4 table as it always did; the rectangle arrives later and separately."""
        data = _anchored_docx()
        detection = detect_format(data, filename="a.docx", declared_mime=None)
        patterns, _ = probe_patterns(data, detection)
        plan = explain_plan(fmt=detection.format, patterns=patterns, patterns_source="probed")
        rows = {task.task: task for task in plan.tasks}
        assert rows[TASK_ANCHOR_SOURCE].outcome == "enqueued"
        assert rows[TASK_RENDER_PDF].outcome == "enqueued"


# ---------------------------------------------------------------------------
# 2. The matcher — what the two sides are expected to agree about
# ---------------------------------------------------------------------------


class TestNormalizeForMatch:
    def test_markdown_markers_come_off_both_sides(self):
        """Both sides are markdown and neither side's markers came from the other: a
        heading the package declares by STYLE becomes ``#`` here, and the same heading in
        the rendition becomes ``#`` only if PyMuPDF judged its font size to be one."""
        assert normalize_for_match("## Retrieval") == normalize_for_match("Retrieval")
        assert normalize_for_match("**bold** text") == normalize_for_match("bold text")

    def test_a_line_break_is_the_renderers_decision_and_not_the_authors(self):
        assert normalize_for_match("one two\nthree") == normalize_for_match("one   two three")

    def test_case_folds_because_a_renderer_may_have_capitalised_it(self):
        """The one lossy step, and it earns its place: a heading styled all-caps renders as
        capitals while the package holds mixed case, so a case-sensitive test would refuse to
        locate exactly the headings a reader most wants located."""
        assert normalize_for_match("RETRIEVAL") == normalize_for_match("Retrieval")


class TestAlignToRendition:
    def test_a_repeated_block_resolves_to_the_copy_in_document_order(self):
        """Part 5's first named failure mode — "the text appears more than once" — answered
        rather than reported. It is ambiguous for a search over one string and unambiguous
        for a walk over a sequence."""
        source = ["alpha", "the heading", "beta", "the heading", "gamma"]
        rendered = ["alpha", "the heading", "beta", "running head", "the heading", "gamma"]
        assert align_to_rendition(source, rendered) == [0, 1, 2, 4, 5]

    def test_page_furniture_no_source_block_produced_is_walked_past(self):
        source = ["alpha", "beta"]
        rendered = ["running head", "alpha", "running head", "beta"]
        assert align_to_rendition(source, rendered) == [1, 3]

    def test_a_block_the_rendition_does_not_hold_is_none_and_costs_nothing_after_it(self):
        """The cursor only moves on a match, so a block that is missing does not desynchronise
        the blocks below it."""
        source = ["alpha", "missing entirely", "beta"]
        rendered = ["alpha", "beta"]
        assert align_to_rendition(source, rendered) == [0, None, 1]

    def test_a_paragraph_split_across_a_page_break_matches_its_first_fragment(self):
        """Part 5's third failure mode. The first fragment is exactly the rectangle an anchor
        wants: the reader who clicks a citation is taken to where the passage starts."""
        whole = "a sentence long enough to be split across a page break by the renderer"
        assert len(whole[:40]) >= MIN_PREFIX_CHARACTERS
        assert align_to_rendition([whole], [whole[:40]]) == [0]

    def test_a_short_block_is_not_matched_by_prefix(self):
        """``yes`` inside ``yes, and the following quarter`` is a coincidence, not a location.
        Exactness is still allowed at any length — it is only the prefix arm that has a
        floor."""
        assert align_to_rendition(["yes"], ["yes, and the following quarter as well"]) == [None]
        assert align_to_rendition(["yes"], ["yes"]) == [0]

    def test_it_never_looks_backwards(self):
        source = ["beta", "alpha"]
        rendered = ["alpha", "beta"]
        assert align_to_rendition(source, rendered) == [1, None]


# ---------------------------------------------------------------------------
# 3. Inverting a span — the shape, and the three absences
# ---------------------------------------------------------------------------


def _blocks(*spec) -> tuple[list[dict], list[int]]:
    """``(blocks, starts)`` from ``(start, end, page, bbox)``; ``page=None`` for a block the
    alignment did not locate."""
    blocks = [
        {"start": start, "end": end, **({"page": page, "bbox": bbox} if page is not None else {})}
        for start, end, page, bbox in spec
    ]
    return blocks, [block["start"] for block in blocks]


class TestAnchorForRenditionSpan:
    def test_a_span_inside_one_located_block_takes_that_blocks_rectangle(self):
        blocks, starts = _blocks((0, 20, 1, [10, 20, 30, 40]))
        anchor, code = anchor_for_rendition_span(blocks, starts, 4, 12)
        assert code is None
        assert anchor == {
            "kind": ANCHOR_KIND_PDF,
            "of": ANCHOR_OF_RENDITION,
            "page": 1,
            "bbox": [10, 20, 30, 40],
        }

    def test_the_rectangle_is_the_union_on_the_first_page_only(self):
        """``anchor_for_span``'s rule, kept word for word so that a consumer which can draw
        one can draw the other: the page is where the passage BEGINS, and a rectangle that
        looked complete when it was not would be the real failure."""
        blocks, starts = _blocks(
            (0, 10, 0, [10, 20, 30, 40]),
            (12, 20, 0, [5, 50, 80, 90]),
            (22, 30, 1, [1, 1, 99, 99]),
        )
        anchor, code = anchor_for_rendition_span(blocks, starts, 0, 30)
        assert code is None
        assert anchor["page"] == 0
        assert anchor["bbox"] == [5, 20, 80, 90]
        assert anchor["continues"] == [1]

    def test_a_span_whose_blocks_are_all_unlocated_says_so(self):
        """The blocks are in the source and the rendition does not hold them. Reaching past
        them to the nearest located block is the plausible-and-wrong answer Part 5 rejects."""
        blocks, starts = _blocks((0, 10, None, None), (12, 20, 0, [1, 2, 3, 4]))
        anchor, code = anchor_for_rendition_span(blocks, starts, 0, 10)
        assert anchor is None
        assert code == UNRESOLVED_NOT_IN_RENDITION

    def test_a_span_touching_no_block_at_all_is_a_different_absence(self):
        """ "Nothing in the source accounts for this chunk" and "the rendition does not hold
        what does" are different facts, and a consumer told only the second would look for a
        conversion problem that is not there."""
        blocks, starts = _blocks((0, 10, 0, [1, 2, 3, 4]))
        anchor, code = anchor_for_rendition_span(blocks, starts, 40, 50)
        assert anchor is None
        assert code == UNRESOLVED_NO_BLOCK


# ---------------------------------------------------------------------------
# 4. Over a real package and a real stored rendition
# ---------------------------------------------------------------------------


class TestOverARealRendition:
    def test_every_anchored_chunk_s_rectangle_holds_its_own_words(
        self, db_session, evidence, anchored, rendition_pdf
    ):
        """The check a plausible-but-wrong rectangle cannot pass, and the office half of
        ``test_citation_task``'s: the chunk's own words are looked for on the page the anchor
        names, in the RENDITION, and the rectangle they are found in has to be inside the one
        that was recorded."""
        node, _ = anchored
        chunks = _chunks(db_session, node.id)
        assert chunks, "the structure rung wrote no chunks"

        placed = 0
        source = pymupdf.open(stream=rendition_pdf, filetype="pdf")
        try:
            for chunk in chunks:
                anchor = evidence(chunk).get(RENDER_ANCHOR_NAME)
                if anchor is None:
                    continue
                placed += 1
                assert anchor["kind"] == ANCHOR_KIND_PDF
                assert anchor["of"] == ANCHOR_OF_RENDITION

                page = source[anchor["page"]]
                needle = _first_words(chunk.content.lstrip("# "))
                hits = page.search_for(needle)
                assert hits, f"{needle!r} is not on rendition page {anchor['page']}"
                recorded = pymupdf.Rect(anchor["bbox"]) + (-0.5, -0.5, 0.5, 0.5)
                assert recorded.contains(hits[0]), (
                    f"chunk {chunk.id}'s rectangle {anchor['bbox']} does not contain the "
                    f"rectangle {list(hits[0])} its own words occupy"
                )
        finally:
            source.close()
        assert placed, "no chunk of the document was given a rectangle"

    def test_the_repeated_heading_does_not_send_a_later_passage_to_an_earlier_page(
        self, db_session, evidence, anchored
    ):
        """The fixture's point, end to end. ``Retrieval`` heads two sections of the package
        and appears twice in the rendition, on different pages. If the alignment had matched
        the SECOND occurrence to the page-0 copy, its cursor would be behind and the passage
        under it would land on page 0; if it had matched the FIRST to the page-1 copy, the
        cursor would be ahead and the passage under THAT one could not be located at all. So
        the two passages' pages are what say the walk stayed in document order, and a
        per-string search — Part 5's proposal — has no way to get both right."""
        node, _ = anchored
        pages = {}
        for chunk in _chunks(db_session, node.id):
            anchor = evidence(chunk).get(RENDER_ANCHOR_NAME)
            if anchor is None:
                continue
            for name, text in (("second", PROSE_TWO), ("fourth", PROSE_FOUR)):
                if chunk.content.startswith(text):
                    pages[name] = anchor
        assert set(pages) == {"second", "fourth"}, f"expected both passages anchored, got {pages}"
        # Under the FIRST `Retrieval`, and it runs on into the next page — the chunk covers
        # `PROSE_THREE` as well, which the rendition puts overleaf.
        assert pages["second"]["page"] == 0
        assert pages["second"]["continues"] == [1]
        # Under the SECOND `Retrieval`, which is on page 1 and nowhere else.
        assert pages["fourth"]["page"] == 1
        assert "continues" not in pages["fourth"]

    def test_the_attempt_detail_carries_the_two_recovery_rates_and_the_renderer(
        self, db_session, anchored
    ):
        """Two ratios rather than one, because they fail in different places.
        ``located_blocks/source_blocks`` is how much of the document the ALIGNMENT found;
        ``anchored/chunks`` is how much of it a reader can click. And ``renderer_version`` is
        what says which pagination these coordinates belong to — Part 6: two LibreOffice
        versions paginate differently."""
        node, outcome = anchored
        detail = outcome.detail
        assert detail["chunks"] == len(_chunks(db_session, node.id))
        assert detail["anchored"] + sum(detail["unresolved"].values()) == detail["chunks"]
        assert detail["rendition_pages"] == len(RENDITION_PAGES)
        # The one paragraph the fixture left out of the rendition. If this were equal, the
        # alignment would be claiming to have found text that is not there.
        assert detail["located_blocks"] < detail["source_blocks"]
        assert "renderer_version" in detail

    def test_bytes_with_no_description_beside_them_are_refused(
        self, db_session, anchored, evidence
    ):
        """``render:pdf`` writes the row and the bytes in one transaction and
        ``RenditionRepository.delete`` removes them in one, so bytes with no ``rendition``
        row is a state neither half can produce. Anchoring through it would mean recovering
        rectangles against a pagination nothing on the node describes."""
        node, _ = anchored
        EvidenceRepository(db_session).delete(node.id, "rendition")
        db_session.flush()

        with pytest.raises(ValueError, match="evidence row describing them"):
            _run_task(db_session, TASK_ANCHOR_RENDER, node.id, WRITE_SUBTREE, run_anchor_render)

    def test_a_page_count_that_disagrees_with_the_bytes_is_refused(self, db_session, anchored):
        """The other direction. ``page_count`` is what a caller scales a rectangle against, so
        a row that disagrees with the object it describes would put boxes on pages nobody
        sees."""
        node, _ = anchored
        repository = EvidenceRepository(db_session)
        record = dict(repository.read(node.id, "rendition"))
        record["page_count"] = 99
        repository.write(node.id, "rendition", record)
        db_session.flush()

        with pytest.raises(ValueError, match="come apart"):
            _run_task(db_session, TASK_ANCHOR_RENDER, node.id, WRITE_SUBTREE, run_anchor_render)

    def test_a_source_pdf_is_refused_by_name_rather_than_re_recovered(self, db_session):
        """For a document whose own bytes are the PDF, Part 5's two anchors are the same
        object. ``citation`` retains the geometry extraction already computed; matching the
        text against itself would write a second row saying the same thing less exactly."""
        from tests.corpus.fixtures import minimal_pdf

        response = IngestService(db_session).upload_file(
            UploadedFile(data=minimal_pdf(), filename="a.pdf", content_type="application/pdf")
        )
        drain_ingest_queue(db_session, max_tasks=200)
        node = DocumentRepository(db_session).get(response.document_id)

        with pytest.raises(ValueError, match="citation"):
            _run_task(db_session, TASK_ANCHOR_RENDER, node.id, WRITE_SUBTREE, run_anchor_render)

    def test_a_pdf_resolved_from_the_blob_is_refused_even_when_the_reader_agreed(
        self, db_session, anchored
    ):
        """The same refusal reached through the OTHER door, which is the one that needed a
        guard. ``pdf_for`` forks on ``file.detected_mime`` and the reader table forks on
        ``extraction.source``, so a node whose ``file`` block says PDF and whose extraction
        says ``docx`` passes the first check and resolves to its own uploaded package. Handing
        those bytes to the alignment as "the rendition" would recover rectangles from a PDF
        that is not one."""
        node, _ = anchored
        repository = EvidenceRepository(db_session)
        block = dict(repository.read(node.id, "file"))
        block["detected_mime"] = "application/pdf"
        repository.write(node.id, "file", block)
        db_session.flush()

        with pytest.raises(ValueError, match="origin 'source'"):
            _run_task(db_session, TASK_ANCHOR_RENDER, node.id, WRITE_SUBTREE, run_anchor_render)

    def test_a_chunk_with_no_source_span_is_reported_rather_than_failed(
        self, db_session, evidence, anchored
    ):
        node, _ = anchored
        chunks = _chunks(db_session, node.id)
        EvidenceRepository(db_session).delete(chunks[0].id, "source_span")
        db_session.flush()

        # Enqueued directly rather than through the planner: this is a RE-RUN, and the
        # planner has already recorded this rendition's fingerprint. What is under test is
        # the handler's answer for a chunk it cannot invert, not the schedule.
        outcome = _run_task(
            db_session, TASK_ANCHOR_RENDER, node.id, WRITE_SUBTREE, run_anchor_render
        )
        assert outcome.detail["unresolved"][UNRESOLVED_NO_SPAN] == 1
        found = evidence(chunks[0])
        assert RENDER_ANCHOR_NAME not in found
        assert found[RENDER_ANCHOR_UNRESOLVED_NAME]["code"] == UNRESOLVED_NO_SPAN


# ---------------------------------------------------------------------------
# 5. Two anchors, two rows — the failure this step is most able to cause
# ---------------------------------------------------------------------------


class TestTheTwoAnchorsStayApart:
    def test_the_exact_address_survives_the_recovered_rectangle(
        self, db_session, evidence, anchored
    ):
        """Both tasks are ``subtree``-scoped over the same chunks and both write an anchor.
        One row name would have made this task the sole survivor, silently — and the thing
        lost would be the EXACT address, in favour of the approximate one."""
        node, _ = anchored
        for chunk in _chunks(db_session, node.id):
            found = evidence(chunk)
            source = found.get(ANCHOR_NAME)
            assert source is not None, "anchor:source's address is gone"
            assert source["kind"] == ANCHOR_KIND_OOXML
            assert source["part"] == "word/document.xml"

    def test_a_reader_can_tell_a_recovered_rectangle_from_an_exact_one(
        self, db_session, evidence, anchored
    ):
        """The contract half. Both rows parse with the same ``parse_anchor``, and what keeps
        them apart once parsed is ``of`` — which is why the field exists and why it defaults
        to ``source`` rather than being required."""
        node, _ = anchored
        rectangles = 0
        for chunk in _chunks(db_session, node.id):
            found = evidence(chunk)
            assert isinstance(parse_anchor(found[ANCHOR_NAME]), OoxmlAnchor)
            row = found.get(RENDER_ANCHOR_NAME)
            if row is None:
                continue
            rectangles += 1
            parsed = parse_anchor(row)
            assert isinstance(parsed, PdfAnchor)
            assert parsed.of == ANCHOR_OF_RENDITION
        assert rectangles

    def test_an_anchor_written_before_this_field_existed_still_reads_as_a_source_anchor(self):
        """``of`` defaults rather than being required, and this is the row that default is
        about: every ``source_anchor`` in every existing database was written by ``citation``
        from a PDF's own text layer. Requiring the field would have changed what all of them
        mean."""
        old = {"kind": "pdf", "page": 3, "bbox": [72.0, 118.4, 540.0, 262.9]}
        assert parse_anchor(old).of == ANCHOR_OF_SOURCE

    def test_the_two_rows_are_registered_under_names_of_their_own(self):
        """The spellings the handler writes and the spellings a client reads, pinned
        together — they live in two distributions and cannot import each other."""
        assert RENDER_ANCHOR_NAME == RENDER_ANCHOR_ROW
        assert RENDER_ANCHOR_UNRESOLVED_NAME == RENDER_ANCHOR_UNRESOLVED_ROW
        assert RENDER_ANCHOR_NAME != ANCHOR_NAME


# ---------------------------------------------------------------------------
# 6. The schedule — offered when there is a rendition, and never otherwise
# ---------------------------------------------------------------------------


class TestTheSettlingBoundaryOffersIt:
    def test_a_document_with_no_rendition_is_offered_nothing(self, db_session):
        """THE TEST THAT MAKES EVERY TEST ABOVE MEAN SOMETHING. An appliance with no
        LibreOffice — which is the default, and CI — must not be queueing a task whose input
        does not exist; and a fixture that passed because nothing ran would be worse than no
        fixture."""
        response = _upload(db_session, _anchored_docx())
        drain_ingest_queue(db_session, max_tasks=200)
        node = DocumentRepository(db_session).get(response.document_id)

        assert RenditionRepository(db_session).get(node.id, KIND_PDF) is None
        offered = {spec.task_type for spec in IngestRollupPlanner()(db_session, node)}
        assert TASK_ANCHOR_RENDER not in offered

    def test_a_document_with_a_rendition_is_offered_the_task(
        self, db_session, rendition_pdf, monkeypatch
    ):
        response = _upload(db_session, _anchored_docx())
        drain_ingest_queue(db_session, max_tasks=200)
        node = DocumentRepository(db_session).get(response.document_id)
        _install_rendition(db_session, node, rendition_pdf, monkeypatch)

        specs = IngestRollupPlanner()(db_session, node)
        assert [spec.task_type for spec in specs] == [TASK_ANCHOR_RENDER]
        spec = specs[0]
        assert spec.write_mode == WRITE_SUBTREE
        # 6.1's diff key. The hash is what moves when the rendition is replaced — which is
        # what `SPRINT_0_6_0.md` 4.9 says happens after a LibreOffice upgrade — so a new
        # pagination is re-anchored and an unchanged one is not.
        assert spec.params["rendition_hash"]

    def test_it_is_not_offered_twice_for_the_same_rendition(
        self, db_session, anchored, rendition_pdf
    ):
        """The walk revisits a node many times. The attempt log is the diff's authority, so a
        completed run is not offered again — and neither is a failed one, which is why this
        cannot become a loop on an appliance where the alignment finds nothing."""
        node, _ = anchored
        offered = {spec.task_type for spec in IngestRollupPlanner()(db_session, node)}
        assert TASK_ANCHOR_RENDER not in offered

    def test_the_document_settles_on_an_appliance_that_cannot_render(self, db_session):
        """The hazard the whole placement decision is about, asserted end to end: with no
        converter, ``render:pdf`` fails permanently, nothing waits on it, and the ``.docx``
        settles with its text, its chunks and its exact addresses."""
        response = _upload(db_session, _anchored_docx())
        drain_ingest_queue(db_session, max_tasks=200)
        node = DocumentRepository(db_session).get(response.document_id)

        assert node.settled == SETTLED_SETTLED
        assert node.content
        assert _chunks(db_session, node.id)


# ---------------------------------------------------------------------------
# 7. What a viewer gets — the question this step had to answer
# ---------------------------------------------------------------------------


class TestWhatTheViewerIsGiven:
    def test_before_the_rendition_the_verb_names_the_exact_address_it_cannot_draw(self, db_session):
        """An office chunk IS addressed and still has no rectangle, and those are different
        facts from "no anchor". The message names the element."""
        response = _upload(db_session, _anchored_docx())
        drain_ingest_queue(db_session, max_tasks=200)
        node = DocumentRepository(db_session).get(response.document_id)
        chunk = _chunks(db_session, node.id)[0]

        with pytest.raises(RegionNotAddressable) as caught:
            DocumentService(db_session)._pdf_anchor(chunk.id)
        assert "word/document.xml" in str(caught.value)
        assert "anchor:render" in str(caught.value)

    def test_after_it_the_verb_draws_the_rectangle_and_says_which_pdf_it_is_in(
        self, db_session, evidence, anchored
    ):
        """THE ANSWER TO "WHICH ONE DOES A VIEWER DRAW". They are not competing answers to
        one question: the exact address cannot be drawn and the rectangle can, so the
        rectangle is drawn and the address is what the rectangle is a picture OF. Neither
        row is hidden — ``GET /documents/{id}/evidence`` serves both — and the ``of`` on the
        anchor is how the viewer knows to say "page 1 of our rendering" rather than "page 1"."""
        node, _ = anchored
        chunk = next(c for c in _chunks(db_session, node.id) if evidence(c).get(RENDER_ANCHOR_NAME))
        anchor = DocumentService(db_session)._pdf_anchor(chunk.id)
        assert isinstance(anchor, PdfAnchor)
        assert anchor.of == ANCHOR_OF_RENDITION
        # And the exact address is still there to be named beside it.
        assert evidence(chunk)[ANCHOR_NAME]["kind"] == ANCHOR_KIND_OOXML

    def test_a_passage_the_rendition_could_not_place_says_that_rather_than_nothing(
        self, db_session, evidence, anchored
    ):
        """Partial recovery is a completed task, so "we looked and could not find it" is an
        ordinary outcome and has to be reported as one. A viewer told only "no rectangle"
        would report a missing conversion that is not missing."""
        node, _ = anchored
        chunk = _chunks(db_session, node.id)[0]
        repository = EvidenceRepository(db_session)
        repository.delete(chunk.id, RENDER_ANCHOR_NAME)
        repository.write(
            chunk.id,
            RENDER_ANCHOR_UNRESOLVED_NAME,
            {"code": UNRESOLVED_NOT_IN_RENDITION, "reason": "the renderer reflowed it"},
        )
        db_session.flush()

        with pytest.raises(RegionNotAddressable) as caught:
            DocumentService(db_session)._pdf_anchor(chunk.id)
        assert "reflowed" in str(caught.value)
