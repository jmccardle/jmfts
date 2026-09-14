"""A workbook's rows must be findable by BM25. ``docs/SPRINT_0_6_0.md`` Block B step 7.

``docs/STRESS_CORPUS.md`` 4.4 measured the defect on a real corpus: **zero ``index:bm25``
tasks exist anywhere in a 21-workbook subtree**, and the ``xlsx`` index held 0 documents, 0
postings and 0 terms while 4,601 ``record`` nodes under it carried real text averaging 457
characters. A spreadsheet was retrievable by vector and MaxSim and by nothing else — absent
from its own collection's index and from ``universal``, so a full-text query over
"everything" silently excluded every spreadsheet. The index existed, answered, and returned
nothing.

**THREE CAUSES, AND THE THIRD IS THE ONE THAT DECIDES THE SHAPE OF THE FIX.** 4.4 named
two, both in one ``TASK_ROWS`` row: ``after_any`` named the three prose structure rungs and
not ``structure:sheets``, and ``requires=(HAS_TEXT_LAYER,)`` is a PDF property no workbook
can satisfy. 4.4c repaired exactly those two and measured the result — ``index:bm25 added 0
documents while the subtree holds 3 record nodes; it saw a subtree of 2`` — because **a
workbook's leaves are not written by its structure rung and every other format's are**.
``structure:sheets`` writes the ``sheet`` containers; ``profile:sheet`` and ``extract:sheet``
are scoped to those children, one rung further down. :class:`TestTheRungIsNotTheLeafWriter`
is that third cause, asserted from the tree rather than cited.

So the row cannot express the dependency — ``after``/``after_any`` order two rows on ONE
node (``ingest_tasks._check_task_rows`` rule 2) and a second row under the same name is
refused by rule 1 — and ``docs/MEASURE_BM25_BOUNDARY.md`` §5.2 is where the write moved
instead: to the settling boundary, over a set difference. ``tests/test_atom_declarations.py``
predicted that move in those words before it happened, as ``EXPECTED_DIVERGENCE``.

**Scope: ``record`` nodes, not ``cell`` nodes** (``docs/SPRINT_0_6_0.md`` question 4.3).
:class:`TestCellNodesAreHeldOut` is the half of that decision this file can assert without
building an over-window row; ``tests/test_profile_usetype.py`` is the pattern.

FIXTURES ARE BUILT BY ``openpyxl``, for ``tests/test_sheet_tasks.py``'s reason: the
hand-assembled packages in ``tests/corpus`` are minimal OOXML for the PROBER, which reads
the ZIP directory with the standard library, and a workbook rich enough to reach
``extract:sheet`` needs a real writer. Nothing binary is committed.
"""

from __future__ import annotations

import io

import pytest
from sqlalchemy import select

pytest.importorskip("openpyxl", reason="the office extra is not installed")

import openpyxl  # noqa: E402

import jmfts_core.ingest_tasks  # noqa: E402,F401  (import order; see the module cycle)
from jmfts_client.contracts.upload import UploadedFile  # noqa: E402
from jmfts_core.config import get_settings  # noqa: E402
from jmfts_core.models.document import (  # noqa: E402
    Document,
    SETTLED_SETTLED,
    USETYPE_CELL,
    USETYPE_RECORD,
    USETYPE_SHEET,
)
from jmfts_core.models.search_index import SearchIndex, SearchIndexEntry  # noqa: E402
from jmfts_core.repositories.document import DocumentRepository  # noqa: E402
from jmfts_core.repositories.search import SearchRepository  # noqa: E402
from jmfts_core.rollup_tasks import IngestRollupPlanner  # noqa: E402
from jmfts_core.services.ingest_service import IngestService  # noqa: E402
from jmfts_core.sheet_tasks import TASK_EXTRACT_SHEET, TASK_STRUCTURE_SHEETS  # noqa: E402
from tests.conftest import drain_ingest_queue  # noqa: E402

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

INDEX = "workbook-corpus"

#: One token that appears in exactly one cell of the fixture and nowhere else in the suite.
#: `docs/STRESS_CORPUS.md` 4.4d searched the reference corpus for `quadraphonic` for the same
#: reason — a term whose only possible source is a spreadsheet cell.
NEEDLE = "quadraphonic"


def _workbook_bytes() -> bytes:
    """Two sheets with a header row apiece, so ``INGEST_SPEC.md`` 8.3 finds ``header_row``.

    Small on purpose: every ``record`` node gets a real ``embed`` task on the way to
    settling, and the boundary this file is about is only reached once they all have.
    """
    workbook = openpyxl.Workbook()
    instruments = workbook.active
    instruments.title = "Instruments"
    instruments.append(["Designator", "Device", "Footprint"])
    instruments.append(["R59", "0603WAF4700T5E", "R0603"])
    instruments.append(["U12", f"{NEEDLE} decoder module", "SOIC16"])
    instruments.append(["C7", "Ceramic 100nF", "C0402"])

    census = workbook.create_sheet("Census")
    census.append(["Month", "Time", "Population"])
    census.append(["1975-01-01", 277, 214931])
    census.append(["1975-02-01", 278, 215502])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def workbook_bytes() -> bytes:
    return _workbook_bytes()


@pytest.fixture
def indexed_folder(db_session):
    """A folder that a BM25 index names as a root. ``INGEST_SPEC.md`` 11.5.

    Membership is a property of a SUBTREE ROOT, so a workbook is only ever a candidate for
    an index that already covers its ancestry. Without this the correct answer to every
    question below is "nothing was indexed", which is the answer a broken pipeline also
    gives — the fixture is what tells the two apart.
    """
    folder = DocumentRepository(db_session).create(
        title="An indexed folder", content=None, auto_embed=False
    )
    db_session.flush()
    search = SearchRepository(db_session)
    search.create_index(INDEX, description="SPRINT_0_6_0.md Block B step 7")
    assert search.add_root_to_index(INDEX, folder.id)
    db_session.flush()
    return folder


@pytest.fixture
def ingested(db_session, indexed_folder, workbook_bytes):
    """The workbook, uploaded under the indexed folder and drained dry. Returns its node.

    ``planner=IngestRollupPlanner()`` and not ``NO_ROLLUP``, because that is what
    ``jmfts-worker`` and ``IngestService`` both run (``ingest_worker.py:587``,
    ``services/ingest_service.py:317``) and because the boundary is the whole subject.
    """
    response = IngestService(db_session).upload_file(
        UploadedFile(data=workbook_bytes, filename="instruments.xlsx", content_type=XLSX_MIME),
        parent_id=indexed_folder.id,
    )
    drain_ingest_queue(db_session, planner=IngestRollupPlanner(), max_tasks=400)
    node = DocumentRepository(db_session).get(response.document_id)
    assert node.settled == SETTLED_SETTLED, "the queue did not drain the workbook"
    return node


def _index_entries(session) -> set[int]:
    index_id = session.execute(select(SearchIndex.id).where(SearchIndex.name == INDEX)).scalar_one()
    return set(
        session.execute(
            select(SearchIndexEntry.document_id).where(SearchIndexEntry.index_id == index_id)
        )
        .scalars()
        .all()
    )


def _subtree(session, node_id: int, usetype: str) -> list[Document]:
    return list(
        session.execute(
            select(Document)
            .where(Document.path.contains([node_id]))
            .where(Document.usetype == usetype)
            .order_by(Document.id)
        )
        .scalars()
        .all()
    )


# ---------------------------------------------------------------------------
# The entry condition
# ---------------------------------------------------------------------------


class TestAWorkbookReachesBM25:
    """Ingest one workbook, search BM25 for a string that is in a cell, get it back."""

    def test_a_cell_value_is_findable(self, db_session, ingested):
        records = _subtree(db_session, ingested.id, USETYPE_RECORD)
        assert records, "the fixture produced no record nodes; 8.3 found no header row"

        results = SearchRepository(db_session).bm25_search(NEEDLE, index_name=INDEX, limit=10)

        found = {r.document.id for r in results}
        assert found, f"BM25 returned nothing for {NEEDLE!r}, which is in a cell of this book"
        hit = [r for r in records if r.id in found]
        assert hit, (
            f"BM25 returned {found} and none of them is a record node of this workbook "
            f"({[r.id for r in records]})"
        )
        assert NEEDLE in (hit[0].content or "").lower()

    def test_every_record_with_text_joined_the_index(self, db_session, ingested):
        """Not only the one the needle is in. A partial index answers, and answers wrongly."""
        records = [r for r in _subtree(db_session, ingested.id, USETYPE_RECORD) if r.content]
        entries = _index_entries(db_session)

        missing = [r.id for r in records if r.id not in entries]
        assert not missing, f"{len(missing)} of {len(records)} record nodes carry no entry"

    def test_nothing_was_written_to_an_index_that_does_not_cover_the_workbook(
        self, db_session, ingested
    ):
        """11.5: membership is READ from the tree, never chosen by the task. The
        boundary rule is offered wherever the walk goes, so "which index" has to keep
        coming from ``covering_indexes`` and not from the set of indexes that exist."""
        records = {r.id for r in _subtree(db_session, ingested.id, USETYPE_RECORD)}
        elsewhere = [
            index_id
            for index_id, in db_session.execute(
                select(SearchIndexEntry.index_id).where(SearchIndexEntry.document_id.in_(records))
            ).all()
        ]
        covering = db_session.execute(
            select(SearchIndex.id).where(SearchIndex.name == INDEX)
        ).scalar_one()
        assert set(elsewhere) == {covering}


# ---------------------------------------------------------------------------
# The third cause, asserted rather than cited
# ---------------------------------------------------------------------------


class TestTheRungIsNotTheLeafWriter:
    """``docs/STRESS_CORPUS.md`` 4.4c: why repairing the two exclusions was not the fix.

    Every prose format's structure rung writes its own chunks, so "the rung completed" and
    "the subtree exists" are the same statement. A workbook's are two statements, and the
    row could only ever have waited for the first.
    """

    def test_the_structure_rung_wrote_sheets_and_the_leaves_came_from_below(
        self, db_session, ingested
    ):
        sheets = _subtree(db_session, ingested.id, USETYPE_SHEET)
        records = _subtree(db_session, ingested.id, USETYPE_RECORD)

        assert {s.produced_by for s in sheets} == {TASK_STRUCTURE_SHEETS}
        assert {r.produced_by for r in records} == {TASK_EXTRACT_SHEET}
        # And the leaves hang off the sheets, not off the file node: the row's scope was
        # the file node, so at the instant `structure:sheets` completed its subtree was the
        # file node and its sheets — 4.4c's "it saw a subtree of 2".
        assert {r.parent_id for r in records} == {s.id for s in sheets}


# ---------------------------------------------------------------------------
# Question 4.3's answer, in the one place that decides it
# ---------------------------------------------------------------------------


class TestCellNodesAreHeldOut:
    """``docs/SPRINT_0_6_0.md`` 4.3: ``record`` nodes reach BM25 and ``cell`` nodes do not.

    A ``cell`` node exists only for a row too long to embed whole (``39d5f20``), and it
    carries one column's share of the row's prose — ``"Designator: R59."``, from
    ``sheet_records.cell_content``. Admitting them puts a population of five-token documents
    into the inverted index and moves every IDF in the index it joins, which this codebase
    has measured going wrong twice.

    The decision lives in ``Settings.bm25_exclude_usetypes`` and nowhere else, which is
    where ``entity``, ``entities``, ``summary`` and ``derived`` already state the same kind
    of thing. That keeps the settling-boundary rule format-blind — it names no usetype, no
    rung and no format — and puts the exclusion at the one point that writes postings.
    """

    def test_cell_is_excluded_from_bm25(self):
        assert USETYPE_CELL in get_settings().bm25_exclude_usetypes

    def test_cell_is_still_retrievable_by_vector(self):
        """The exclusion is about the inverted index, not about the node. A split row's
        columns are embedded and answer vector search; only their postings are declined."""
        assert USETYPE_CELL not in get_settings().search_exclude_usetypes

    def test_record_is_excluded_from_neither(self):
        settings = get_settings()
        assert USETYPE_RECORD not in settings.bm25_exclude_usetypes
        assert USETYPE_RECORD not in settings.search_exclude_usetypes

    def test_index_document_refuses_a_cell_node(self, db_session):
        """The config is the decision and this is the code that honours it."""
        doc = Document(
            title="Instruments row 4 · Designator",
            content="Designator: R59.",
            usetype=USETYPE_CELL,
            settled=SETTLED_SETTLED,
        )
        db_session.add(doc)
        db_session.flush()

        repo = SearchRepository(db_session)
        repo.create_index("cell-exclusion")

        assert repo.index_document(doc.id, "cell-exclusion") is False
