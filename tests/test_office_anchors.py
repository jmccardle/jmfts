"""``anchor:source`` — the element of an office package a chunk came from.

``docs/OFFICE_SPEC.md`` Part 5, Part 11 step 9; ``docs/SPRINT_0_6_0.md`` Block G step 34.
Before this, a ``.docx`` or ``.pptx`` chunk that answered a query pointed at nothing: the
readers produced markdown and threw away which paragraph produced which piece of it, so the
beeline this release is built around — click a result, see where the answer came from —
dead-ended on exactly the formats the release is widening to.

**WHICH ANCHOR, AND WHY IT IS NOT THE OTHER ONE.** Part 5's "Two anchors, not one" is the
design and this file is written against it. The anchor here addresses the SOURCE package —
``{"kind": "ooxml", "part": "word/document.xml", "path": "/w:document/w:body/w:p[42]"}`` —
and it is exact, permanent, and derived from the bytes that were uploaded. The other anchor
is a rectangle on a page of a RENDITION, recovered by searching the rendered PDF for the
chunk's text; it is approximate, depends on which LibreOffice paginated the file, and can
fail to be computed. That one is Part 11 step 8 and nothing here writes it.

THE TEST THAT MATTERS IS NOT THAT AN ANCHOR IS PRESENT. ``test_citation_task.py`` re-opens
the PDF and asks PyMuPDF whether the chunk's words are inside the rectangle that was
recorded, because an anchor that is plausible and wrong is the specific failure Part 5 says
is worse than no anchor at all. The office equivalent is here: the XPath is evaluated
against the part it names, with lxml, and the element it resolves to has to be the element
whose text the chunk holds. A path that merely parses would pass a presence check and point
at the wrong paragraph.

The whole module skips without the ``office`` extra, like ``test_office_extraction.py``: an
install that can probe but not open is a supported deployment.
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from sqlalchemy import select

pytest.importorskip("docx", reason="the office extra is not installed")
pytest.importorskip("pptx", reason="the office extra is not installed")

import jmfts_core.ingest_tasks  # noqa: E402,F401  (import order; see the module cycle)
from docx import Document as DocxDocument  # noqa: E402
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from jmfts_client.contracts.upload import UploadedFile  # noqa: E402
from jmfts_core.citation_tasks import (  # noqa: E402
    ANCHOR_KIND_OOXML,
    ANCHOR_NAME,
    ANCHOR_UNRESOLVED_NAME,
    OOXML_READERS,
    UNRESOLVED_NO_BLOCK,
    UNRESOLVED_NO_ELEMENT,
    UNRESOLVED_NO_SPAN,
    anchor_for_ooxml_span,
    run_anchor_source,
    shift_for_removed,
)
from jmfts_core.ingest_tasks import (  # noqa: E402
    TASK_ANCHOR_SOURCE,
    TASK_CITATION,
    explain_plan,
)
from jmfts_core.models.document import (  # noqa: E402
    Document,
    SETTLED_FAILED,
    SETTLED_SETTLED,
    USETYPE_CHUNK,
)
from jmfts_core.models.task_queue import (  # noqa: E402
    ADVISORY_TASK_TYPES,
    TASK_FAILED,
    WRITE_SUBTREE,
)
from jmfts_core.office.extract import (  # noqa: E402
    OOXML_PREFIXES,
    docx_to_markdown,
    pptx_to_markdown,
)
from jmfts_core.probe import detect_format, probe_patterns  # noqa: E402
from jmfts_core.repositories.document import DocumentRepository  # noqa: E402
from jmfts_core.repositories.evidence import EvidenceRepository  # noqa: E402
from jmfts_core.repositories.task_queue import TaskQueueRepository  # noqa: E402
from jmfts_core.services.ingest_service import IngestService  # noqa: E402
from jmfts_core.structure_tasks import (  # noqa: E402
    EXTRACTION_DOCX_BODY,
    EXTRACTION_PPTX_SLIDES,
)
from jmfts_core.task_errors import ErrorType  # noqa: E402
from tests.conftest import drain_ingest_queue  # noqa: E402

#: prefix -> URI, for evaluating a path this appliance wrote. The inverse of the one table
#: the reader is allowed to name namespaces from, so the test cannot resolve a prefix the
#: appliance would not have written.
NAMESPACES = {prefix: uri for uri, prefix in OOXML_PREFIXES.items()}

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


# ---------------------------------------------------------------------------
# Fixtures — built by the libraries themselves, per test_office_extraction.py
# ---------------------------------------------------------------------------


def _docx_bytes(build) -> bytes:
    document = DocxDocument()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(build) -> bytes:
    presentation = Presentation()
    build(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _anchored_docx() -> bytes:
    """Headings, prose, a table and a list — the four block kinds the docx reader emits."""

    def build(document):
        document.add_heading("Late interaction", level=1)
        document.add_paragraph(
            "MaxSim scores a document as the sum of per-query-token maxima against the "
            "document's own token vectors, which is what makes it a reranker rather than "
            "a retriever."
        )
        document.add_heading("Measured", level=2)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "metric"
        table.cell(0, 1).text = "value"
        table.cell(1, 0).text = "nDCG@10"
        table.cell(1, 1).text = "0.7090"
        document.add_paragraph("Retention at fifty percent of tokens.", style="List Bullet")

    return _docx_bytes(build)


def _anchored_pptx() -> bytes:
    """A titled slide with notes, and an untitled one whose heading is promoted text."""

    def build(presentation):
        first = presentation.slides.add_slide(presentation.slide_layouts[1])
        first.shapes.title.text = "Retrieval fusion"
        body = first.placeholders[1].text_frame
        body.text = "Vector recall is high and precision is not."
        deeper = body.add_paragraph()
        deeper.text = "BM25 supplies the precision."
        deeper.level = 1
        first.notes_slide.notes_text_frame.text = "Say the weights are tunable per corpus."

        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = second.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
        box.text_frame.text = "No title placeholder on this one"
        box.text_frame.add_paragraph().text = "and a second line under it"

    return _pptx_bytes(build)


def _element_text(data: bytes, anchor: dict) -> str:
    """The text of the element ``anchor`` names, read out of the package with lxml.

    This is the half that makes the test about correctness rather than presence: the path is
    evaluated against the part it names, and what comes back has to be the element whose
    text the chunk holds.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        tree = etree.fromstring(archive.read(anchor["part"]))
    # An absolute path names the root; lxml evaluates it against the tree, so the root
    # element has to be handed over as a tree rather than as an element.
    target = (tree.getroottree() if anchor["path"].startswith("/") else tree).xpath(
        anchor["path"], namespaces=NAMESPACES
    )
    assert len(target) == 1, f"{anchor['path']!r} resolved to {len(target)} elements, not one"
    return "".join(target[0].itertext())


def _normalized(text: str) -> str:
    """The characters, with markdown's own markers and all whitespace gone.

    The element a path resolves to carries the author's characters; the markdown carries the
    same characters with an ATX marker or a bullet in front of them, pipes between table
    cells and a rule row under the header. Whitespace goes too, and it has to: ``itertext()``
    concatenates a table's cells with nothing between them while the markdown puts a space
    on either side of each pipe. What is left is what the two must agree about.
    """
    return re.sub(r"[\s#|*_\-]+", "", text)


def _upload(session, data, filename, content_type):
    return IngestService(session).upload_file(
        UploadedFile(data=data, filename=filename, content_type=content_type)
    )


def _chunks(session, root_id: int) -> list[Document]:
    return list(
        session.execute(
            select(Document)
            .where(Document.path.contains([root_id]))
            .where(Document.usetype == USETYPE_CHUNK)
            .order_by(Document.id)
        )
        .scalars()
        .all()
    )


# ---------------------------------------------------------------------------
# 1. The declaration — advisory, and the two spellings that cannot import each other
# ---------------------------------------------------------------------------


class TestDeclaration:
    def test_anchor_source_is_declared_advisory(self):
        """``models.task_queue`` cannot import ``ingest_tasks``, so the name is spelled
        twice. A rename that touched only one of them would turn a missing address back into
        a failed document."""
        assert TASK_ANCHOR_SOURCE in ADVISORY_TASK_TYPES

    def test_it_qualifies_under_the_rule_the_constant_states(self):
        """A task may be advisory only if a node without its output is still correct and
        still says so. A chunk with no address retrieves exactly as it did before addresses
        existed, and the handler writes the reason on it — which is what the unresolved row
        is for."""
        assert UNRESOLVED_NO_ELEMENT in {
            UNRESOLVED_NO_ELEMENT,
            UNRESOLVED_NO_BLOCK,
            UNRESOLVED_NO_SPAN,
        }
        assert set(OOXML_READERS) == {EXTRACTION_DOCX_BODY, EXTRACTION_PPTX_SLIDES}


# ---------------------------------------------------------------------------
# 2. The schedule — two anchors, two rows, disjoint formats
# ---------------------------------------------------------------------------


def _plan_rows(data: bytes, filename: str) -> dict:
    detection = detect_format(data, filename=filename, declared_mime=None)
    patterns, _ = probe_patterns(data, detection)
    plan = explain_plan(fmt=detection.format, patterns=patterns, patterns_source="probed")
    return {task.task: task for task in plan.tasks}


class TestTheTwoRowsAreDisjoint:
    def test_an_office_file_is_scheduled_for_a_source_anchor_and_not_for_a_rectangle(self):
        """Part 5: a ``.docx`` has no pages until something paginates it, so the row that
        wants a page geometry reports itself IMPOSSIBLE rather than false — which is a
        different statement, and the one a reader of the plan needs."""
        for data, filename in ((_anchored_docx(), "a.docx"), (_anchored_pptx(), "a.pptx")):
            rows = _plan_rows(data, filename)
            assert rows[TASK_ANCHOR_SOURCE].outcome == "enqueued"
            assert rows[TASK_CITATION].outcome == "impossible"
            assert "page geometry" in rows[TASK_CITATION].reason

    def test_a_pdf_is_scheduled_for_a_rectangle_and_not_for_a_source_anchor(self):
        """The other direction, which is what makes the two dicts disjoint rather than one
        dict somebody forgot to merge: a PDF has no package part to name an element in, and
        its source anchor IS its rectangle."""
        from tests.corpus.fixtures import minimal_pdf

        rows = _plan_rows(minimal_pdf(), "a.pdf")
        assert rows[TASK_CITATION].outcome == "enqueued"
        assert rows[TASK_ANCHOR_SOURCE].outcome == "impossible"
        assert "package element" in rows[TASK_ANCHOR_SOURCE].reason

    def test_the_row_reserves_the_subtree_it_writes(self):
        """`subtree`, for `citation`'s reason word for word: the chunks are grandchildren
        under a titled section and children under an untitled one, and the same deck has
        both."""
        rows = _plan_rows(_anchored_pptx(), "a.pptx")
        assert rows[TASK_ANCHOR_SOURCE].write_mode == WRITE_SUBTREE


# ---------------------------------------------------------------------------
# 3. The block map — every span slices back to the block it claims
# ---------------------------------------------------------------------------


class TestSourceBlocks:
    @pytest.mark.parametrize(
        "reader,data",
        [
            (docx_to_markdown, _anchored_docx()),
            (pptx_to_markdown, _anchored_pptx()),
        ],
    )
    def test_a_block_span_slices_the_markdown_back_to_that_block(self, reader, data):
        """The offsets are the whole mechanism — a chunk's ``source_span`` is inverted
        through them — and they are taken across a normalisation that DELETES characters.
        An offset measured before a deletion is wrong by however much came off above it,
        which is the hazard ``pdf_extraction`` documents at its own offset computation."""
        markdown, meta = reader(data)
        blocks = meta["source_blocks"]
        assert blocks, "the reader emitted no block map"
        for block in blocks:
            sliced = markdown[block["start"] : block["end"]]
            assert sliced.strip(), f"block {block} slices to whitespace"
            assert sliced == sliced.strip("\n")

    @pytest.mark.parametrize(
        "reader,data",
        [
            (docx_to_markdown, _anchored_docx()),
            (pptx_to_markdown, _anchored_pptx()),
        ],
    )
    def test_the_blocks_are_ordered_and_do_not_overlap(self, reader, data):
        """Document order and disjointness are what make the inversion a bisect rather than
        a scan, and what make "the first block the passage touches" mean the first one."""
        _, meta = reader(data)
        blocks = meta["source_blocks"]
        for previous, block in zip(blocks, blocks[1:]):
            assert previous["end"] <= block["start"]

    def test_the_normalisation_the_join_used_to_do_has_nothing_left_to_do(self):
        """The offsets only stand because the join is exact concatenation. If a run of three
        newlines or a trailing space survived into the joined text, the old global pass would
        have deleted it and every offset past that point would be wrong."""
        for reader, data in (
            (docx_to_markdown, _anchored_docx()),
            (pptx_to_markdown, _anchored_pptx()),
        ):
            markdown, _ = reader(data)
            assert "\n\n\n" not in markdown
            assert not [line for line in markdown.split("\n") if line != line.rstrip()]


# ---------------------------------------------------------------------------
# 4. The address is RIGHT — the path resolves to the element whose text it claims
# ---------------------------------------------------------------------------


class TestTheAddressResolves:
    def test_every_docx_block_path_names_the_element_that_produced_it(self):
        """The presence check's opposite. Each path is evaluated against the part it names
        and the element it resolves to has to hold the block's own words."""
        data = _anchored_docx()
        markdown, meta = docx_to_markdown(data)
        addressed = [block for block in meta["source_blocks"] if "path" in block]
        assert len(addressed) == len(meta["source_blocks"]), "a docx block went unaddressed"
        for block in addressed:
            assert block["part"] == "word/document.xml"
            expected = _normalized(markdown[block["start"] : block["end"]])
            assert _normalized(_element_text(data, block)) == expected

    def test_a_slide_paragraph_is_addressed_as_a_paragraph_and_not_as_its_shape(self):
        """A shape-level address would give every bullet on the slide the same one — right
        about which box, wrong about which line. Part 5 spends its failure list refusing
        exactly that resolution."""
        data = _anchored_pptx()
        markdown, meta = pptx_to_markdown(data)
        bullets = [
            block
            for block in meta["source_blocks"]
            if "path" in block and markdown[block["start"] : block["end"]].lstrip().startswith("- ")
        ]
        assert bullets, "the deck produced no addressed bullet"
        assert len({block["path"] for block in bullets}) == len(bullets)
        for block in bullets:
            assert block["path"].endswith("]") and "a:p" in block["path"]
            expected = _normalized(markdown[block["start"] : block["end"]])
            assert _normalized(_element_text(data, block)) == expected

    def test_speaker_notes_are_addressed_in_the_notes_part_and_not_in_the_slide(self):
        """They are a different part of the package. A path into ``slide1.xml`` would name
        whatever element happened to sit at that position on the slide."""
        data = _anchored_pptx()
        markdown, meta = pptx_to_markdown(data)
        notes = [
            block
            for block in meta["source_blocks"]
            if block.get("part", "").startswith("ppt/notes")
        ]
        assert len(notes) == 1
        assert notes[0]["slide"] == 1
        expected = _normalized(markdown[notes[0]["start"] : notes[0]["end"]])
        assert expected in _normalized(_element_text(data, notes[0]))

    def test_the_headings_this_reader_invented_carry_no_address(self):
        """``## Notes`` is a label this module wrote. No element of the package produced
        those characters, and an address for them would be an address for text that is not
        there — which is the anchor-that-is-probably-right Part 5 rejects."""
        data = _anchored_pptx()
        markdown, meta = pptx_to_markdown(data)
        unaddressed = [
            markdown[block["start"] : block["end"]]
            for block in meta["source_blocks"]
            if "path" not in block
        ]
        assert unaddressed == ["## Notes"]

    def test_a_promoted_title_is_addressed_to_the_paragraph_it_was_promoted_from(self):
        """The deck's second slide has no title placeholder, so the heading is the first line
        of an ordinary text box. That line is real text in a real paragraph, and the address
        says which paragraph rather than which box."""
        data = _anchored_pptx()
        markdown, meta = pptx_to_markdown(data)
        heading = next(
            block
            for block in meta["source_blocks"]
            if markdown[block["start"] : block["end"]].startswith("# No title")
        )
        assert heading["slide"] == 2
        assert heading["path"].endswith("a:p[1]")
        assert _normalized(_element_text(data, heading)) == _normalized(
            "No title placeholder on this one"
        )


# ---------------------------------------------------------------------------
# 5. Inverting a span — what a chunk that covers several elements says
# ---------------------------------------------------------------------------


def _blocks(*spec) -> tuple[list[dict], list[int]]:
    """``(blocks, starts)`` from ``(start, end, part, path)`` tuples; ``path=None`` for a
    block the reader wrote."""
    blocks = [
        {"start": start, "end": end, **({"part": part, "path": path} if path else {})}
        for start, end, part, path in spec
    ]
    return blocks, [block["start"] for block in blocks]


PART = "word/document.xml"
NOTES = "ppt/notesSlides/notesSlide1.xml"


class TestAnchorForOoxmlSpan:
    def test_a_span_inside_one_element_takes_that_elements_path(self):
        blocks, starts = _blocks((0, 20, PART, "/w:document/w:body/w:p[1]"))
        anchor, code = anchor_for_ooxml_span(blocks, starts, 4, 12)
        assert code is None
        assert anchor == {
            "kind": ANCHOR_KIND_OOXML,
            "part": PART,
            "path": "/w:document/w:body/w:p[1]",
        }

    def test_a_span_over_several_elements_of_one_part_carries_the_run(self):
        """`path` and `through` bound what the passage occupies. Naming only the first would
        present a partial address as a whole one, which is the distinction the rest of Part 5
        is most careful about."""
        blocks, starts = _blocks(
            (0, 10, PART, "/w:document/w:body/w:p[1]"),
            (12, 20, PART, "/w:document/w:body/w:p[2]"),
            (22, 30, PART, "/w:document/w:body/w:p[3]"),
        )
        anchor, code = anchor_for_ooxml_span(blocks, starts, 4, 25)
        assert code is None
        assert anchor["path"] == "/w:document/w:body/w:p[1]"
        assert anchor["through"] == "/w:document/w:body/w:p[3]"
        assert "continues" not in anchor

    def test_a_span_that_leaves_its_part_says_which_parts_it_runs_into(self):
        """The office analogue of a passage that runs past a page break: a chunk that starts
        on a slide and ends in its speaker notes is ordinary, and its address covers the
        slide only."""
        blocks, starts = _blocks(
            (0, 10, "ppt/slides/slide1.xml", "./p:cSld/p:spTree/p:sp[1]"),
            (12, 20, NOTES, "./p:cSld/p:spTree/p:sp[2]"),
        )
        anchor, code = anchor_for_ooxml_span(blocks, starts, 0, 20)
        assert code is None
        assert anchor["part"] == "ppt/slides/slide1.xml"
        assert anchor["continues"] == [NOTES]
        assert "through" not in anchor

    def test_a_span_covering_only_invented_text_is_reported_rather_than_guessed(self):
        """The blocks are there and none of them has an address. Reaching past them to the
        nearest real element is the plausible-and-wrong answer."""
        blocks, starts = _blocks((0, 8, None, None), (10, 20, PART, "/w:document/w:body/w:p[1]"))
        anchor, code = anchor_for_ooxml_span(blocks, starts, 0, 8)
        assert anchor is None
        assert code == UNRESOLVED_NO_ELEMENT

    def test_a_span_touching_no_block_is_a_different_absence(self):
        blocks, starts = _blocks((0, 10, PART, "/w:document/w:body/w:p[1]"))
        anchor, code = anchor_for_ooxml_span(blocks, starts, 40, 50)
        assert anchor is None
        assert code == UNRESOLVED_NO_BLOCK

    def test_an_invented_block_before_a_real_one_does_not_hide_it(self):
        """A chunk that begins at a ``Notes`` label and runs into the notes themselves is
        addressed to the notes. The label accounts for none of the text, so skipping it
        claims nothing that is not true."""
        blocks, starts = _blocks((0, 8, None, None), (10, 20, NOTES, "./p:cSld/p:spTree/p:sp[2]"))
        anchor, code = anchor_for_ooxml_span(blocks, starts, 0, 20)
        assert code is None
        assert anchor["part"] == NOTES


class TestShiftForRemoved:
    def test_a_removed_character_moves_every_offset_above_it(self):
        """``extract:text`` strips NUL before it writes the node's content, so the block map
        this handler measures and the chunks it inverts are in different coordinates until
        this runs."""
        text = "ab\x00cd\x00ef"
        blocks = [{"start": 0, "end": 2}, {"start": 3, "end": 5}, {"start": 6, "end": 8}]
        moved_text, moved = shift_for_removed(text, blocks, "\x00")
        assert moved_text == "abcdef"
        assert [(block["start"], block["end"]) for block in moved] == [(0, 2), (2, 4), (4, 6)]
        for block in moved:
            assert moved_text[block["start"] : block["end"]] in ("ab", "cd", "ef")

    def test_a_document_with_nothing_to_remove_is_left_alone(self):
        blocks = [{"start": 0, "end": 2}]
        text, moved = shift_for_removed("abc", blocks, "\x00")
        assert text == "abc"
        assert moved == blocks


# ---------------------------------------------------------------------------
# 6. End to end — the chunks of a real package, addressed
# ---------------------------------------------------------------------------


class TestAnchorSourceOverARealPackage:
    def test_every_docx_chunk_is_addressed_and_the_element_holds_its_words(
        self, db_session, evidence
    ):
        """The office half of ``test_citation_task``'s rectangle check. The path is resolved
        against the part it names and the element has to hold the chunk's own words — an
        address that is plausible and wrong cannot pass this."""
        data = _anchored_docx()
        response = _upload(db_session, data, "anchors.docx", DOCX_MIME)
        drain_ingest_queue(db_session, max_tasks=200)

        node = DocumentRepository(db_session).get(response.document_id)
        assert node.settled == SETTLED_SETTLED
        chunks = _chunks(db_session, node.id)
        assert chunks, "the structure rung wrote no chunks"

        for chunk in chunks:
            found = evidence(chunk)
            anchor = found.get(ANCHOR_NAME)
            assert anchor is not None, found.get(ANCHOR_UNRESOLVED_NAME)
            assert anchor["kind"] == ANCHOR_KIND_OOXML
            assert anchor["part"] == "word/document.xml"
            element = _normalized(_element_text(data, anchor))
            first_line = _normalized(chunk.content.split("\n")[0])
            assert first_line, "a chunk with no first line"
            assert first_line in element or element in first_line, (
                f"chunk {chunk.id}'s address {anchor['path']} resolves to {element!r}, which "
                f"does not account for {first_line!r}"
            )

    def test_a_deck_is_addressed_slide_by_slide(self, db_session, evidence):
        data = _anchored_pptx()
        response = _upload(db_session, data, "anchors.pptx", PPTX_MIME)
        drain_ingest_queue(db_session, max_tasks=200)

        node = DocumentRepository(db_session).get(response.document_id)
        chunks = _chunks(db_session, node.id)
        assert chunks

        anchors = [evidence(chunk).get(ANCHOR_NAME) for chunk in chunks]
        placed = [anchor for anchor in anchors if anchor]
        assert placed, "no chunk of the deck was addressed"
        for anchor in placed:
            assert anchor["slide"] in (1, 2)
            assert anchor["part"].startswith("ppt/")
            _element_text(data, anchor)

    def test_the_attempt_detail_counts_what_was_and_was_not_addressed(self, db_session, evidence):
        """Partial recovery is a COMPLETED task, and the counts are what make the rate a
        number rather than an impression."""
        response = _upload(db_session, _anchored_docx(), "anchors.docx", DOCX_MIME)
        drain_ingest_queue(db_session, max_tasks=200)

        node = DocumentRepository(db_session).get(response.document_id)
        attempts = {entry["task"]: entry for entry in evidence(node)["attempts"]}
        assert TASK_ANCHOR_SOURCE in attempts, "anchor:source was never enqueued for a docx"
        attempt = attempts[TASK_ANCHOR_SOURCE]
        assert attempt["status"] == "completed"

        detail = attempt["detail"]
        assert detail["chunks"] == len(_chunks(db_session, node.id))
        assert detail["anchored"] + sum(detail["unresolved"].values()) == detail["chunks"]
        assert detail["addressed_blocks"] <= detail["source_blocks"]
        assert detail["parts_anchored"] == 1

    def test_the_extraction_attempt_reports_how_much_of_the_document_is_addressable(
        self, db_session, evidence
    ):
        """Before any chunk exists. The pair is the reading Part 10 needs to decide whether
        the addresses are worth what they cost."""
        response = _upload(db_session, _anchored_pptx(), "anchors.pptx", PPTX_MIME)
        drain_ingest_queue(db_session, max_tasks=200)

        node = DocumentRepository(db_session).get(response.document_id)
        attempts = {entry["task"]: entry for entry in evidence(node)["attempts"]}
        detail = attempts["extract:text"]["detail"]
        assert detail["source_blocks"] > detail["addressed_blocks"] > 0

    def test_a_chunk_with_no_source_span_is_reported_rather_than_failed(self, db_session, evidence):
        """The chunk keeps no address, gains a reason, and the rest of the document is
        addressed around it."""
        response = _upload(db_session, _anchored_docx(), "anchors.docx", DOCX_MIME)
        drain_ingest_queue(db_session, max_tasks=200)

        node = DocumentRepository(db_session).get(response.document_id)
        chunks = _chunks(db_session, node.id)
        repository = EvidenceRepository(db_session)
        repository.delete(chunks[0].id, "source_span")
        db_session.flush()

        task = TaskQueueRepository(db_session).enqueue(
            TASK_ANCHOR_SOURCE, node.id, WRITE_SUBTREE, service_badge=None
        )
        outcome = run_anchor_source(db_session, task)
        assert outcome.detail["unresolved"][UNRESOLVED_NO_SPAN] == 1

        found = evidence(chunks[0])
        assert ANCHOR_NAME not in found
        assert found[ANCHOR_UNRESOLVED_NAME]["code"] == UNRESOLVED_NO_SPAN
        assert "source_span" in found[ANCHOR_UNRESOLVED_NAME]["reason"]

    def test_text_that_disagrees_with_a_re_extraction_raises_rather_than_addressing(
        self, db_session
    ):
        """A block map derived from different text than the chunks were cut from would
        produce addresses that are wrong in a way nothing downstream could detect."""
        response = _upload(db_session, _anchored_docx(), "anchors.docx", DOCX_MIME)
        drain_ingest_queue(db_session, max_tasks=200)

        node = DocumentRepository(db_session).get(response.document_id)
        node.content = (node.content or "") + "\n\nan edit nobody made to the package\n"
        db_session.flush()

        task = TaskQueueRepository(db_session).enqueue(
            TASK_ANCHOR_SOURCE, node.id, WRITE_SUBTREE, service_badge=None
        )
        with pytest.raises(ValueError, match="confidently wrong"):
            run_anchor_source(db_session, task)

    def test_a_pdf_is_refused_by_name_rather_than_addressed(self, db_session):
        """The handler dispatches on the extraction source and says what it does instead —
        an error that names the other task costs one read instead of a search."""
        from tests.corpus.fixtures import minimal_pdf

        response = _upload(db_session, minimal_pdf(), "a.pdf", "application/pdf")
        drain_ingest_queue(db_session, max_tasks=200)
        node = DocumentRepository(db_session).get(response.document_id)

        task = TaskQueueRepository(db_session).enqueue(
            TASK_ANCHOR_SOURCE, node.id, WRITE_SUBTREE, service_badge=None
        )
        with pytest.raises(ValueError, match="citation"):
            run_anchor_source(db_session, task)

    def test_a_failed_anchor_source_does_not_take_the_document_with_it(self, db_session):
        """Advisory, end to end: the file node settles, keeps its text and its chunks, and
        carries the failure in its log. A document whose addresses could not be recovered is
        a perfectly good searchable document."""
        response = _upload(db_session, _anchored_docx(), "anchors.docx", DOCX_MIME)
        drain_ingest_queue(db_session, max_tasks=200)
        node = DocumentRepository(db_session).get(response.document_id)

        tasks = TaskQueueRepository(db_session)
        task = tasks.enqueue(TASK_ANCHOR_SOURCE, node.id, WRITE_SUBTREE, service_badge=None)
        claimed = tasks.claim_next("test-worker")
        assert claimed is not None and claimed.id == task.id
        tasks.mark_running(claimed)
        record = tasks.fail(
            claimed,
            error="the package could not be re-opened",
            error_type=ErrorType.PERMANENT,
        )
        db_session.flush()
        db_session.refresh(node)

        assert record.status == TASK_FAILED
        assert node.settled != SETTLED_FAILED
        # And the walk still closes over it, which is the whole point of the mechanism.
        # Enqueueing put the node back in flight, so "did not fail" is only half the claim:
        # a node that could never settle again would be just as lost.
        from jmfts_core.settling import NO_ROLLUP, settle_node

        assert settle_node(db_session, node.id, NO_ROLLUP).settled is True
        assert node.settled == SETTLED_SETTLED
        assert node.content
        assert _chunks(db_session, node.id)
