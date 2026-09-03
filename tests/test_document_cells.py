"""``GET /documents/{id}/cells`` — a spreadsheet region. ``OFFICE_SPEC.md`` Part 7.

Four things are under test and they are different in kind.

The **address**: :func:`~jmfts_core.office.cells.parse_ref` turns an A1 string into a
rectangle, and refuses the three families of string that are not one. No workbook.

The **reader**: :func:`~jmfts_core.office.cells.read_region` reads that rectangle out of
the bytes, priced in cells, with the formulas a read-only load cannot carry. Bytes only,
no database.

The **choice of rectangle**: :func:`~jmfts_core.services.document_service._cells_bounds`,
which is the part with a decision in it — what the caller named, then the node's own
anchor, then the sheet's measured used range. Pure; no session and no reader.

The **verb**, end to end on the real tree: upload a workbook, drain the queue, and read a
region back off the sheet node the ingest wrote — including through the mounted route, so
that the query parameter binding and the four failure statuses are the ones a client sees.

FIXTURES ARE BUILT BY ``openpyxl``, for the reason ``tests/test_sheet_records.py`` gives:
``tests/corpus`` is hand-assembled minimal OOXML for the PROBER, and a workbook rich enough
to exercise a reader needs a writer.

**A formula written by openpyxl has no cached value.** ``deals["C5"] = "=C2+C3"`` writes
``<f>C2+C3</f>`` and leaves ``<v/>`` empty, because openpyxl does not evaluate. So a formula
cell here has a note and no value — which is exactly the case the region reader must not
drop, and it is why the reader tests a row's CELLS rather than its values.
"""

from __future__ import annotations

import datetime
import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

# The reader is the `office` extra. `dev` implies it (pyproject.toml), so the suite job
# runs this file and the base-install job skips it.
pytest.importorskip("openpyxl", reason="the office extra is not installed")

import openpyxl  # noqa: E402

import jmfts_core.ingest_tasks  # noqa: E402,F401  (import order; see the module cycle)
from jmfts_client.contracts.upload import UploadedFile  # noqa: E402
from jmfts_core.database import get_db  # noqa: E402
from jmfts_core.models.document import Document  # noqa: E402
from jmfts_core.office.cells import (  # noqa: E402
    CELLS_READ_MAX,
    SHEET_MAX_COL,
    SHEET_MAX_ROW,
    BadCellRef,
    CellRange,
    CellRegion,
    TooManyCells,
    parse_ref,
    read_region,
    render_region,
    used_range,
)
from jmfts_core.repositories.document import DocumentRepository  # noqa: E402
from jmfts_core.repositories.evidence import EvidenceRepository  # noqa: E402
from jmfts_core.rest.main import app  # noqa: E402
from jmfts_core.services.document_service import (  # noqa: E402
    ANCHOR_NAME,
    ANCHOR_KIND_CELLS,
    CELLS_REF_ANCHOR,
    CELLS_REF_REQUEST,
    CELLS_REF_USED_RANGE,
    DocumentService,
    NotASheetNode,
    SheetSourceUnavailable,
    _cells_bounds,
)
from jmfts_core.models.document import USETYPE_SHEET  # noqa: E402
from jmfts_core.services.ingest_service import IngestService  # noqa: E402
from tests.conftest import drain_ingest_queue  # noqa: E402

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _workbook_bytes() -> bytes:
    """One sheet of records, one sheet whose data does not start at A1, one empty sheet.

    ``Deals`` holds every type a cell can be, a blank row in the middle so that a row NUMBER
    cannot be confused with a position, a forced-text identifier and a formula. ``Offset``
    starts at C3 so that a region's first column is provably not column A. ``Notes`` is never
    written to, which is the sheet with no used range at all.
    """
    workbook = openpyxl.Workbook()
    deals = workbook.active
    deals.title = "Deals"
    deals.append(["Deal ID", "Account", "Value", "Close Date", "Won"])
    deals.append(["D-4471", "Northwind Freight", 128000, datetime.datetime(2026, 9, 30), True])
    deals.append(["D-4472", "Contoso", 96500.5, datetime.datetime(2026, 10, 15), False])
    deals.append([None, None, None, None, None])
    deals["A5"] = "0012345"
    # The apostrophe Excel shows in the formula bar. A style flag, not part of the value.
    deals["A5"].quotePrefix = True
    deals["B5"] = "Fabrikam"
    deals["C5"] = "=C2+C3"
    deals["E5"] = True

    offset = workbook.create_sheet("Offset")
    offset["C3"] = "Region"
    offset["D3"] = "Owner"
    offset["C4"] = "Northeast"
    offset["D4"] = "Ap | Kaur"

    workbook.create_sheet("Notes")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def workbook_bytes() -> bytes:
    return _workbook_bytes()


def _upload(session, data: bytes, filename: str = "cells.xlsx"):
    return IngestService(session).upload_file(
        UploadedFile(data=data, filename=filename, content_type=XLSX_MIME)
    )


def _ingest(session, data: bytes, filename: str = "cells.xlsx") -> Document:
    response = _upload(session, data, filename)
    drain_ingest_queue(session)
    return DocumentRepository(session).get(response.document_id)


def _children(session, node_id: int, usetype: str) -> list:
    return list(
        session.execute(
            select(Document)
            .where(Document.parent_id == node_id, Document.usetype == usetype)
            .order_by(Document.id)
        )
        .scalars()
        .all()
    )


def _sheet_node(session, file_node: Document, name: str) -> Document:
    repo = EvidenceRepository(session)
    return next(
        node
        for node in _children(session, file_node.id, USETYPE_SHEET)
        if (repo.read(node.id, "sheet") or {}).get("name") == name
    )


def _row(response, number: int) -> list:
    return next(row.values for row in response.rows if row.row == number)


@pytest.fixture
def client_with_db(db_session):
    """TestClient bound to the savepoint-wrapped session (the house pattern)."""

    def _override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    from tests.conftest import AUTH_HEADERS

    client = TestClient(app, headers=AUTH_HEADERS)
    yield client
    app.dependency_overrides.pop(get_db, None)


# ---------------------------------------------------------------------------
# The address
# ---------------------------------------------------------------------------


class TestParseRef:
    def test_a_rectangle_parses_to_its_four_bounds(self):
        assert parse_ref("B4:H120") == CellRange(min_row=4, min_col=2, max_row=120, max_col=8)

    def test_a_single_cell_is_a_one_by_one_rectangle(self):
        assert parse_ref("B4") == CellRange(min_row=4, min_col=2, max_row=4, max_col=2)

    def test_the_ends_are_put_in_order(self):
        """`H120:B4` names the same rectangle. The pair is unordered, not malformed."""
        assert parse_ref("H120:B4") == parse_ref("B4:H120")

    def test_absolute_markers_are_dropped(self):
        """`$B$4` and `B4` are the same cell; the `$` is what Excel writes into a copy."""
        assert parse_ref("$B$4:$H$120") == parse_ref("B4:H120")

    def test_lowercase_is_the_same_address(self):
        assert parse_ref("b4:h120") == parse_ref("B4:H120")

    def test_the_served_ref_is_always_the_two_ended_form(self):
        """So a caller comparing what it asked for against what it got need not normalise."""
        assert parse_ref("B4").ref == "B4:B4"
        assert parse_ref("H120:B4").ref == "B4:H120"

    def test_a_whole_column_reference_is_refused(self):
        """`B:H` means `as far as the sheet goes`, which is a bound the request did not
        state. Resolving it against the used range would make one string name different
        rectangles on the same sheet at different times."""
        with pytest.raises(BadCellRef, match="Whole-column"):
            parse_ref("B:H")

    def test_a_whole_row_reference_is_refused(self):
        with pytest.raises(BadCellRef, match="not a cell reference"):
            parse_ref("4:120")

    def test_a_sheet_qualified_reference_is_refused(self):
        """The sheet is the document being read. A second opinion about it would let a
        caller address any sheet of the workbook through a node that names one."""
        with pytest.raises(BadCellRef, match="names a sheet"):
            parse_ref("Deals!B4:H120")

    def test_a_typo_is_refused_rather_than_read_as_an_empty_region(self):
        with pytest.raises(BadCellRef):
            parse_ref("B4:H12O")

    def test_an_empty_reference_is_refused(self):
        with pytest.raises(BadCellRef, match="names no region"):
            parse_ref("   ")

    def test_a_row_past_the_formats_limit_is_refused(self):
        with pytest.raises(BadCellRef, match="rows 1 to"):
            parse_ref(f"A1:A{SHEET_MAX_ROW + 1}")

    def test_a_column_past_the_formats_limit_is_refused(self):
        with pytest.raises(BadCellRef, match="columns A to"):
            parse_ref("A1:ZZZ1")

    def test_the_last_addressable_cell_is_addressable(self):
        """The limits are inclusive: XFD1048576 is a cell, and refusing it would be off by
        one in the direction that hides a real region."""
        bounds = parse_ref(f"XFD{SHEET_MAX_ROW}")
        assert (bounds.max_row, bounds.max_col) == (SHEET_MAX_ROW, SHEET_MAX_COL)

    def test_the_area_is_what_the_limit_is_measured_against(self):
        assert parse_ref("B4:H120").cells == 7 * 117


class TestUsedRange:
    def test_it_is_a1_to_the_last_measured_cell(self):
        assert used_range(5, 3) == CellRange(min_row=1, min_col=1, max_row=5, max_col=3)
        assert used_range(5, 3).ref == "A1:C5"

    def test_a_sheet_measured_to_hold_nothing_has_no_used_range(self):
        """`None`, not `A1:A1`. A 1x1 rectangle over a cell that holds nothing would answer
        a question about the sheet with a shape invented here."""
        assert used_range(0, 0) is None


# ---------------------------------------------------------------------------
# The reader
# ---------------------------------------------------------------------------


class TestReadRegion:
    def test_values_come_back_typed_and_in_region_order(self, workbook_bytes):
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("A1:E3"), max_cells=CELLS_READ_MAX
        )
        row = next(row for row in region.rows if row.index == 2)

        assert row.values == ("D-4471", "Northwind Freight", 128000, "2026-09-30", True)

    def test_a_region_that_does_not_start_at_a1_is_positioned_from_its_own_first_column(
        self, workbook_bytes
    ):
        """`values[0]` is the REGION's first column. It is column A only when the region
        starts there, which is the one place this differs from `read_rows`."""
        region = read_region(
            workbook_bytes, "Offset", bounds=parse_ref("C3:D4"), max_cells=CELLS_READ_MAX
        )

        assert [row.values for row in region.rows] == [
            ("Region", "Owner"),
            ("Northeast", "Ap | Kaur"),
        ]

    def test_a_region_is_clipped_to_its_columns(self, workbook_bytes):
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("B2:C2"), max_cells=CELLS_READ_MAX
        )

        assert [row.values for row in region.rows] == [("Northwind Freight", 128000)]

    def test_a_row_holding_nothing_yields_nothing_and_the_numbers_survive(self, workbook_bytes):
        """Row 4 is blank. It is absent rather than present-and-empty, and the rows around
        it keep the numbers Excel shows down the left edge."""
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("A1:E5"), max_cells=CELLS_READ_MAX
        )

        assert [row.index for row in region.rows] == [1, 2, 3, 5]

    def test_a_formula_cell_keeps_its_row_even_with_no_cached_value(self, workbook_bytes):
        """C5 is `=C2+C3` and openpyxl cached no result, so the value pass sees `None`. A
        row held to its values alone would drop the note that says the formula is there."""
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("C5:C5"), max_cells=CELLS_READ_MAX
        )

        assert [row.index for row in region.rows] == [5]
        assert region.rows[0].values == (None,)
        assert region.rows[0].notes[3].formula == "C2+C3"

    def test_notes_are_keyed_by_absolute_column(self, workbook_bytes):
        """So a note's address does not depend on where the rectangle happens to start."""
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("B5:E5"), max_cells=CELLS_READ_MAX
        )

        assert set(region.rows[0].notes) == {3}

    def test_a_forced_text_cell_is_reported_as_one(self, workbook_bytes):
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("A5:A5"), max_cells=CELLS_READ_MAX
        )

        assert region.rows[0].values == ("0012345",)
        assert region.rows[0].notes[1].text_forced is True

    def test_notes_outside_the_rectangle_are_not_reported(self, workbook_bytes):
        """The note pass walks past the rows above `min_row` to reach the region. What it
        picked up on the way belongs to cells nobody asked for."""
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("D1:E5"), max_cells=CELLS_READ_MAX
        )

        assert all(not row.notes for row in region.rows)

    def test_with_notes_false_says_that_nothing_looked(self, workbook_bytes):
        region = read_region(
            workbook_bytes,
            "Deals",
            bounds=parse_ref("A1:E5"),
            max_cells=CELLS_READ_MAX,
            with_notes=False,
        )

        assert region.notes_read is False
        assert all(not row.notes for row in region.rows)

    def test_a_region_beyond_the_sheets_data_holds_no_rows(self, workbook_bytes):
        """An empty rectangle is a fact about the sheet, and the ref is echoed beside it."""
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("A40:C60"), max_cells=CELLS_READ_MAX
        )

        assert region.rows == ()
        assert region.bounds.ref == "A40:C60"

    def test_a_region_of_an_empty_sheet_holds_no_rows(self, workbook_bytes):
        region = read_region(
            workbook_bytes, "Notes", bounds=parse_ref("A1:C5"), max_cells=CELLS_READ_MAX
        )

        assert region.rows == ()

    def test_a_region_too_large_fails_and_names_the_limit(self, workbook_bytes):
        """6.6's rule in the unit a region is priced in: a named limit that FAILS, not a
        silent truncation."""
        with pytest.raises(TooManyCells, match="CELLS_READ_MAX"):
            read_region(
                workbook_bytes,
                "Deals",
                bounds=parse_ref("A1:XFD400"),
                max_cells=CELLS_READ_MAX,
            )

    def test_the_limit_is_checked_before_the_package_is_opened(self):
        """Which is why a million-cell reference costs one multiplication. The bytes here
        are not a workbook at all; reaching openpyxl would raise something else."""
        with pytest.raises(TooManyCells):
            read_region(b"not a workbook", "Deals", bounds=parse_ref("A1:Z100000"), max_cells=1000)

    def test_a_caller_cannot_ask_for_more_than_the_readers_ceiling(self, workbook_bytes):
        with pytest.raises(ValueError, match="CELLS_READ_MAX"):
            read_region(
                workbook_bytes,
                "Deals",
                bounds=parse_ref("A1:B2"),
                max_cells=CELLS_READ_MAX + 1,
            )

    def test_an_unknown_sheet_raises(self, workbook_bytes):
        with pytest.raises(ValueError, match="names no sheet"):
            read_region(
                workbook_bytes, "Nowhere", bounds=parse_ref("A1:B2"), max_cells=CELLS_READ_MAX
            )


class TestRenderRegion:
    def test_it_renders_a_grid_addressed_the_way_the_sheet_is(self, workbook_bytes):
        """Column letters across the top and the worksheet row number down the side. The
        only rendering that stays honest when the rectangle does not start at A1."""
        region = read_region(
            workbook_bytes, "Offset", bounds=parse_ref("C3:D4"), max_cells=CELLS_READ_MAX
        )

        assert render_region(region).splitlines() == [
            "|  | C | D |",
            "| --- | --- | --- |",
            "| 3 | Region | Owner |",
            "| 4 | Northeast | Ap   Kaur |",
        ]

    def test_a_pipe_in_a_value_does_not_close_the_cell(self, workbook_bytes):
        region = read_region(
            workbook_bytes, "Offset", bounds=parse_ref("D4:D4"), max_cells=CELLS_READ_MAX
        )

        assert render_region(region).splitlines()[-1] == "| 4 | Ap   Kaur |"

    def test_a_region_holding_nothing_renders_no_row_claiming_to_be_one(self):
        region = CellRegion(sheet="Deals", bounds=parse_ref("A1:B2"), rows=(), notes_read=True)

        assert render_region(region).splitlines() == ["|  | A | B |", "| --- | --- | --- |"]

    def test_values_render_the_way_every_other_rendering_in_the_package_does(self, workbook_bytes):
        """TRUE, not True; 128000, not 128000.0. `canonical` is the one spelling."""
        region = read_region(
            workbook_bytes, "Deals", bounds=parse_ref("C2:E2"), max_cells=CELLS_READ_MAX
        )

        assert render_region(region).splitlines()[-1] == "| 2 | 128000 | 2026-09-30 | TRUE |"


# ---------------------------------------------------------------------------
# Which rectangle gets served
# ---------------------------------------------------------------------------


def _sheet_block(*, rows: int = 5, cols: int = 3, name: str = "Deals") -> dict:
    return {"name": name, "measurements": {"rows": rows, "cols": cols}}


class TestCellsBounds:
    def test_what_the_caller_named_wins(self):
        bounds, source = _cells_bounds(7, {}, _sheet_block(), "B2:C3")

        assert (bounds.ref, source) == ("B2:C3", CELLS_REF_REQUEST)

    def test_the_nodes_own_anchor_is_next(self):
        found = {ANCHOR_NAME: {"kind": ANCHOR_KIND_CELLS, "sheet": "Deals", "ref": "B4:H120"}}

        bounds, source = _cells_bounds(7, found, _sheet_block(), None)

        assert (bounds.ref, source) == ("B4:H120", CELLS_REF_ANCHOR)

    def test_the_measured_used_range_is_the_default(self):
        bounds, source = _cells_bounds(7, {}, _sheet_block(rows=5, cols=3), None)

        assert (bounds.ref, source) == ("A1:C5", CELLS_REF_USED_RANGE)

    def test_an_anchor_of_another_kind_is_an_error_and_not_an_ignored_anchor(self):
        """Something wrote an address for a region that is not a region of cells. Serving a
        different rectangle instead would hide it."""
        found = {ANCHOR_NAME: {"kind": "pdf", "page": 3, "bbox": [1, 2, 3, 4]}}

        with pytest.raises(SheetSourceUnavailable, match="kind 'pdf'"):
            _cells_bounds(7, found, _sheet_block(), None)

    def test_an_anchor_naming_another_sheet_is_an_error(self):
        found = {ANCHOR_NAME: {"kind": ANCHOR_KIND_CELLS, "sheet": "Lookup", "ref": "A1:B2"}}

        with pytest.raises(SheetSourceUnavailable, match="addresses sheet 'Lookup'"):
            _cells_bounds(7, found, _sheet_block(name="Deals"), None)

    def test_an_anchor_with_no_ref_is_an_error(self):
        found = {ANCHOR_NAME: {"kind": ANCHOR_KIND_CELLS, "sheet": "Deals"}}

        with pytest.raises(SheetSourceUnavailable, match="with no"):
            _cells_bounds(7, found, _sheet_block(), None)

    def test_an_unmeasured_sheet_has_no_default_and_says_which_task_is_missing(self):
        with pytest.raises(SheetSourceUnavailable, match="profile:sheet"):
            _cells_bounds(7, {}, {"name": "Deals"}, None)

    def test_a_sheet_measured_as_empty_has_no_default_and_names_the_way_round_it(self):
        with pytest.raises(SheetSourceUnavailable, match="`ref`"):
            _cells_bounds(7, {}, _sheet_block(rows=0, cols=0), None)

    def test_an_anchor_that_is_not_an_object_is_an_error(self):
        with pytest.raises(SheetSourceUnavailable, match="not an object"):
            _cells_bounds(7, {ANCHOR_NAME: "B4:H120"}, _sheet_block(), None)

    def test_a_caller_named_ref_is_still_parsed(self):
        with pytest.raises(BadCellRef):
            _cells_bounds(7, {}, _sheet_block(), "not a ref")


class TestConstantsAgreeWithTheirWriters:
    """The string this service spells rather than imports.

    Importing it would pull in ``jmfts_core.citation_tasks``, and through it
    ``jmfts_core.ingest_tasks``, whose module scope REGISTERS every task handler as a side
    effect — a read verb on the query path must not change what the worker dispatches
    merely by being imported. This is the guard that import would have been.

    THERE WERE TWO OF THESE. ``USETYPE_SHEET`` was the other, and it is gone: the ingest
    usetypes moved to ``jmfts_core.models.document`` in ``SPRINT_JOBS.md`` Phase 3 — Part
    4's rule table names node kinds and cannot import the handler modules — and this
    service already imports the model, which registers nothing. So the reasoning above was
    right about ``sheet_tasks`` and was never a reason to keep a second spelling of the
    string. The import is the guard now, and the test that stood in for it is deleted
    rather than left asserting that a name equals itself.
    """

    def test_the_anchor_name_is_the_one_citation_writes(self):
        from jmfts_core.citation_tasks import ANCHOR_NAME as WRITTEN

        assert ANCHOR_NAME == WRITTEN


# ---------------------------------------------------------------------------
# The verb, on the real tree
# ---------------------------------------------------------------------------


class TestCellsVerb:
    def test_it_serves_the_measured_used_range_by_default(self, db_session, workbook_bytes):
        """The whole claim, from bytes to a table: upload, drain, and the cells are there."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        response = DocumentService(db_session).get_document_cells(sheet.id)

        assert response.sheet == "Deals"
        assert response.ref == "A1:E5"
        assert response.ref_source == CELLS_REF_USED_RANGE
        assert response.columns == ["A", "B", "C", "D", "E"]
        assert _row(response, 2) == [
            "D-4471",
            "Northwind Freight",
            128000,
            "2026-09-30",
            True,
        ]

    def test_the_priced_number_is_the_area_and_not_the_filled_cells(
        self, db_session, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        response = DocumentService(db_session).get_document_cells(sheet.id)

        assert response.cell_count == 25
        assert response.row_count == 4  # row 4 holds nothing

    def test_a_named_ref_serves_that_rectangle(self, db_session, workbook_bytes):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Offset")

        response = DocumentService(db_session).get_document_cells(sheet.id, ref="C3:D4")

        assert (response.ref, response.ref_source) == ("C3:D4", CELLS_REF_REQUEST)
        assert response.columns == ["C", "D"]
        assert _row(response, 4) == ["Northeast", "Ap | Kaur"]

    def test_formulas_come_back_keyed_by_cell(self, db_session, workbook_bytes):
        """The declared relationship. It is the one place in a workbook where the author
        states one outright, and a read-only load cannot carry it beside the value."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        response = DocumentService(db_session).get_document_cells(sheet.id)

        assert response.cells["C5"].formula == "C2+C3"
        assert response.cells["A5"].text_forced is True
        assert "B5" not in response.cells

    def test_the_markdown_is_the_region_as_a_grid(self, db_session, workbook_bytes):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Offset")

        response = DocumentService(db_session).get_document_cells(sheet.id, ref="C3:D4")

        assert response.markdown.splitlines()[0] == "|  | C | D |"
        assert response.markdown.splitlines()[-1] == "| 4 | Northeast | Ap   Kaur |"

    def test_a_node_anchor_is_served_when_the_caller_names_nothing(
        self, db_session, workbook_bytes
    ):
        """NOTHING IN THIS TREE WRITES A `cells` ANCHOR YET — Part 5 specifies where one
        lives and every writer of one is still unbuilt. The anchor is written here by hand
        so that the branch the spec's default names is covered the day one exists."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")
        EvidenceRepository(db_session).write(
            sheet.id,
            ANCHOR_NAME,
            {"kind": ANCHOR_KIND_CELLS, "sheet": "Deals", "ref": "A2:C3"},
        )
        db_session.flush()

        response = DocumentService(db_session).get_document_cells(sheet.id)

        assert (response.ref, response.ref_source) == ("A2:C3", CELLS_REF_ANCHOR)
        assert [row.row for row in response.rows] == [2, 3]

    def test_a_document_that_is_not_a_sheet_says_so(self, db_session, workbook_bytes):
        """Not an empty table. The caller cannot tell `this sheet is blank` from `you asked
        the wrong node` if both answer the same way."""
        file_node = _ingest(db_session, workbook_bytes)

        with pytest.raises(NotASheetNode, match="not a worksheet"):
            DocumentService(db_session).get_document_cells(file_node.id)

    def test_an_empty_sheet_has_no_used_range_to_serve(self, db_session, workbook_bytes):
        """`Notes` was never written to. It has no used range, so there is no region to
        default to, and the error names `ref` as the way to read one anyway."""
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Notes")

        with pytest.raises(SheetSourceUnavailable, match="no cells"):
            DocumentService(db_session).get_document_cells(sheet.id)

    def test_an_empty_sheet_still_answers_a_named_region(self, db_session, workbook_bytes):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Notes")

        response = DocumentService(db_session).get_document_cells(sheet.id, ref="A1:C5")

        assert (response.ref, response.rows) == ("A1:C5", [])

    def test_a_missing_document_is_a_lookup_error(self, db_session):
        with pytest.raises(LookupError):
            DocumentService(db_session).get_document_cells(10_000_000)

    def test_a_malformed_ref_is_refused(self, db_session, workbook_bytes):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        with pytest.raises(BadCellRef):
            DocumentService(db_session).get_document_cells(sheet.id, ref="B:H")

    def test_a_ref_naming_a_million_cells_is_refused_by_name(self, db_session, workbook_bytes):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        with pytest.raises(TooManyCells, match=f"{CELLS_READ_MAX:,}"):
            DocumentService(db_session).get_document_cells(sheet.id, ref="A1:XFD400")


# ---------------------------------------------------------------------------
# The wire
# ---------------------------------------------------------------------------


class TestCellsRoute:
    def test_the_route_serves_the_region_and_binds_ref_as_a_query_parameter(
        self, client_with_db, db_session, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Offset")

        response = client_with_db.get(f"/documents/{sheet.id}/cells", params={"ref": "C3:D4"})

        assert response.status_code == 200
        body = response.json()
        assert body["ref"] == "C3:D4"
        assert body["ref_source"] == CELLS_REF_REQUEST
        assert body["rows"] == [
            {"row": 3, "values": ["Region", "Owner"]},
            {"row": 4, "values": ["Northeast", "Ap | Kaur"]},
        ]

    def test_a_malformed_ref_is_a_400_that_names_the_problem(
        self, client_with_db, db_session, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        response = client_with_db.get(f"/documents/{sheet.id}/cells", params={"ref": "B:H"})

        assert response.status_code == 400
        assert "Whole-column" in response.json()["detail"]

    def test_an_over_large_region_is_a_413_that_names_the_limit(
        self, client_with_db, db_session, workbook_bytes
    ):
        file_node = _ingest(db_session, workbook_bytes)
        sheet = _sheet_node(db_session, file_node, "Deals")

        response = client_with_db.get(f"/documents/{sheet.id}/cells", params={"ref": "A1:XFD400"})

        assert response.status_code == 413
        assert f"{CELLS_READ_MAX:,}" in response.json()["detail"]

    def test_a_node_that_is_not_a_sheet_is_a_409(self, client_with_db, db_session, workbook_bytes):
        file_node = _ingest(db_session, workbook_bytes)

        response = client_with_db.get(f"/documents/{file_node.id}/cells")

        assert response.status_code == 409
        assert "not a worksheet" in response.json()["detail"]

    def test_a_missing_document_is_a_404(self, client_with_db):
        assert client_with_db.get("/documents/10000000/cells").status_code == 404
