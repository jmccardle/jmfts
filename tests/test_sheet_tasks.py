"""A workbook's declared rung is its sheet list. ``INGEST_SPEC.md`` 8.1.

Three things are under test and they are different in kind.

The **reader** (:mod:`jmfts_core.office.workbook`) turns bytes into the list the workbook
names, in workbook order, with the state each sheet declares. No database.

The **schedule** — that ``probe``'s ``has_sheets`` is what makes ``structure:sheets``
eligible, that no other format can ever reach the row, and that the two text rungs are
untouched by any of it. Also no database: ``explain_plan`` is a pure function of
``(format, patterns, options)`` and that is the property being relied on.

The **task**, end to end, on the real queue: upload, drain, and look at the tree. That is
the half that would notice a node written with the wrong usetype, an order that does not
match the workbook, or a sheet node left in flight with nothing able to settle it.

FIXTURES ARE BUILT BY ``openpyxl``, not taken from ``tests/corpus``. The corpus packages
are hand-assembled minimal OOXML for the PROBER, which reads the ZIP directory with the
standard library and needs no reader at all; that is a guarantee worth keeping, and a
workbook rich enough to exercise this one needs ``openpyxl`` to build. So it is built here,
and the whole module skips where the ``office`` extra is not installed.
"""

from __future__ import annotations

import io

import pytest
from sqlalchemy import select

# The reader is the `office` extra. `dev` implies it (pyproject.toml), so the suite job
# runs this file and the base-install job skips it.
pytest.importorskip("openpyxl", reason="the office extra is not installed")

import openpyxl  # noqa: E402

import jmfts_core.ingest_tasks  # noqa: E402,F401  (import order; see the module cycle)
from jmfts_client.contracts.upload import UploadedFile  # noqa: E402
from jmfts_core.ingest_tasks import (  # noqa: E402
    OUTCOME_ENQUEUED,
    OUTCOME_IMPOSSIBLE,
    TASK_STRUCTURE_DECLARED,
    TASK_STRUCTURE_INFERRED,
    TASK_STRUCTURE_SHEETS,
    explain_plan,
)
from jmfts_core.models.document import Document, SETTLED_SETTLED  # noqa: E402
from jmfts_core.office.workbook import read_sheets  # noqa: E402
from jmfts_core.probe import detect_format, probe_patterns  # noqa: E402
from jmfts_core.repositories.document import DocumentRepository  # noqa: E402
from jmfts_core.services.ingest_service import IngestService  # noqa: E402
from jmfts_core.sheet_tasks import (  # noqa: E402
    SOURCE_WORKBOOK_SHEETS,
    USETYPE_SHEET,
    run_structure_sheets,
)
from jmfts_core.structure_tasks import RUNG_DECLARED  # noqa: E402
from tests.conftest import drain_ingest_queue  # noqa: E402

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _workbook_bytes() -> bytes:
    """Four sheets, one of them hidden, and the hidden one is not last.

    The shape is the one 8.1 describes a real workbook as having — "a data table, a pivot
    summary, a lookup list, and a notes page" — and the hidden sheet sits in the middle so
    that an index taken from enumeration and an index taken from "the visible ones" cannot
    accidentally agree.
    """
    workbook = openpyxl.Workbook()
    deals = workbook.active
    deals.title = "Deals"
    deals.append(["Deal ID", "Account", "Value"])
    deals.append(["D-4471", "Northwind Freight", 128000])

    lookup = workbook.create_sheet("Lookup")
    lookup.sheet_state = "hidden"
    lookup.append(["Code", "Region"])
    lookup.append(["NW", "Northwest"])

    matrix = workbook.create_sheet("Coverage")
    matrix.append([None, "Chassis Rev C"])
    matrix.append(["Payload Handling", "X"])

    workbook.create_sheet("Notes")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def workbook_bytes() -> bytes:
    return _workbook_bytes()


def _upload(session, data: bytes, filename: str = "pipeline.xlsx"):
    return IngestService(session).upload_file(
        UploadedFile(data=data, filename=filename, content_type=XLSX_MIME)
    )


def _ingest(session, data: bytes, filename: str = "pipeline.xlsx") -> Document:
    """Upload and run the queue dry. Returns the file node."""
    response = _upload(session, data, filename)
    drain_ingest_queue(session)
    return DocumentRepository(session).get(response.document_id)


def _children(session, node_id: int) -> list[Document]:
    return list(
        session.execute(
            select(Document)
            .where(Document.parent_id == node_id)
            .order_by(Document.position, Document.id)
        )
        .scalars()
        .all()
    )


def _attempt(node: Document, task: str) -> dict:
    return next(e for e in node.structured_content["attempts"] if e["task"] == task)


class _ScopedTask:
    """The one field a handler reads off its queue row.

    A real ``TaskQueue`` would need a claim, a write mode and a transaction to go with it,
    and none of that is what the test below is about — it drives the handler directly to
    reach a contradiction the front door cannot produce.
    """

    def __init__(self, document_id: int):
        self.scope_document_id = document_id
        self.params: dict = {}


# ---------------------------------------------------------------------------
# The reader
# ---------------------------------------------------------------------------


class TestReadSheets:
    def test_it_returns_every_sheet_in_workbook_order(self, workbook_bytes):
        sheets = read_sheets(workbook_bytes)

        assert [sheet.name for sheet in sheets] == ["Deals", "Lookup", "Coverage", "Notes"]
        assert [sheet.index for sheet in sheets] == [0, 1, 2, 3]

    def test_a_hidden_sheet_is_still_a_sheet_and_says_so(self, workbook_bytes):
        """Hiding a lookup table is how workbooks are built. Dropping it would remove the
        half of the workbook the visible half references."""
        states = {sheet.name: sheet.state for sheet in read_sheets(workbook_bytes)}

        assert states == {
            "Deals": "visible",
            "Lookup": "hidden",
            "Coverage": "visible",
            "Notes": "visible",
        }

    def test_it_opens_the_workbook_read_only(self, workbook_bytes, monkeypatch):
        """Not a preference: Part 8 profiles sheets of a million rows, and without it
        openpyxl materialises every cell as a Python object before a name can be read."""
        seen: dict = {}
        real = openpyxl.load_workbook

        def spy(*args, **kwargs):
            seen.update(kwargs)
            return real(*args, **kwargs)

        monkeypatch.setattr(openpyxl, "load_workbook", spy)
        read_sheets(workbook_bytes)

        assert seen.get("read_only") is True

    def test_a_sheet_with_no_cells_is_reported_like_any_other(self, workbook_bytes):
        """`Notes` was never written to. It is still a sheet the workbook declares, and
        whether it holds anything is a MEASUREMENT — 8.3's business, not this rung's."""
        assert read_sheets(workbook_bytes)[3].name == "Notes"


# ---------------------------------------------------------------------------
# The schedule
# ---------------------------------------------------------------------------


def _plan_for(data: bytes, filename: str):
    """What the scheduler would do with these bytes, from the patterns probe measures."""
    detection = detect_format(data, filename=filename, declared_mime=None)
    patterns, _ = probe_patterns(data, detection)
    plan = explain_plan(fmt=detection.format, patterns=patterns, patterns_source="probed")
    return patterns, {task.task: task for task in plan.tasks}


class TestSchedule:
    def test_probe_reports_the_pattern_the_row_reads(self, workbook_bytes):
        """Tier 1 measures it, with no optional dependency, and this rung does not
        recompute it."""
        patterns, _ = _plan_for(workbook_bytes, "pipeline.xlsx")

        assert patterns["has_sheets"] is True
        assert patterns["sheet_count"] == 4

    def test_a_workbook_enqueues_the_sheet_rung(self, workbook_bytes):
        _, tasks = _plan_for(workbook_bytes, "pipeline.xlsx")

        assert tasks[TASK_STRUCTURE_SHEETS].outcome == OUTCOME_ENQUEUED
        assert tasks[TASK_STRUCTURE_SHEETS].requires == ("has_sheets",)

    def test_the_row_waits_for_nothing(self, workbook_bytes):
        """Its evidence is the workbook part, which is in the uploaded bytes. That is what
        makes it a different row from the two rungs that consume `extract:text`."""
        _, tasks = _plan_for(workbook_bytes, "pipeline.xlsx")

        assert tasks[TASK_STRUCTURE_SHEETS].after == ()
        assert tasks[TASK_STRUCTURE_SHEETS].after_any == ()

    def test_the_text_rungs_stay_out_of_it(self, workbook_bytes):
        """A workbook has no text layer and does not pretend to. Both rungs report the
        missing extraction rather than anything about sheets."""
        _, tasks = _plan_for(workbook_bytes, "pipeline.xlsx")

        for name in (TASK_STRUCTURE_DECLARED, TASK_STRUCTURE_INFERRED):
            assert "extract:text is not eligible" in tasks[name].reason

    @pytest.mark.parametrize("fmt", ["pdf", "docx", "pptx", "text", "html"])
    def test_no_other_format_can_ever_reach_the_row(self, fmt):
        """`impossible`, not `not_applicable`: no other format has a sheet-list pattern at
        all, so no bytes of one could satisfy the requirement. Asked of the format rather
        than of bytes, because that is the property 11.2 relies on."""
        plan = explain_plan(fmt=fmt, patterns={"has_text_layer": True}, patterns_source="supplied")
        row = next(task for task in plan.tasks if task.task == TASK_STRUCTURE_SHEETS)

        assert row.outcome == OUTCOME_IMPOSSIBLE
        assert row.reason == f"format {fmt!r} names no worksheet list a declared rung could read"
        # The sentinel resolves to nothing, so it is dropped rather than emitted as a null.
        assert row.requires == ()


# ---------------------------------------------------------------------------
# The task
# ---------------------------------------------------------------------------


class TestStructureSheets:
    def test_one_node_per_sheet_under_the_file_node(self, db_session, workbook_bytes):
        node = _ingest(db_session, workbook_bytes)
        children = _children(db_session, node.id)

        assert [child.title for child in children] == ["Deals", "Lookup", "Coverage", "Notes"]
        assert {child.usetype for child in children} == {USETYPE_SHEET}

    def test_the_declared_tree_stops_at_the_sheet(self, db_session, workbook_bytes):
        """8.1's whole claim. A sheet node has no content of its own: its cells are not its
        text, and rendering them here would be 8.4's representation decision taken by the
        rung the spec says makes none.

        What is BELOW a sheet is no longer nothing — `profile:sheet` writes one profile
        node per sheet (8.5) — and that is measured, not declared. `tests/test_sheet_profile.py`
        holds it."""
        node = _ingest(db_session, workbook_bytes)

        for child in _children(db_session, node.id):
            assert child.content is None

    def test_a_sheet_node_carries_its_name_index_and_state(self, db_session, workbook_bytes):
        """The three facts THIS rung writes. `profile:sheet` adds its measurements to the
        same block (8.3 stores them under `sheet.measurements`), so the assertion is on
        these three keys rather than on the whole block."""
        node = _ingest(db_session, workbook_bytes)
        sheets = [child.structured_content["sheet"] for child in _children(db_session, node.id)]
        declared = [{key: block[key] for key in ("index", "name", "state")} for block in sheets]

        assert declared == [
            {"index": 0, "name": "Deals", "state": "visible"},
            {"index": 1, "name": "Lookup", "state": "hidden"},
            {"index": 2, "name": "Coverage", "state": "visible"},
            {"index": 3, "name": "Notes", "state": "visible"},
        ]

    def test_a_sheet_node_carries_the_declared_rung(self, db_session, workbook_bytes):
        """The task name diverges from 8.2; the RUNG does not. A reader asking whether the
        declared rung ran for this workbook gets its answer from the node."""
        node = _ingest(db_session, workbook_bytes)
        structure = _children(db_session, node.id)[0].structured_content["structure"]

        assert structure == {
            "primary_rung": RUNG_DECLARED,
            "source": SOURCE_WORKBOOK_SHEETS,
        }

    def test_the_file_node_records_what_the_rung_built(self, db_session, workbook_bytes):
        node = _ingest(db_session, workbook_bytes)

        assert node.structured_content["structure"] == {
            "primary_rung": RUNG_DECLARED,
            "source": SOURCE_WORKBOOK_SHEETS,
            "node_count": 4,
            "max_depth": 1,
        }

    def test_a_sheet_node_is_settled_once_the_work_beneath_it_is(self, db_session, workbook_bytes):
        """This test carried a claim with a date on it, and the date has passed.

        Step 5 created a sheet node SETTLED because nothing was queued beneath it — a node
        left in flight with no task able to release it parks the whole workbook. Now
        `profile:sheet` is registered and is enqueued per sheet, so the node is created IN
        FLIGHT and the settle walk releases it, which is the contract every container has.
        The state after a full drain is the same either way, and that is what is asserted
        here; `tests/test_sheet_profile.py` asserts the state in between."""
        node = _ingest(db_session, workbook_bytes)

        assert all(child.settled == SETTLED_SETTLED for child in _children(db_session, node.id))
        assert node.settled == SETTLED_SETTLED

    def test_the_attempt_carries_the_probe_count_beside_its_own(self, db_session, workbook_bytes):
        """The paired measurement. Probe counted `<sheet>` elements with the standard
        library and openpyxl resolved each one to a part; two numbers that should agree and
        do not are what says the workbook holds something other than worksheets."""
        node = _ingest(db_session, workbook_bytes)
        detail = _attempt(node, TASK_STRUCTURE_SHEETS)["detail"]

        assert detail["sheets"] == 4
        assert detail["sheets_probed"] == 4
        assert detail["hidden_sheets"] == 1

    def test_both_of_8_2s_per_sheet_tasks_are_queued_in_dependency_order(
        self, db_session, workbook_bytes
    ):
        """3.4: a rung that ran and stopped where the spec says to stop must not look like
        one that never ran — and the converse, which is what this now asserts.

        Both of 8.2's tasks have handlers. `extract:sheet` runs 8.4's `records` shape on
        `header_row`, which is a measured boolean rather than one of 8.8's unset
        thresholds, so it no longer waits on the calibration corpus. The order matters and
        is asserted: it reads the header labels off what `profile:sheet` stored."""
        node = _ingest(db_session, workbook_bytes)
        detail = _attempt(node, TASK_STRUCTURE_SHEETS)["detail"]

        assert detail["queued_per_sheet"] == ["profile:sheet", "extract:sheet"]
        assert detail["deferred"] == {}

    def test_an_empty_sheet_list_raises_rather_than_settling(
        self, db_session, workbook_bytes, monkeypatch
    ):
        """A contradiction the front door cannot produce, which is why it is driven here.

        `has_sheets` is what schedules the task and probe read it from the same
        `xl/workbook.xml` the reader does. The two disagreeing means the bytes changed
        underneath. Writing zero nodes and completing would settle a workbook as structured
        when nothing structured it.
        """
        node = _ingest(db_session, workbook_bytes)
        monkeypatch.setattr("jmfts_core.sheet_tasks.read_sheets", lambda data: [])

        with pytest.raises(ValueError) as raised:
            run_structure_sheets(db_session, _ScopedTask(node.id))

        assert "names no sheets" in str(raised.value)
        # The disagreement is quantified, not merely asserted.
        assert "probe counted 4" in str(raised.value)

    def test_a_missing_blob_raises(self, db_session, workbook_bytes):
        """The same failure `extract:text` has for the same reason: a task enqueued for
        bytes that are no longer there cannot complete quietly."""
        node = _ingest(db_session, workbook_bytes)
        orphan = DocumentRepository(db_session).create(title="no bytes", content=None)

        with pytest.raises(ValueError) as raised:
            run_structure_sheets(db_session, _ScopedTask(orphan.id))

        assert "has no stored blob" in str(raised.value)
        assert node.id != orphan.id
